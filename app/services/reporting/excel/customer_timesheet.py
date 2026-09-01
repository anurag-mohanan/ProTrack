"""Excel export for customer weekly/monthly timesheet packs."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.schemas.reporting import CustomerTimesheetPackPayload
from app.services.reporting.excel.letterhead import (
    set_print_layout,
    style_total_row,
    style_utilization_header,
    write_report_letterhead,
)
from app.services.reporting.excel.styles import (
    autofit_columns,
    freeze_and_filter,
    style_body_rows,
    style_header_row,
)


def generate_customer_timesheet_excel(payload: CustomerTimesheetPackPayload) -> bytes:
    workbook = Workbook()
    _write_associate_sheet(workbook, payload)
    _write_tool_sheet(workbook, payload)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_associate_sheet(workbook: Workbook, payload: CustomerTimesheetPackPayload) -> None:
    sheet = workbook.active
    sheet.title = (
        "Weekly Timesheet"[:31]
        if payload.period.period_type == "weekly"
        else "Monthly Timesheet"
    )

    headers = [
        "Sl No.",
        "Name of the associate",
        "Designation",
        "Productive hours",
        "Non Productive hours",
        "TOTAL",
        "Utilization",
        "Remarks",
    ]
    period_label = (
        f"Period: {payload.period.start_date.strftime('%d %b %Y')} to "
        f"{payload.period.end_date.strftime('%d %b %Y')}"
    )
    extras = [
        f"Customer: {payload.customer_name}",
        f"Working hours target: {float(payload.working_hours_target):.0f}",
    ]
    if payload.week_number is not None:
        extras.insert(1, f"Week of {payload.week_number}")

    next_row = write_report_letterhead(
        sheet,
        company_name=payload.company_name,
        report_title=payload.title,
        period_label=period_label,
        extra_lines=extras,
        col_span=len(headers),
    )

    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, len(headers))
    style_utilization_header(sheet, header_row, 7)

    for offset, row in enumerate(payload.associates):
        excel_row = header_row + 1 + offset
        values = [
            row.serial_no,
            row.associate_name,
            row.designation or "",
            float(row.productive_hours),
            float(row.non_productive_hours),
            float(row.total_hours),
            float(row.utilization_percent) / 100.0,
            row.remarks or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col, value=value)
            if col == 7:
                cell.number_format = "0%"
                cell.alignment = Alignment(horizontal="right")
            elif col in {4, 5, 6}:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "0.00"

    total_row = header_row + 1 + len(payload.associates)
    sheet.cell(row=total_row, column=2, value="TOTAL")
    sheet.cell(row=total_row, column=4, value=float(payload.total_productive_hours))
    sheet.cell(row=total_row, column=5, value=float(payload.total_non_productive_hours))
    sheet.cell(row=total_row, column=6, value=float(payload.total_hours))
    util_cell = sheet.cell(
        row=total_row,
        column=7,
        value=float(payload.overall_utilization_percent) / 100.0,
    )
    util_cell.number_format = "0%"
    style_total_row(sheet, total_row, len(headers), emphasize_cols={2, 4, 5, 6, 7})

    if payload.associates:
        style_body_rows(sheet, header_row + 1, total_row - 1, len(headers))

    freeze_and_filter(sheet, header_row, len(headers))
    autofit_columns(sheet, min_width=12)
    set_print_layout(
        sheet,
        company_name=payload.company_name,
        audience="customer",
        header_rows=f"{header_row}:{header_row}",
    )


def _write_tool_sheet(workbook: Workbook, payload: CustomerTimesheetPackPayload) -> None:
    sheet = workbook.create_sheet("By Tool")
    headers = ["TOOL No.", "Sum of HOURS", "Comments"]
    next_row = write_report_letterhead(
        sheet,
        company_name=payload.company_name,
        report_title=f"{payload.title} — Tool Rollup",
        period_label=f"Customer: {payload.customer_name}",
        extra_lines=[f"Period: {payload.period.label}"],
        col_span=len(headers),
    )

    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, len(headers))

    for offset, row in enumerate(payload.tools):
        excel_row = header_row + 1 + offset
        sheet.cell(row=excel_row, column=1, value=row.tool_number)
        hours_cell = sheet.cell(row=excel_row, column=2, value=float(row.hours))
        hours_cell.number_format = "0.00"
        hours_cell.alignment = Alignment(horizontal="right")
        sheet.cell(row=excel_row, column=3, value=row.comments or "")

    total_row = header_row + 1 + len(payload.tools)
    sheet.cell(row=total_row, column=1, value="TOTAL")
    sheet.cell(
        row=total_row,
        column=2,
        value=float(sum((row.hours for row in payload.tools), Decimal("0"))),
    ).number_format = "0.00"
    style_total_row(sheet, total_row, len(headers), emphasize_cols={1, 2})

    if payload.tools:
        style_body_rows(sheet, header_row + 1, total_row - 1, len(headers))

    freeze_and_filter(sheet, header_row, len(headers))
    autofit_columns(sheet, min_width=14)
    set_print_layout(sheet, company_name=payload.company_name, audience="customer")

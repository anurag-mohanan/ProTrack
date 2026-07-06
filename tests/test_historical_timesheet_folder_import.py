import io
import time
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.models.models import Project, TimesheetEntry
from app.services.historical_timesheet_folder_import_service import (
    parse_prosohm_workbook,
    scan_folder_source,
)
from tests.conftest import login


def build_prosohm_workbook(
    *,
    designer: str,
    month: str,
    rows: list[list],
    header_row: int = 5,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timesheet"
    sheet["B1"] = designer
    sheet["B2"] = month
    headers = ["SL No", "Date", "Project", "Customer", "Task", "Billable", "Hours", "Notes"]
    for col, header in enumerate(headers, start=1):
        sheet.cell(header_row, col, header)
    for offset, row in enumerate(rows, start=1):
        for col, value in enumerate(row, start=1):
            sheet.cell(header_row + offset, col, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def prosohm_productive_bytes():
    return build_prosohm_workbook(
        designer="Binil JR",
        month="Jun-26",
        rows=[
            [1, "2026-06-23", "T-100", "Acme", "Design", "Yes", 8, "Folder import test"],
        ],
    )


def test_parse_prosohm_workbook(tmp_path, prosohm_productive_bytes):
    path = tmp_path / "Binil JR" / "Jun-26.xlsx"
    path.parent.mkdir(parents=True)
    path.write_bytes(prosohm_productive_bytes)
    result = parse_prosohm_workbook(path, relative_path="Binil JR/Jun-26.xlsx")
    assert result.designer == "Binil JR"
    assert result.month_label == "Jun-26"
    assert len(result.rows) == 1
    assert result.rows[0].project_number == "T-100"
    assert result.rows[0].hours == Decimal("8")


def test_scan_folder_batch(tmp_path, prosohm_productive_bytes):
    root = tmp_path / "Historical Timesheets"
    designer_dir = root / "Binil JR"
    designer_dir.mkdir(parents=True)
    (designer_dir / "Jun-26.xlsx").write_bytes(prosohm_productive_bytes)

    scan = scan_folder_source(source_path=str(root))
    assert scan.file_count == 1
    assert scan.estimated_entries == 1
    assert scan.designers[0].designer == "Binil JR"
    assert scan.designers[0].files == 1


def test_folder_upload_and_scan(client, prosohm_productive_bytes):
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/folder/upload",
        headers=headers,
        files={"files": ("Jun-26.xlsx", prosohm_productive_bytes, "application/vnd.ms-excel")},
        data={"paths": "Binil JR/Jun-26.xlsx"},
    )
    assert upload.status_code == 200, upload.text
    batch_id = upload.json()["batch_id"]

    scan = client.post(
        "/api/v1/imports/historical-timesheets/folder/scan",
        headers=headers,
        json={"batch_id": batch_id},
    )
    assert scan.status_code == 200
    body = scan.json()
    assert body["file_count"] == 1
    assert body["estimated_entries"] == 1


def _poll_folder_job(client, headers, job_id: str):
    for _ in range(40):
        response = client.get(
            f"/api/v1/imports/historical-timesheets/folder/jobs/{job_id}",
            headers=headers,
        )
        job = response.json()
        if job["status"] in {"completed", "failed", "cancelled"}:
            return job
        time.sleep(0.1)
    raise AssertionError("Folder import job did not complete")


def test_folder_import_creates_entries(client, prosohm_productive_bytes, test_session_factory):
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/folder/upload",
        headers=headers,
        files={"files": ("Jun-26.xlsx", prosohm_productive_bytes, "application/vnd.ms-excel")},
        data={"paths": "Binil JR/Jun-26.xlsx"},
    )
    batch_id = upload.json()["batch_id"]

    run = client.post(
        "/api/v1/imports/historical-timesheets/folder/run",
        headers=headers,
        json={"batch_id": batch_id},
    )
    assert run.status_code == 200, run.text
    job = _poll_folder_job(client, headers, run.json()["job_id"])
    assert job["status"] == "completed"
    assert job["summary"]["rows_imported"] == 1

    db = test_session_factory()
    try:
        project = db.scalar(select(Project).where(Project.tool_number == "T-100"))
        assert project is not None
        assert project.actual_hours == Decimal("8")
        entries = db.scalars(select(TimesheetEntry).where(TimesheetEntry.project_id == project.id)).all()
        assert len(entries) == 1
    finally:
        db.close()


def test_folder_import_allows_consecutive_days_same_project(client, test_session_factory):
    workbook_bytes = build_prosohm_workbook(
        designer="Binil JR",
        month="May-26",
        rows=[
            [1, "2026-05-01", "T-100", "Acme", "Design", "Yes", 8, "Design for intermediates"],
            [2, "2026-05-02", "T-100", "Acme", "Design", "Yes", 8, "Design for intermediates"],
        ],
    )
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/folder/upload",
        headers=headers,
        files={"files": ("May-26.xlsx", workbook_bytes, "application/vnd.ms-excel")},
        data={"paths": "Binil JR/May-26.xlsx"},
    )
    batch_id = upload.json()["batch_id"]
    run = client.post(
        "/api/v1/imports/historical-timesheets/folder/run",
        headers=headers,
        json={"batch_id": batch_id},
    )
    job = _poll_folder_job(client, headers, run.json()["job_id"])
    assert job["status"] == "completed"
    assert job["summary"]["rows_imported"] == 2
    assert job["summary"]["duplicates_skipped"] == 0

    db = test_session_factory()
    try:
        project = db.scalar(select(Project).where(Project.tool_number == "T-100"))
        entries = db.scalars(select(TimesheetEntry).where(TimesheetEntry.project_id == project.id)).all()
        assert len(entries) == 2
    finally:
        db.close()


def test_folder_import_allows_same_day_different_notes(client, test_session_factory):
    workbook_bytes = build_prosohm_workbook(
        designer="Binil JR",
        month="May-26",
        rows=[
            [1, "2026-05-01", "T-100", "Acme", "Design", "Yes", 8, "Model update"],
            [2, "2026-05-01", "T-100", "Acme", "Design", "Yes", 8, "Design for intermediates"],
        ],
    )
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/folder/upload",
        headers=headers,
        files={"files": ("May-26.xlsx", workbook_bytes, "application/vnd.ms-excel")},
        data={"paths": "Binil JR/May-26.xlsx"},
    )
    batch_id = upload.json()["batch_id"]
    run = client.post(
        "/api/v1/imports/historical-timesheets/folder/run",
        headers=headers,
        json={"batch_id": batch_id},
    )
    job = _poll_folder_job(client, headers, run.json()["job_id"])
    assert job["summary"]["rows_imported"] == 2
    assert job["summary"]["duplicates_skipped"] == 0


def test_folder_import_skips_exact_duplicate(client, prosohm_productive_bytes):
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/folder/upload",
        headers=headers,
        files={"files": ("Jun-26.xlsx", prosohm_productive_bytes, "application/vnd.ms-excel")},
        data={"paths": "Binil JR/Jun-26.xlsx"},
    )
    batch_id = upload.json()["batch_id"]

    first = client.post(
        "/api/v1/imports/historical-timesheets/folder/run",
        headers=headers,
        json={"batch_id": batch_id},
    )
    _poll_folder_job(client, headers, first.json()["job_id"])

    second = client.post(
        "/api/v1/imports/historical-timesheets/folder/run",
        headers=headers,
        json={"batch_id": batch_id},
    )
    job = _poll_folder_job(client, headers, second.json()["job_id"])
    assert job["summary"]["duplicates_skipped"] == 1
    assert job["summary"]["rows_imported"] == 0

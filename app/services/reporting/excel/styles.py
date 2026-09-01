"""ProTrack Excel report branding theme — shared visual system for all exports.

Logo resolution (production-safe):
  1. Uploaded company logo under ``COMPANY_LOGO_DIR`` (``uploads/company/``)
     Preferred filenames: ``company-logo.png|.jpg|.jpeg|.webp``
  2. Fallback filenames in the same folder: ``Prosohm_Logo.png``, ``logo.png``
  3. Optional env override: ``PROTRACK_REPORT_LOGO_PATH`` (absolute file path)

Do not hardcode developer machine paths. Upload the logo via Admin → Company settings.
"""

from __future__ import annotations

from dataclasses import dataclass
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# ── Brand colours (Prosohm / ProTrack) ──────────────────────────────────────
PRIMARY = "0066B3"
PRIMARY_DARK = "004C86"
PRIMARY_SOFT = "D9EAF7"
PRIMARY_MUTED = "E8F1FA"
SECONDARY = "1E293B"
TEXT = "1F2937"
TEXT_MUTED = "64748B"
TEXT_LIGHT = "FFFFFF"
BORDER = "D0D7DE"
BORDER_STRONG = "94A3B8"
ALT_ROW = "F6F8FA"
CARD_FILL = "F8FAFC"
CARD_HEADER_FILL = "EEF5FB"
TOTAL_FILL = "E8F1FA"
SUCCESS = "1B7F4B"
WARNING = "C47F00"
DANGER = "B42318"
UTIL_HIGH_FILL = PatternFill("solid", fgColor="E8F8EF")
UTIL_LOW_FILL = PatternFill("solid", fgColor="FFF4E5")

FONT_NAME = "Calibri"

# ── Typography hierarchy ────────────────────────────────────────────────────
COMPANY_FONT = Font(name=FONT_NAME, size=22, bold=True, color=PRIMARY)
REPORT_TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True, color=SECONDARY)
SECTION_FONT = Font(name=FONT_NAME, size=12, bold=True, color=PRIMARY)
CARD_TITLE_FONT = Font(name=FONT_NAME, size=9, bold=True, color=PRIMARY)
CARD_VALUE_FONT = Font(name=FONT_NAME, size=12, bold=True, color=SECONDARY)
CARD_DETAIL_FONT = Font(name=FONT_NAME, size=9, color=TEXT_MUTED)
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color=TEXT_LIGHT)
TITLE_FONT = COMPANY_FONT  # backward-compat alias used by letterhead callers
SUBTITLE_FONT = Font(name=FONT_NAME, size=11, color=TEXT_MUTED)
BODY_FONT = Font(name=FONT_NAME, size=10, color=TEXT)
BODY_BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True, color=TEXT)
KPI_LABEL_FONT = Font(name=FONT_NAME, size=8, bold=True, color=TEXT_MUTED)
KPI_VALUE_FONT = Font(name=FONT_NAME, size=14, bold=True, color=PRIMARY)
TOTAL_FONT = Font(name=FONT_NAME, size=10, bold=True, color=PRIMARY)
TOTAL_EMPHASIS_FONT = Font(name=FONT_NAME, size=10, bold=True, color=TEXT_LIGHT)
FOOTER_FONT = Font(name=FONT_NAME, size=8, color=TEXT_MUTED)
NOTES_FONT = Font(name=FONT_NAME, size=9, color=TEXT_MUTED)

# ── Fills / borders ─────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
ALT_FILL = PatternFill("solid", fgColor=ALT_ROW)
KPI_FILL = PatternFill("solid", fgColor=PRIMARY_SOFT)
CARD_BG = PatternFill("solid", fgColor=CARD_FILL)
CARD_HEADER_BG = PatternFill("solid", fgColor=CARD_HEADER_FILL)
TOTAL_BG = PatternFill("solid", fgColor=TOTAL_FILL)
TOTAL_EMPHASIS_FILL = PatternFill("solid", fgColor=PRIMARY)
ACCENT_FILL = PatternFill("solid", fgColor=PRIMARY)
SOFT_BAND_FILL = PatternFill("solid", fgColor=PRIMARY_SOFT)

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER),
    right=Side(style="thin", color=BORDER),
    top=Side(style="thin", color=BORDER),
    bottom=Side(style="thin", color=BORDER),
)
CARD_BORDER = Border(
    left=Side(style="thin", color=BORDER_STRONG),
    right=Side(style="thin", color=BORDER_STRONG),
    top=Side(style="thin", color=BORDER_STRONG),
    bottom=Side(style="thin", color=BORDER_STRONG),
)
TOTAL_TOP_BORDER = Border(
    left=Side(style="thin", color=BORDER),
    right=Side(style="thin", color=BORDER),
    top=Side(style="medium", color=PRIMARY),
    bottom=Side(style="thin", color=BORDER),
)

# ── Alignments ──────────────────────────────────────────────────────────────
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_NO_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

# ── Number formats ──────────────────────────────────────────────────────────
HOURS_FORMAT = "0.0"
HOURS_FORMAT_2 = "0.00"
PERCENT_FORMAT = "0.0"
PERCENT_FORMAT_PCT = "0.0%"

# ── Layout defaults ─────────────────────────────────────────────────────────
COMPANY_ROW_HEIGHT = 28
TITLE_ROW_HEIGHT = 22
SECTION_ROW_HEIGHT = 20
BODY_ROW_HEIGHT = 18
HEADER_ROW_HEIGHT = 28
CARD_ROW_HEIGHT = 16
KPI_VALUE_ROW_HEIGHT = 24
ACCENT_BAR_HEIGHT = 5
SOFT_BAND_HEIGHT = 6

REPORT_VERSION = "v1.0"
DEFAULT_FOOTER_PRODUCT = "ProTrack | Engineering Management Platform"
DEFAULT_CONFIDENTIALITY = "Confidential - Internal Use Only"
CUSTOMER_FACING_CONFIDENTIALITY = "Confidential - For Intended Recipient Only"

# Column width guidance (characters)
COLUMN_WIDTH_HINTS = {
    "team": (22, 30),
    "designer": (20, 28),
    "hours": (12, 16),
    "utilization": (12, 15),
    "projects": (10, 14),
    "tool": (12, 16),
    "customer": (18, 28),
}


@dataclass(frozen=True)
class ReportTheme:
    """Central report branding configuration."""

    company_name_fallback: str = "Prosohm Projects Pvt Ltd"
    primary_color: str = PRIMARY
    secondary_color: str = SECONDARY
    font_name: str = FONT_NAME
    company_font_size: int = 22
    title_font_size: int = 16
    body_font_size: int = 10
    header_font_size: int = 10
    report_version: str = REPORT_VERSION
    footer_product: str = DEFAULT_FOOTER_PRODUCT
    confidentiality_internal: str = DEFAULT_CONFIDENTIALITY
    confidentiality_customer: str = CUSTOMER_FACING_CONFIDENTIALITY


DEFAULT_THEME = ReportTheme()


def style_header_row(sheet, row: int, col_count: int) -> None:
    sheet.row_dimensions[row].height = HEADER_ROW_HEIGHT
    for col in range(1, col_count + 1):
        cell = sheet.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def style_body_rows(sheet, start_row: int, end_row: int, col_count: int) -> None:
    for row in range(start_row, end_row + 1):
        sheet.row_dimensions[row].height = BODY_ROW_HEIGHT
        for col in range(1, col_count + 1):
            cell = sheet.cell(row=row, column=col)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            if (row - start_row) % 2 == 1:
                cell.fill = ALT_FILL


def autofit_columns(sheet, min_width: int = 10, max_width: int = 42) -> None:
    from openpyxl.utils import get_column_letter

    for col_cells in sheet.columns:
        letter = get_column_letter(col_cells[0].column)
        length = max((len(str(cell.value or "")) for cell in col_cells), default=0)
        sheet.column_dimensions[letter].width = min(max(length + 2, min_width), max_width)


def apply_column_widths(
    sheet,
    widths: dict[int, float],
    *,
    min_width: float = 8,
    max_width: float = 42,
) -> None:
    from openpyxl.utils import get_column_letter

    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = min(
            max(width, min_width), max_width
        )


def freeze_and_filter(
    sheet, header_row: int, last_col: int, *, end_row: int | None = None
) -> None:
    from openpyxl.utils import get_column_letter

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    last = end_row if end_row is not None else sheet.max_row
    if last >= header_row:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(last_col)}{last}"
        )

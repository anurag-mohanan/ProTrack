"""Branded Excel letterhead — thin compatibility layer over the shared report template.

Prefer ``app.services.reporting.excel.template`` for new report work.
This module keeps historical imports working for all existing exporters.
"""

from __future__ import annotations

from app.services.reporting.excel.template import (
    render_report_header,
    render_report_logo,
    resolve_company_logo_path,
    set_print_layout,
    style_total_row,
    write_simple_extras,
)
from app.services.reporting.excel.styles import (
    HEADER_FONT,
    THIN_BORDER,
)
from openpyxl.styles import Alignment, PatternFill

# Re-exports for existing imports / tests
add_company_logo_top_right = render_report_logo


def write_report_letterhead(
    sheet,
    *,
    company_name: str,
    report_title: str,
    period_label: str | None = None,
    extra_lines: list[str] | None = None,
    col_span: int = 8,
) -> int:
    """
    Write the standard ProTrack brand header and optional subtitle lines.

    Returns the next free row index (1-based) for table content.
    """
    next_row = render_report_header(
        sheet,
        company_name=company_name,
        report_title=report_title,
        col_span=col_span,
    )
    lines: list[str] = []
    if period_label:
        lines.append(period_label)
    lines.extend(extra_lines or [])
    if lines:
        next_row = write_simple_extras(
            sheet, start_row=next_row, lines=lines, col_span=col_span
        )
    return next_row


def style_utilization_header(sheet, row: int, col: int) -> None:
    cell = sheet.cell(row=row, column=col)
    cell.fill = PatternFill("solid", fgColor="C47F00")
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER

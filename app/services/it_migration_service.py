"""IT spreadsheet migration — analyze / preview / import without storing secrets."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.access_control import MODULE_IT_OPERATIONS
from app.core.config import UPLOAD_DIR
from app.core.exceptions import ProTrackValidationError
from app.models.enums import ActivityAction, EntityType
from app.models.it_operations import (
    Asset,
    AssetType,
    Computer,
    InventoryItem,
    ITSupplier,
    ITUserAccount,
    SoftwareAssignment,
    SoftwareCatalog,
    SoftwareLicensePool,
)
from app.models.models import Customer, User
from app.services import it_asset_service
from app.services.activity_service import log_activity

MODULE = MODULE_IT_OPERATIONS

SOURCE_TYPES = frozenset(
    {
        "hardware",
        "accounts",
        "software",
        "inventory",
        "consumables",
        "suppliers",
    }
)

_SENSITIVE_HEADER_RE = re.compile(
    r"(pass|pwd|secret|credential|teamviewer|windows\s*license|license\s*#)",
    re.I,
)

_SESSIONS_DIR = UPLOAD_DIR / "it_migration"

_SHEET_HINTS: dict[str, list[str]] = {
    "hardware": ["1.hardwares", "hardwares", "hardware"],
    "accounts": ["2.user_credentials", "user_credentials", "credentials"],
    "software": ["3.softwares", "softwares", "licenses"],
    "inventory": ["inventory list", "inventory"],
    "consumables": ["consumables"],
    "suppliers": ["supplier list", "suppliers"],
}

_HEADER_ROW_OVERRIDES: dict[str, int] = {
    "inventory": 3,
    "consumables": 3,
    "suppliers": 3,
}


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _is_sensitive_header(header: str) -> bool:
    return bool(_SENSITIVE_HEADER_RE.search(header or ""))


def _norm_id(value: str) -> str:
    return re.sub(r"\s+", "", (value or "").upper())


def _norm_name(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def _parse_date(value: str) -> date | None:
    raw = (value or "").strip()
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:19], fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        return None


def _parse_int(value: str) -> int | None:
    raw = (value or "").strip()
    if not raw:
        return None
    try:
        return int(float(raw))
    except ValueError:
        return None


def _parse_decimal(value: str) -> Decimal | None:
    raw = (value or "").strip().replace(",", "")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def _sessions_dir() -> Path:
    _SESSIONS_DIR.mkdir(parents=True, exist_ok=True)
    return _SESSIONS_DIR


def _session_path(session_id: str) -> Path:
    safe = re.sub(r"[^a-zA-Z0-9_-]", "", session_id)
    if not safe:
        raise ProTrackValidationError("Invalid migration session id.")
    return _sessions_dir() / f"{safe}.json"


def _save_session(payload: dict[str, Any]) -> str:
    session_id = str(uuid4())
    _session_path(session_id).write_text(json.dumps(payload), encoding="utf-8")
    return session_id


def _load_session(session_id: str) -> dict[str, Any]:
    path = _session_path(session_id)
    if not path.exists():
        raise ProTrackValidationError("Migration session not found or expired.")
    return json.loads(path.read_text(encoding="utf-8"))


def _resolve_sheet_name(sheetnames: list[str], source_type: str) -> str | None:
    hints = _SHEET_HINTS.get(source_type, [])
    lowered = {s.lower().strip(): s for s in sheetnames}
    for hint in hints:
        for key, original in lowered.items():
            if hint in key:
                return original
    return None


def _find_header_row(rows: list[tuple[Any, ...]], source_type: str) -> int:
    if source_type in _HEADER_ROW_OVERRIDES:
        return _HEADER_ROW_OVERRIDES[source_type]
    for i, row in enumerate(rows[:40]):
        vals = [_cell_str(c) for c in (row or ())]
        nonempty = [v for v in vals if v]
        if len(nonempty) >= 3:
            return i
    return 0


def _read_sheet_rows(
    content: bytes, source_type: str
) -> tuple[str, list[str], list[dict[str, str]], list[str]]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        sheet_name = _resolve_sheet_name(wb.sheetnames, source_type)
        if sheet_name is None:
            raise ProTrackValidationError(
                f"Could not find a sheet for source type '{source_type}'. "
                f"Available: {', '.join(wb.sheetnames)}"
            )
        ws = wb[sheet_name]
        raw_rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not raw_rows:
        return sheet_name, [], [], []

    header_idx = _find_header_row(raw_rows, source_type)
    headers_raw = [_cell_str(c) for c in (raw_rows[header_idx] or ())]
    while headers_raw and not headers_raw[-1]:
        headers_raw.pop()

    headers: list[str] = []
    seen: dict[str, int] = {}
    for h in headers_raw:
        key = h or f"col_{len(headers)}"
        if key in seen:
            seen[key] += 1
            key = f"{key}_{seen[key]}"
        else:
            seen[key] = 1
        headers.append(key)

    sensitive = [h for h in headers if _is_sensitive_header(h)]
    safe_headers = [h for h in headers if h not in sensitive]

    rows: list[dict[str, str]] = []
    for row in raw_rows[header_idx + 1 :]:
        vals = [_cell_str(c) for c in (row or ())]
        if len(vals) < len(headers):
            vals += [""] * (len(headers) - len(vals))
        record = {headers[i]: vals[i] for i in range(len(headers))}
        meaningful = [v for k, v in record.items() if v and k not in sensitive]
        if not meaningful:
            continue
        rows.append({k: record.get(k, "") for k in safe_headers})

    return sheet_name, safe_headers, rows, sensitive


def _customer_index(db: Session) -> dict[str, Customer]:
    customers = list(db.scalars(select(Customer).where(Customer.is_active.is_(True))).all())
    index: dict[str, Customer] = {}
    for c in customers:
        index[_norm_name(c.name)] = c
        if c.code:
            index[_norm_name(c.code)] = c
    return index


def _user_index(db: Session) -> dict[str, User]:
    users = list(db.scalars(select(User).where(User.is_deleted.is_(False))).all())
    index: dict[str, User] = {}
    for u in users:
        full = f"{u.first_name} {u.last_name}".strip()
        if full:
            index[_norm_name(full)] = u
        if u.email:
            index[_norm_name(u.email)] = u
            index[_norm_name(u.email.split("@")[0])] = u
    return index


def _match_customer(index: dict[str, Customer], label: str) -> Customer | None:
    key = _norm_name(label)
    if not key:
        return None
    if key in {"prosohm", "prosohm eng", "prosohm admin", "generic", "admin", "organization"}:
        return None
    if key in index:
        return index[key]
    for k, cust in index.items():
        if key in k or k in key:
            return cust
    return None


def _match_user(index: dict[str, User], label: str) -> User | None:
    key = _norm_name(label)
    if not key:
        return None
    if key in {
        "all",
        "open",
        "n/a",
        "servprosohm",
        "conference room",
        "planning board",
        "eng planning board",
        "prosohm eng team",
    }:
        return None
    if key in index:
        return index[key]
    base = re.sub(r"\(.*?\)", "", key).strip()
    if base in index:
        return index[base]
    for k, user in index.items():
        if base and (base in k or k in base):
            return user
    return None


def _existing_asset_keys(db: Session) -> dict[str, set[str]]:
    assets = list(db.scalars(select(Asset).where(Asset.is_deleted.is_(False))).all())
    return {
        "legacy": {_norm_id(a.legacy_asset_number or "") for a in assets if a.legacy_asset_number},
        "serial": {_norm_id(a.serial_number or "") for a in assets if a.serial_number},
        "service_tag": {_norm_id(a.service_tag or "") for a in assets if a.service_tag},
        "asset_number": {_norm_id(a.asset_number or "") for a in assets if a.asset_number},
    }


def _get(row: dict[str, str], *candidates: str) -> str:
    lower_map = {k.lower(): k for k in row}
    for cand in candidates:
        key = lower_map.get(cand.lower())
        if key and row.get(key):
            return row[key]
    for cand in candidates:
        for k, v in row.items():
            if cand.lower() in k.lower() and v:
                return v
    return ""


def list_source_types() -> list[dict[str, str]]:
    return [
        {
            "id": "hardware",
            "label": "Hardware (IT Records)",
            "sheet_hint": "1.HARDWARES",
            "description": "Serialized computers, servers, firewalls, printers.",
        },
        {
            "id": "accounts",
            "label": "Users & Accounts",
            "sheet_hint": "2.USER_CREDENTIALS",
            "description": "Account metadata only — passwords are never imported.",
        },
        {
            "id": "software",
            "label": "Software & Licenses",
            "sheet_hint": "3.SOFTWARES_&_LICENSES",
            "description": "Software catalog, seats, ownership, assignments.",
        },
        {
            "id": "inventory",
            "label": "Asset Inventory",
            "sheet_hint": "INVENTORY LIST",
            "description": "General / IT inventory list with ownership signals.",
        },
        {
            "id": "consumables",
            "label": "Consumables",
            "sheet_hint": "CONSUMABLES",
            "description": "Stock quantities (cables, mice, keyboards, …).",
        },
        {
            "id": "suppliers",
            "label": "Suppliers",
            "sheet_hint": "SUPPLIER LIST",
            "description": "Vendor master (sheet may be empty).",
        },
    ]


def analyze_upload(
    db: Session,
    *,
    actor: User,
    content: bytes,
    filename: str,
    source_type: str,
) -> dict[str, Any]:
    source_type = (source_type or "").strip().lower()
    if source_type not in SOURCE_TYPES:
        raise ProTrackValidationError(
            f"Invalid source_type. Allowed: {', '.join(sorted(SOURCE_TYPES))}"
        )
    if not content:
        raise ProTrackValidationError("Uploaded file is empty.")

    sheet_name, headers, rows, sensitive = _read_sheet_rows(content, source_type)
    customers = _customer_index(db)
    users = _user_index(db)
    existing = _existing_asset_keys(db)

    exceptions: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    preview_rows: list[dict[str, str]] = []
    new_count = 0
    skip_count = 0
    review_count = 0
    seen_ids: dict[str, int] = {}

    for idx, row in enumerate(rows, start=1):
        row_exceptions: list[dict[str, Any]] = []
        action = "import"
        match_key = None

        if source_type in {"hardware", "inventory"}:
            legacy = _get(row, "ID", "ID Tag", "Id Tag")
            serial = _get(row, "SERIAL NUMBER", "Serial Number", "Service/Serial")
            service = _get(row, "SERVICE TAG#", "Service Tag", "SERVICE TAG", "Service/Serial")
            if not serial and service:
                serial = service

            norm_legacy = _norm_id(legacy)
            norm_serial = _norm_id(serial)
            norm_service = _norm_id(service)
            if norm_legacy:
                seen_ids[norm_legacy] = seen_ids.get(norm_legacy, 0) + 1

            strong = None
            if norm_service and norm_service in existing["service_tag"]:
                strong = "service_tag"
            elif norm_serial and norm_serial in existing["serial"]:
                strong = "serial_number"
            elif norm_legacy and (
                norm_legacy in existing["legacy"] or norm_legacy in existing["asset_number"]
            ):
                strong = "legacy_asset_number"

            if strong:
                action = "skip_duplicate"
                skip_count += 1
                duplicates.append(
                    {
                        "row": idx,
                        "match_type": "strong",
                        "matched_on": strong,
                        "legacy_id": legacy or None,
                        "serial": serial or None,
                        "service_tag": service or None,
                    }
                )
            else:
                new_count += 1

            if source_type == "hardware":
                lid_u = legacy.upper().replace(" ", "")
                if lid_u.startswith("IT-"):
                    row_exceptions.append(
                        {
                            "code": "OWNERSHIP_UNCLEAR",
                            "severity": "review",
                            "detail": f"ID {legacy} does not encode ownership; requires admin review.",
                        }
                    )
                    review_count += 1
                cust_used = _get(row, "CUSTOMER USED FOR")
                if cust_used and _match_customer(customers, cust_used) is None:
                    if _norm_name(cust_used) not in {
                        "generic",
                        "prosohm eng",
                        "prosohm admin",
                        "prosohm",
                    }:
                        row_exceptions.append(
                            {
                                "code": "CUSTOMER_UNKNOWN",
                                "severity": "review",
                                "detail": f"CUSTOMER USED FOR '{cust_used}' not matched.",
                            }
                        )
                        review_count += 1
                assignee = _get(row, "CURRENT USER")
                if assignee and _match_user(users, assignee) is None:
                    row_exceptions.append(
                        {
                            "code": "USER_UNKNOWN",
                            "severity": "warning",
                            "detail": f"CURRENT USER '{assignee}' not matched to an employee.",
                        }
                    )
                if not service and not serial:
                    row_exceptions.append(
                        {
                            "code": "SERIAL_MISSING",
                            "severity": "warning",
                            "detail": "No service tag / serial on hardware row.",
                        }
                    )

            if source_type == "inventory":
                cust_paid = _get(row, "CUST  PAID  ?", "CUST PAID ?", "Customer Paid")
                if cust_paid and _match_customer(customers, cust_paid) is None:
                    row_exceptions.append(
                        {
                            "code": "CUSTOMER_UNKNOWN",
                            "severity": "review",
                            "detail": f"CUST PAID '{cust_paid}' not matched to Customer.",
                        }
                    )
                    review_count += 1
                elif not cust_paid:
                    row_exceptions.append(
                        {
                            "code": "OWNERSHIP_UNCLEAR",
                            "severity": "review",
                            "detail": "CUST PAID blank — ownership unclear.",
                        }
                    )
                    review_count += 1

            match_key = legacy or serial or service

        elif source_type == "accounts":
            name = _get(row, "EMPLOYEE NAME", "Employee")
            email = _get(row, "EMAIL", "Email")
            username = _get(row, "USERNAME", "Username")
            user = (
                _match_user(users, email)
                or _match_user(users, name)
                or _match_user(users, username)
            )
            if user is None:
                action = "skip"
                skip_count += 1
                row_exceptions.append(
                    {
                        "code": "USER_UNKNOWN",
                        "severity": "review",
                        "detail": f"Employee '{name or username}' not matched.",
                    }
                )
                review_count += 1
            else:
                new_count += 1
                action = "import_metadata"
            if sensitive:
                row_exceptions.append(
                    {
                        "code": "SENSITIVE_EXCLUDED",
                        "severity": "info",
                        "detail": "Password columns excluded; credential_status=migration_required.",
                    }
                )
            match_key = username or email or name

        elif source_type == "software":
            software = _get(row, "SOFTWARE", "Software")
            purchased = _get(row, "PURCHASED BY", "Purchased By")
            assignee = _get(row, "USER", "User")
            if not software:
                action = "skip"
                skip_count += 1
            else:
                new_count += 1
            if purchased and _norm_name(purchased) not in {"prosohm", "organization"}:
                if _match_customer(customers, purchased) is None:
                    row_exceptions.append(
                        {
                            "code": "CUSTOMER_UNKNOWN",
                            "severity": "review",
                            "detail": f"PURCHASED BY '{purchased}' not matched.",
                        }
                    )
                    review_count += 1
            if assignee and _match_user(users, assignee) is None:
                row_exceptions.append(
                    {
                        "code": "USER_UNKNOWN",
                        "severity": "warning",
                        "detail": f"Assigned USER '{assignee}' not matched.",
                    }
                )
            match_key = software

        elif source_type == "consumables":
            name = _get(row, "Name", "NAME")
            if not name:
                action = "skip"
                skip_count += 1
            else:
                new_count += 1
                qty = _parse_int(_get(row, "QTY", "Qty"))
                issued = _parse_int(_get(row, "IN USE", "In Use"))
                available = _parse_int(_get(row, "AVAILABLE", "Available"))
                if (
                    qty is not None
                    and issued is not None
                    and available is not None
                    and qty != issued + available
                ):
                    row_exceptions.append(
                        {
                            "code": "QTY_INCONSISTENT",
                            "severity": "warning",
                            "detail": f"QTY {qty} != IN USE {issued} + AVAILABLE {available}",
                        }
                    )
            match_key = name

        elif source_type == "suppliers":
            name = _get(row, "SUPPLIER NAME", "Supplier Name", "Name")
            if not name:
                action = "skip"
                skip_count += 1
            else:
                new_count += 1
            match_key = name

        for exc in row_exceptions:
            exceptions.append(
                {
                    "id": f"EX-{source_type[:3].upper()}-{idx:04d}-{exc['code']}",
                    "row": idx,
                    "source_key": match_key,
                    **exc,
                }
            )

        if len(preview_rows) < 15 and action != "skip":
            preview_rows.append({k: v for k, v in row.items() if v})

    for lid, count in seen_ids.items():
        if count > 1:
            exceptions.append(
                {
                    "id": f"EX-HW-DUP-{lid}",
                    "row": None,
                    "source_key": lid,
                    "code": "LEGACY_ID_COLLISION",
                    "severity": "blocker",
                    "detail": f"Legacy ID {lid} appears {count} times in the source file.",
                }
            )
            review_count += 1

    if source_type == "suppliers" and new_count == 0:
        exceptions.append(
            {
                "id": "EX-SUP-EMPTY",
                "row": None,
                "source_key": None,
                "code": "SHEET_EMPTY",
                "severity": "warning",
                "detail": "Supplier sheet has no data.",
            }
        )

    session_payload = {
        "source_type": source_type,
        "filename": filename,
        "sheet_name": sheet_name,
        "headers": headers,
        "sensitive_columns_excluded": sensitive,
        "rows": rows,
        "created_at": datetime.utcnow().isoformat() + "Z",
        "actor_id": str(actor.id),
    }
    session_id = _save_session(session_payload)

    log_activity(
        db,
        user=actor,
        entity_type=EntityType.it_settings,
        entity_id=actor.id,
        action=ActivityAction.it_migration_analyzed,
        new_value={
            "source_type": source_type,
            "filename": filename,
            "records": len(rows),
            "sensitive_columns_excluded": sensitive,
            "session_id": session_id,
        },
        outcome="success",
        module=MODULE,
        commit=True,
    )

    return {
        "session_id": session_id,
        "source_type": source_type,
        "filename": filename,
        "sheet_name": sheet_name,
        "records_found": len(rows),
        "new_records": new_count,
        "potential_duplicates": len(duplicates),
        "skipped_records": skip_count,
        "requires_review": review_count,
        "sensitive_columns_excluded": sensitive,
        "sensitive_data_excluded_count": len(sensitive),
        "headers": headers,
        "duplicates": duplicates[:100],
        "exceptions": exceptions[:200],
        "preview_rows": preview_rows,
        "confirm_required": True,
        "message": (
            "Preview only — no data written. Secrets were excluded. "
            "Call import with session_id and confirm=true to commit."
        ),
    }


def _resolve_asset_type(db: Session, type_label: str) -> AssetType:
    code_map = {
        "laptop": "LAPTOP",
        "desktop": "DESKTOP",
        "workstation": "DESKTOP",
        "monitor": "MONITOR",
        "server": "SERVER",
        "firewall": "OTHER",
        "printer": "OTHER",
        "keyboard": "KEYBOARD",
        "mouse": "MOUSE",
        "it": "OTHER",
    }
    label = (type_label or "OTHER").strip().lower()
    code = code_map.get(label, "OTHER")
    at = db.scalar(select(AssetType).where(AssetType.code == code))
    if at is None:
        at = db.scalar(select(AssetType).where(AssetType.code == "OTHER"))
    if at is None:
        raise ProTrackValidationError("No asset types configured.")
    return at


def _infer_purchased_by(
    *,
    legacy_id: str,
    cust_paid: str,
    customers: dict[str, Customer],
) -> tuple[str, UUID | None, str | None]:
    if cust_paid:
        cust = _match_customer(customers, cust_paid)
        if cust:
            return "customer", cust.id, None
        return "organization", None, f"Unmatched CUST PAID '{cust_paid}'"

    lid = (legacy_id or "").upper().replace(" ", "")
    if lid.startswith("SY-"):
        return "customer", None, "SY- prefix suggests customer ownership"
    return "organization", None, None


def _get_or_create_supplier(db: Session, name: str) -> UUID | None:
    norm = _norm_name(name)
    if not norm:
        return None
    existing = list(db.scalars(select(ITSupplier)).all())
    for s in existing:
        if _norm_name(s.name) == norm:
            return s.id
    for s in existing:
        if norm in _norm_name(s.name) or _norm_name(s.name) in norm:
            return s.id
    supplier = ITSupplier(id=uuid4(), name=name.strip(), is_active=True)
    db.add(supplier)
    db.flush()
    return supplier.id


def commit_import(
    db: Session,
    *,
    actor: User,
    session_id: str,
    confirm: bool,
    skip_duplicates: bool = True,
    skip_review_rows: bool = False,
) -> dict[str, Any]:
    if not confirm:
        raise ProTrackValidationError(
            "Import requires confirm=true after reviewing the analyze preview."
        )
    session = _load_session(session_id)
    source_type = session["source_type"]
    rows: list[dict[str, str]] = session.get("rows") or []
    customers = _customer_index(db)
    users = _user_index(db)
    existing = _existing_asset_keys(db)

    imported = 0
    skipped = 0
    duplicated = 0
    conflicted = 0
    errors: list[str] = []
    owned_org = 0
    owned_cust = 0
    sensitive_excluded = len(session.get("sensitive_columns_excluded") or [])

    try:
        if source_type in {"hardware", "inventory"}:
            for idx, row in enumerate(rows, start=1):
                legacy = _get(row, "ID", "ID Tag")
                serial = _get(row, "SERIAL NUMBER", "Serial Number")
                service = _get(
                    row, "SERVICE TAG#", "Service Tag", "SERVICE TAG", "Service/Serial"
                )
                if not serial and service and source_type == "hardware":
                    serial = service
                make = _get(row, "MAKE", "Make", "Name")
                model = _get(row, "MODEL", "Model", "MODEL NUMBER")
                desc = _get(row, "DESCRIPTION", "Description")
                location = _get(row, "ROOM", "LOCATION")
                type_label = _get(row, "TYPE", "Category")
                cust_used = _get(row, "CUSTOMER USED FOR")
                cust_paid = _get(row, "CUST  PAID  ?", "CUST PAID ?")
                purchase_date = _parse_date(_get(row, "DATE OF PURCHASE", "DATE", "Date"))
                warranty = _parse_date(_get(row, "WARRANTY VALID UPTO", "WARRANTY/EXPIRATION"))
                cost = _parse_decimal(_get(row, "VALUE", "Value"))
                condition = _get(row, "CONDITION", "CURRENT STATUS")
                supplier_name = _get(row, "SUPPLIER", "Supplier")
                invoice = _get(row, "INVOICE NO", "Invoice")
                computer_name = _get(row, "NAME")
                os_name = _get(row, "OS")
                ram = _parse_int(_get(row, "RAM (GB)", "RAM"))
                cpu = _get(row, "CPU (Ghz)", "CPU")
                mac = _get(row, "MAC Address", "MAC")
                current_user_label = _get(row, "CURRENT USER", "Current Designer / user")

                if source_type == "inventory":
                    category = (_get(row, "Category") or "").upper()
                    if category in {"FURNITURE", "KITCHEN", "TOOLS"}:
                        skipped += 1
                        continue

                norm_legacy = _norm_id(legacy)
                norm_serial = _norm_id(serial)
                norm_service = _norm_id(service)
                if (
                    (norm_service and norm_service in existing["service_tag"])
                    or (norm_serial and norm_serial in existing["serial"])
                    or (
                        norm_legacy
                        and (
                            norm_legacy in existing["legacy"]
                            or norm_legacy in existing["asset_number"]
                        )
                    )
                ):
                    duplicated += 1
                    if skip_duplicates:
                        skipped += 1
                        continue
                    conflicted += 1
                    continue

                purchased_by, owner_id, note = _infer_purchased_by(
                    legacy_id=legacy,
                    cust_paid=cust_paid,
                    customers=customers,
                )
                if purchased_by == "customer" and owner_id is None:
                    for key, cust in customers.items():
                        if "sybridge" in key:
                            owner_id = cust.id
                            break
                    if owner_id is None:
                        if skip_review_rows:
                            skipped += 1
                            conflicted += 1
                            continue
                        purchased_by = "organization"
                        note = (note or "") + " ; customer owner unresolved"

                used_for = _match_customer(customers, cust_used) if cust_used else None
                supplier_id = _get_or_create_supplier(db, supplier_name) if supplier_name else None

                status = "available"
                cond_u = (condition or "").upper()
                if "DISPOSE" in cond_u or "SCRAP" in cond_u:
                    status = "disposed"
                elif "NOT WORKING" in cond_u or "DAMAGE" in cond_u:
                    status = "damaged"

                try:
                    at = _resolve_asset_type(db, type_label)
                    asset = it_asset_service.create_asset(
                        db,
                        actor=actor,
                        asset_type_id=at.id,
                        serial_number=serial or None,
                        make=make or None,
                        model=model or None,
                        purchase_date=purchase_date,
                        purchase_cost=cost,
                        warranty_expiry=warranty if warranty and warranty.year < 2090 else None,
                        location=location or None,
                        notes="; ".join(
                            x
                            for x in [desc, note, _get(row, "remarks", "REMARKS/NOTES")]
                            if x
                        )
                        or None,
                        legacy_asset_number=legacy or None,
                        description=desc or make or None,
                        service_tag=service or None,
                        purchased_by=purchased_by,
                        owner_customer_id=owner_id,
                        customer_used_for_id=used_for.id if used_for else None,
                        supplier_id=supplier_id,
                        invoice_number=invoice or None,
                        condition=condition or None,
                        commit=False,
                    )
                    if status != "available":
                        asset.status = status
                        db.add(asset)

                    if source_type == "hardware" and (type_label or "").lower() in {
                        "laptop",
                        "workstation",
                        "server",
                        "desktop",
                    }:
                        cname = (computer_name or legacy or asset.asset_number)[:40]
                        existing_c = db.scalar(
                            select(Computer).where(Computer.computer_name == cname)
                        )
                        if existing_c is None:
                            db.add(
                                Computer(
                                    id=uuid4(),
                                    asset_id=asset.id,
                                    computer_name=cname,
                                    os=os_name or None,
                                    processor=cpu or None,
                                    ram_gb=ram,
                                    mac_address=(mac[:17] if mac else None),
                                )
                            )

                    assignee = (
                        _match_user(users, current_user_label) if current_user_label else None
                    )
                    if assignee and asset.status == "available":
                        try:
                            it_asset_service.assign_asset(
                                db,
                                asset,
                                user_id=assignee.id,
                                by_user=actor,
                                notes="Imported assignment",
                                commit=False,
                            )
                        except Exception:
                            pass

                    if purchased_by == "customer":
                        owned_cust += 1
                    else:
                        owned_org += 1
                    imported += 1
                    if norm_legacy:
                        existing["legacy"].add(norm_legacy)
                    if norm_serial:
                        existing["serial"].add(norm_serial)
                    if norm_service:
                        existing["service_tag"].add(norm_service)
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Row {idx}: {exc}")
                    skipped += 1

        elif source_type == "accounts":
            for row in rows:
                name = _get(row, "EMPLOYEE NAME")
                email = _get(row, "EMAIL")
                username = _get(row, "USERNAME")
                status_raw = _get(row, "STATUS")
                user = (
                    _match_user(users, email)
                    or _match_user(users, name)
                    or _match_user(users, username)
                )
                if user is None:
                    skipped += 1
                    continue
                acct_status = "active"
                if "ex-" in status_raw.lower():
                    acct_status = "deactivated"
                if username:
                    existing_acct = db.scalar(
                        select(ITUserAccount).where(
                            ITUserAccount.user_id == user.id,
                            ITUserAccount.account_type == "domain",
                            ITUserAccount.username == username,
                        )
                    )
                    if existing_acct:
                        skipped += 1
                        duplicated += 1
                    else:
                        db.add(
                            ITUserAccount(
                                id=uuid4(),
                                user_id=user.id,
                                account_type="domain",
                                username=username,
                                display_name=name or None,
                                status=acct_status,
                                credential_status="migration_required",
                                notes="Legacy Credential — Migration Required (passwords not imported)",
                            )
                        )
                        imported += 1
                teams = _get(row, "MS TEAMS ID")
                if teams:
                    existing_t = db.scalar(
                        select(ITUserAccount).where(
                            ITUserAccount.user_id == user.id,
                            ITUserAccount.account_type == "email",
                            ITUserAccount.username == teams,
                        )
                    )
                    if not existing_t:
                        db.add(
                            ITUserAccount(
                                id=uuid4(),
                                user_id=user.id,
                                account_type="email",
                                username=teams,
                                display_name=name or None,
                                status=acct_status,
                                credential_status="migration_required",
                                notes="Imported Teams/email identity — Migration Required",
                            )
                        )
                        imported += 1

        elif source_type == "software":
            for row in rows:
                software = _get(row, "SOFTWARE")
                if not software:
                    skipped += 1
                    continue
                purchased = _get(row, "PURCHASED BY")
                seats = _parse_int(_get(row, "NO OF USER\\LICENSE", "NO OF USER/LICENSE")) or 1
                expiry = _parse_date(_get(row, "LICENSE VALID UPTO"))
                renewal = _get(row, "RENEWAL MODE")
                assignee_label = _get(row, "USER")
                dept = _get(row, "DEPARTMENT")
                comments = _get(row, "COMMENTS")

                catalog = db.scalar(
                    select(SoftwareCatalog).where(
                        func.lower(SoftwareCatalog.name) == software.lower()
                    )
                )
                if catalog is None:
                    catalog = SoftwareCatalog(id=uuid4(), name=software, is_active=True)
                    db.add(catalog)
                    db.flush()

                purchased_by = "organization"
                owner_id = None
                if purchased and _norm_name(purchased) not in {"prosohm", "organization"}:
                    cust = _match_customer(customers, purchased)
                    if cust:
                        purchased_by = "customer"
                        owner_id = cust.id
                        owned_cust += 1
                    else:
                        owned_org += 1
                else:
                    owned_org += 1

                pool = SoftwareLicensePool(
                    id=uuid4(),
                    software_id=catalog.id,
                    purchased_by=purchased_by,
                    owner_customer_id=owner_id,
                    seat_count=seats,
                    expiry_date=expiry,
                    renewal_mode=renewal or None,
                    notes=comments or None,
                )
                db.add(pool)
                db.flush()
                assignee = _match_user(users, assignee_label) if assignee_label else None
                if assignee:
                    db.add(
                        SoftwareAssignment(
                            id=uuid4(),
                            license_pool_id=pool.id,
                            user_id=assignee.id,
                            assigned_date=date.today(),
                            department=dept or None,
                        )
                    )
                imported += 1

        elif source_type == "consumables":
            for row in rows:
                name = _get(row, "Name", "NAME")
                if not name:
                    skipped += 1
                    continue
                total = _parse_int(_get(row, "QTY")) or 0
                issued = _parse_int(_get(row, "IN USE")) or 0
                existing_item = db.scalar(
                    select(InventoryItem).where(
                        func.lower(InventoryItem.name) == name.lower(),
                        InventoryItem.is_deleted.is_(False),
                    )
                )
                if existing_item:
                    duplicated += 1
                    skipped += 1
                    continue
                db.add(
                    InventoryItem(
                        id=uuid4(),
                        name=name,
                        description=_get(row, "Description") or None,
                        category=_get(row, "Category") or None,
                        total_qty=total,
                        issued_qty=issued,
                        location=_get(row, "STORAGE LOCATION") or None,
                        condition=_get(row, "CONDITION") or None,
                        unit_cost=_parse_decimal(_get(row, "VALUE")),
                        status="current",
                        purchased_by="organization",
                    )
                )
                imported += 1
                owned_org += 1

        elif source_type == "suppliers":
            for row in rows:
                name = _get(row, "SUPPLIER NAME", "Name")
                if not name:
                    skipped += 1
                    continue
                before = db.scalar(
                    select(func.count()).select_from(ITSupplier).where(
                        func.lower(ITSupplier.name) == name.lower()
                    )
                )
                _get_or_create_supplier(db, name)
                after_exists = db.scalar(
                    select(ITSupplier).where(func.lower(ITSupplier.name) == name.lower())
                )
                if before and before > 0:
                    duplicated += 1
                elif after_exists:
                    imported += 1

        db.flush()
        log_activity(
            db,
            user=actor,
            entity_type=EntityType.it_settings,
            entity_id=actor.id,
            action=ActivityAction.it_migration_imported,
            new_value={
                "source_type": source_type,
                "session_id": session_id,
                "imported": imported,
                "skipped": skipped,
                "duplicated": duplicated,
                "sensitive_excluded": sensitive_excluded,
            },
            outcome="success",
            module=MODULE,
            commit=False,
        )
        db.commit()
    except Exception:
        db.rollback()
        raise

    try:
        _session_path(session_id).unlink(missing_ok=True)
    except OSError:
        pass

    return {
        "session_id": session_id,
        "source_type": source_type,
        "filename": session.get("filename"),
        "imported": imported,
        "skipped": skipped,
        "duplicated": duplicated,
        "conflicted": conflicted,
        "requires_review": conflicted,
        "sensitive_data_excluded": sensitive_excluded,
        "organization_owned": owned_org,
        "customer_owned": owned_cust,
        "returned_assets": 0,
        "errors": errors[:50],
        "message": "Import committed.",
    }

"""Historical timesheet import from a single master Excel workbook."""

from __future__ import annotations

import json
import logging
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import gettempdir
from typing import Callable

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.enums import WorkCategory
from app.models.models import NonProductiveCode, TimesheetEntry, User
from app.schemas.historical_timesheet_master_import import (
    MasterDesignerScanRow,
    MasterImportLogRow,
    MasterImportSummary,
    MasterScanResponse,
)
from app.services.historical_import_service import (
    KNOWN_NP_CODES,
    _get_or_create_approved_timesheet,
    _normalize_header,
    normalize_special_project_code,
    _resolve_np_code,
    _week_start,
)
from app.services.historical_timesheet_folder_import_service import create_pre_import_backup
from app.services.historical_timesheet_import_service import (
    _cell_text,
    _find_project,
    _find_task_type,
    _match_user_by_name,
    _parse_decimal,
    _parse_date,
    _resolve_task_type_name,
)
from app.services.project_calculation_service import recalculate_project
from app.services.non_productive_entry_service import build_np_timesheet_entry

MASTER_UPLOAD_DIR = Path(gettempdir()) / "protrack_master_timesheet_imports"
BATCH_COMMIT_SIZE = 500
PROGRESS_LOG_EVERY_ROWS = 10
logger = logging.getLogger(__name__)

MASTER_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "designer": ("designer",),
    "sl_no": ("sl no", "sl no.", "s.no", "s no", "#", "sl"),
    "entry_date": ("date",),
    "project": ("project", "project number", "tool no", "tool number", "tool no."),
    "customer": ("customer", "customer name"),
    "task": ("task", "task type", "activity"),
    "billable": ("billable",),
    "hours": ("hours", "hrs", "time"),
    "notes": ("notes", "description", "comments"),
}


@dataclass
class MasterRow:
    row_number: int
    designer: str | None = None
    entry_date: date | None = None
    project_value: str | None = None
    task: str | None = None
    billable: bool | None = None
    hours: Decimal | None = None
    notes: str | None = None
    is_np_row: bool = False
    np_code: str | None = None


def _upload_dir(upload_id: str) -> Path:
    return MASTER_UPLOAD_DIR / upload_id


def save_master_upload(upload_id: str, filename: str, content: bytes) -> Path:
    dest = _upload_dir(upload_id)
    dest.mkdir(parents=True, exist_ok=True)
    (dest / "workbook.xlsx").write_bytes(content)
    (dest / "meta.json").write_text(
        json.dumps({"filename": filename}),
        encoding="utf-8",
    )
    return dest


def get_master_upload(upload_id: str) -> tuple[Path, str]:
    dest = _upload_dir(upload_id)
    workbook = dest / "workbook.xlsx"
    meta_path = dest / "meta.json"
    if not workbook.is_file() or not meta_path.is_file():
        raise FileNotFoundError("Master upload not found or expired.")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    return workbook, meta.get("filename", "workbook.xlsx")


def _is_blank_row(values: list[object]) -> bool:
    return not any(v is not None and str(v).strip() for v in values)


def _detect_master_header(headers: list[str]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    normalized = [_normalize_header(h) for h in headers]
    for field_name, aliases in MASTER_HEADER_ALIASES.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if header in aliases or any(alias in header for alias in aliases):
                header_map[field_name] = idx
                break
    required = {"designer", "entry_date", "project", "task", "billable", "hours"}
    if not required.issubset(header_map):
        missing = ", ".join(sorted(required - set(header_map)))
        raise ValueError(f"Missing required columns: {missing}")
    return header_map


def _parse_billable(value: object | None) -> bool | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    if text in {"yes", "y", "true", "1"}:
        return True
    if text in {"no", "n", "false", "0"}:
        return False
    return None


def _classify_project_value(db: Session, value: str | None) -> tuple[bool, str | None]:
    if not value:
        return False, None
    special = normalize_special_project_code(value)
    if special is not None:
        return True, special
    code = value.strip().upper()
    if code in KNOWN_NP_CODES:
        return True, code
    np_record = db.scalar(select(NonProductiveCode).where(NonProductiveCode.code == code))
    if np_record is not None:
        return True, np_record.code
    if code.startswith("EST"):
        return True, code
    return False, None


def _format_month_year(value: date) -> str:
    return value.strftime("%b %Y")


def _format_date_range(date_from: date | None, date_to: date | None) -> str | None:
    if date_from is None or date_to is None:
        return None
    left = _format_month_year(date_from)
    right = _format_month_year(date_to)
    return left if left == right else f"{left} – {right}"


def parse_master_workbook_bytes(content: bytes, *, db: Session | None = None) -> list[MasterRow]:
    logger.info("Workbook opening from bytes size=%s", len(content))
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        logger.info("Workbook opened successfully.")
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Workbook has no active worksheet.")
        logger.info("Worksheet found title=%s", sheet.title)

        header_map: dict[str, int] | None = None
        header_row = 1
        max_col = max(sheet.max_column or 1, 12)
        for row_idx in range(1, min(sheet.max_row or 1, 40) + 1):
            headers = [sheet.cell(row_idx, col).value for col in range(1, max_col + 1)]
            try:
                header_map = _detect_master_header(
                    [str(h) if h is not None else "" for h in headers]
                )
                header_row = row_idx
                break
            except ValueError:
                continue
        if header_map is None:
            raise ValueError("Could not locate header row in master workbook.")
        logger.info("Header row detected at row=%s", header_row)

        def read_col(values: list[object], field: str) -> object | None:
            idx = header_map.get(field)  # type: ignore[union-attr]
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        rows: list[MasterRow] = []
        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            values = [sheet.cell(row_idx, col).value for col in range(1, max_col + 1)]
            if _is_blank_row(values):
                continue

            project_value = _cell_text(read_col(values, "project"))
            is_np = False
            np_code = None
            if db is not None and project_value:
                is_np, np_code = _classify_project_value(db, project_value)
            elif project_value:
                code = project_value.strip().upper()
                if code in KNOWN_NP_CODES or code.startswith("C") or code.startswith("EST"):
                    is_np = True
                    np_code = code

            rows.append(
                MasterRow(
                    row_number=row_idx,
                    designer=_cell_text(read_col(values, "designer")),
                    entry_date=_parse_date(read_col(values, "entry_date")),
                    project_value=project_value,
                    task=_cell_text(read_col(values, "task")),
                    billable=_parse_billable(read_col(values, "billable")),
                    hours=_parse_decimal(read_col(values, "hours")),
                    notes=_cell_text(read_col(values, "notes")),
                    is_np_row=is_np,
                    np_code=np_code,
                )
            )
        logger.info("Workbook parsing complete rows=%s", len(rows))
        return rows
    finally:
        workbook.close()


def scan_master_workbook(db: Session, upload_id: str) -> MasterScanResponse:
    workbook_path, filename = get_master_upload(upload_id)
    rows = parse_master_workbook_bytes(workbook_path.read_bytes(), db=db)

    by_designer: dict[str, list[MasterRow]] = defaultdict(list)
    all_dates: list[date] = []
    for row in rows:
        designer = row.designer or "Unknown"
        by_designer[designer].append(row)
        if row.entry_date is not None:
            all_dates.append(row.entry_date)

    designers: list[MasterDesignerScanRow] = []
    for designer_name in sorted(by_designer):
        designer_rows = by_designer[designer_name]
        dates = [row.entry_date for row in designer_rows if row.entry_date is not None]
        date_from = min(dates) if dates else None
        date_to = max(dates) if dates else None
        designers.append(
            MasterDesignerScanRow(
                designer=designer_name,
                rows=len(designer_rows),
                date_from=date_from,
                date_to=date_to,
                date_range_label=_format_date_range(date_from, date_to) or "—",
                user_matched=_match_user_by_name(db, designer_name) is not None,
            )
        )

    date_from = min(all_dates) if all_dates else None
    date_to = max(all_dates) if all_dates else None
    return MasterScanResponse(
        upload_id=upload_id,
        filename=filename,
        row_count=len(rows),
        designer_count=len(by_designer),
        date_from=date_from,
        date_to=date_to,
        date_range_label=_format_date_range(date_from, date_to),
        designers=designers,
    )


def master_import_log_to_excel(log_rows: list[MasterImportLogRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import Log"
    sheet.append(["Row", "Designer", "Project", "Error"])
    for row in log_rows:
        sheet.append([row.row, row.designer or "", row.project or "", row.error])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _log_error(
    log_rows: list[MasterImportLogRow],
    *,
    row_number: int,
    designer: str | None,
    project: str | None,
    error: str,
) -> None:
    log_rows.append(
        MasterImportLogRow(
            row=row_number,
            designer=designer,
            project=project,
            error=error,
        )
    )


def run_master_import(
    db: Session,
    *,
    upload_id: str,
    workbook_path_override: str | Path | None = None,
    selected_designers: set[str] | None = None,
    progress_callback: Callable[..., None] | None = None,
    cancel_check: Callable[[], bool] | None = None,
    backup_path: Path | None = None,
) -> tuple[MasterImportSummary, list[MasterImportLogRow]]:
    started = datetime.now(timezone.utc)
    logger.info("Master import started upload_id=%s", upload_id)
    if workbook_path_override is not None:
        workbook_path = Path(workbook_path_override)
    else:
        workbook_path, _ = get_master_upload(upload_id)
    logger.info("Workbook path resolved path=%s", workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook file does not exist: {workbook_path}")
    logger.info("Reading workbook bytes...")
    rows = parse_master_workbook_bytes(workbook_path.read_bytes(), db=db)
    logger.info("Workbook rows loaded count=%s", len(rows))
    if progress_callback:
        progress_callback(
            percent_complete=1,
            rows_processed=0,
            rows_imported=0,
            rows_total=len(rows),
            message="Workbook opened. Matching designers...",
        )

    if selected_designers is not None:
        selected = {name.strip().lower() for name in selected_designers}
        rows = [row for row in rows if (row.designer or "").strip().lower() in selected]
        logger.info("Filtered rows by selected designers count=%s", len(rows))

    summary = MasterImportSummary(
        backup_path=str(backup_path) if backup_path else None,
        duplicate_check_disabled=True,
    )
    log_rows: list[MasterImportLogRow] = []
    pending_entries: list[TimesheetEntry] = []
    affected_projects: set[uuid.UUID] = set()
    rows_total = len(rows)
    rows_processed = 0
    logger.info("Beginning row processing rows_total=%s", rows_total)

    def flush_batch() -> None:
        nonlocal pending_entries
        if not pending_entries:
            return
        try:
            db.add_all(pending_entries)
            db.commit()
            pending_entries = []
        except Exception:
            db.rollback()
            raise

    try:
        for row in rows:
            if cancel_check and cancel_check():
                flush_batch()
                break

            summary.rows_read += 1
            rows_processed += 1
            designer_name = row.designer or "Unknown"

            if progress_callback:
                progress_callback(
                    percent_complete=int((rows_processed / rows_total) * 100) if rows_total else 0,
                    rows_processed=rows_processed,
                    rows_imported=summary.rows_imported,
                    rows_total=rows_total,
                    current_designer=designer_name,
                    message="Importing rows…",
                )
            if rows_processed == 1 or rows_processed % PROGRESS_LOG_EVERY_ROWS == 0:
                logger.info(
                    "Processed row checkpoint processed=%s imported=%s designer=%s",
                    rows_processed,
                    summary.rows_imported,
                    designer_name,
                )

            designer_user = _match_user_by_name(db, designer_name)
            if designer_user is None:
                summary.errors += 1
                summary.rows_skipped += 1
                _log_error(
                    log_rows,
                    row_number=row.row_number,
                    designer=designer_name,
                    project=row.project_value,
                    error="Unknown Designer",
                )
                continue

            if row.entry_date is None:
                summary.errors += 1
                summary.rows_skipped += 1
                _log_error(
                    log_rows,
                    row_number=row.row_number,
                    designer=designer_name,
                    project=row.project_value,
                    error="Invalid Date",
                )
                continue

            is_np, np_code = _classify_project_value(db, row.project_value)
            row.is_np_row = is_np
            row.np_code = np_code

            if row.is_np_row and row.np_code and row.hours is None:
                row.hours = Decimal("0")

            if row.hours is None or (row.hours <= 0 and not row.is_np_row):
                summary.errors += 1
                summary.rows_skipped += 1
                _log_error(
                    log_rows,
                    row_number=row.row_number,
                    designer=designer_name,
                    project=row.project_value,
                    error="Invalid Hours",
                )
                continue

            if row.billable is None and not row.is_np_row:
                summary.errors += 1
                summary.rows_skipped += 1
                _log_error(
                    log_rows,
                    row_number=row.row_number,
                    designer=designer_name,
                    project=row.project_value,
                    error="Invalid Billable value",
                )
                continue

            week = _week_start(row.entry_date)
            timesheet = _get_or_create_approved_timesheet(db, designer_user.id, week)

            if row.is_np_row:
                if not row.np_code:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Unknown Non-Productive Code",
                    )
                    continue
                try:
                    np_record = _resolve_np_code(db, row.np_code)
                except ValueError:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Unknown Non-Productive Code",
                    )
                    continue
                entry = build_np_timesheet_entry(
                    np_code=np_record,
                    timesheet_id=timesheet.id,
                    entry_date=row.entry_date,
                    hours=row.hours,
                    description=row.notes,
                    is_billable=row.billable,
                    allow_billable_override=True,
                )
            else:
                if not row.project_value:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Project not found",
                    )
                    continue
                if not row.task:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Missing Task",
                    )
                    continue

                project = _find_project(db, row.project_value, None)
                if project is None:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Project not found",
                    )
                    continue

                task_name = _resolve_task_type_name(row.task)
                task_type = _find_task_type(db, task_name, project.stream_id)
                if task_type is None:
                    summary.errors += 1
                    summary.rows_skipped += 1
                    _log_error(
                        log_rows,
                        row_number=row.row_number,
                        designer=designer_name,
                        project=row.project_value,
                        error="Unknown Task",
                    )
                    continue

                entry = TimesheetEntry(
                    timesheet_id=timesheet.id,
                    work_category=WorkCategory.productive,
                    project_id=project.id,
                    customer_id=project.customer_id,
                    task_type_id=task_type.id,
                    is_billable=row.billable,
                    entry_date=row.entry_date,
                    hours=row.hours,
                    description=row.notes,
                )
                affected_projects.add(project.id)

            pending_entries.append(entry)
            summary.rows_imported += 1

            if len(pending_entries) >= BATCH_COMMIT_SIZE:
                logger.info("Flushing pending batch size=%s", len(pending_entries))
                if progress_callback:
                    progress_callback(
                        percent_complete=int((rows_processed / rows_total) * 100) if rows_total else 0,
                        rows_processed=rows_processed,
                        rows_imported=summary.rows_imported,
                        rows_total=rows_total,
                        current_designer=designer_name,
                        message="Saving entries...",
                    )
                flush_batch()

        if progress_callback:
            progress_callback(
                percent_complete=99,
                rows_processed=rows_processed,
                rows_imported=summary.rows_imported,
                rows_total=rows_total,
                message="Finalizing import...",
            )
        logger.info("Final batch flush pending=%s", len(pending_entries))
        flush_batch()
        for project_id in affected_projects:
            recalculate_project(db, project_id)

        summary.duration_seconds = int((datetime.now(timezone.utc) - started).total_seconds())
        logger.info(
            "Master import completed rows_read=%s imported=%s skipped=%s errors=%s",
            summary.rows_read,
            summary.rows_imported,
            summary.rows_skipped,
            summary.errors,
        )
        return summary, log_rows
    except Exception:
        db.rollback()
        logger.exception("Master import failed upload_id=%s", upload_id)
        raise

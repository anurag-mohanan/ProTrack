"""Historical timesheet import from Excel/CSV designer workbooks."""

from __future__ import annotations

import csv
import io
import json
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import gettempdir
from typing import Any, Callable

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.permissions import DESIGNER, PROJECT_STAFF_ROLES
from app.core.security import hash_password
from app.models.enums import ExecutionStatus, TimesheetStatus, WorkCategory
from app.models.models import (
    Contact,
    Customer,
    NonProductiveCode,
    Project,
    Role,
    Stream,
    TaskType,
    Timesheet,
    TimesheetEntry,
    TimesheetImportHistory,
    User,
)
from app.schemas.historical_import import ImportSummary
from app.schemas.historical_timesheet_import import (
    CustomerRowResolution,
    DesignerMatchInfo,
    DesignerResolution,
    DesignerResolutionAction,
    DuplicateWeekAction,
    DuplicateWeekInfo,
    ProjectResolutionAction,
    ProjectRowResolution,
    TaskTypeRowResolution,
    TimesheetImportRowPreview,
    TimesheetImportRowStatus,
    TimesheetImportSummary,
    TimesheetImportUploadResponse,
    TimesheetImportValidateResponse,
    TimesheetValidationIssue,
)
from app.services.historical_import_service import (
    DEFAULT_IMPORT_PASSWORD,
    KNOWN_NP_CODES,
    _get_or_create_approved_timesheet,
    _normalize_header,
    _normalize_key,
    _resolve_np_code,
    _resolve_or_create_customer,
    _split_name,
    _unique_email,
    _week_start,
)
from app.services.project_calculation_service import recalculate_project
from app.services.project_template_service import (
    create_milestones_from_template,
    resolve_template_for_import,
)

UPLOAD_DIR = Path(gettempdir()) / "protrack_timesheet_imports"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv"}

TIMESHEET_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "designer": ("designer", "employee", "staff name", "resource"),
    "tool_number": (
        "tool no.",
        "tool no",
        "tool number",
        "project number",
        "project no",
        "tool #",
        "tool",
    ),
    "project_code": ("project code", "code", "job code"),
    "customer": ("customer", "customer name"),
    "task_type": ("task", "task type", "activity", "work type"),
    "np_code": ("np code", "non productive code", "non-productive code", "np"),
    "hours": ("hours", "time", "logged hours", "hrs", "actual hours"),
    "entry_date": ("entry date", "work date", "timesheet date", "date"),
    "description": ("description", "notes", "comments", "part description"),
}

TASK_TYPE_ALIASES: dict[str, str] = {
    "design": "Design",
    "surfacing": "Surfacing",
    "surface": "Surfacing",
    "feasibility": "Feasibility",
    "engineering change": "Engineering Change (EC)",
    "engineering change (ec)": "Engineering Change (EC)",
    "ec": "Engineering Change (EC)",
    "2d drawings": "2D Drawings",
    "2d drawing": "2D Drawings",
    "drawings": "2D Drawings",
}


@dataclass
class ParsedTimesheetRow:
    row_number: int
    designer: str | None = None
    tool_number: str | None = None
    project_code: str | None = None
    customer: str | None = None
    task_type: str | None = None
    np_code: str | None = None
    hours: Decimal | None = None
    entry_date: date | None = None
    description: str | None = None
    is_np_row: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class ImportContext:
    designer_resolution: DesignerResolution | None = None
    duplicate_week_action: DuplicateWeekAction = DuplicateWeekAction.skip
    project_resolutions: dict[int, ProjectRowResolution] = field(default_factory=dict)
    customer_resolutions: dict[int, CustomerRowResolution] = field(default_factory=dict)
    task_type_resolutions: dict[int, TaskTypeRowResolution] = field(default_factory=dict)


def _cell_text(value: Any) -> str | None:
    from app.core.field_normalization import normalize_optional_text

    return normalize_optional_text(value)


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _hours_key(hours: Decimal) -> str:
    return format(hours.quantize(Decimal("0.01")), "f")


def _notes_key(notes: str | None) -> str:
    return (notes or "").strip()


def timesheet_duplicate_key(
    *,
    user_id: uuid.UUID,
    entry_date: date,
    project_number: str | None = None,
    np_code: str | None = None,
    task: str | None = None,
    hours: Decimal,
    is_billable: bool,
    notes: str | None = None,
) -> tuple:
    """Duplicate when designer, date, project/NP, task, hours, billable, notes, month, and year all match."""
    project_key = (np_code or project_number or "").strip().upper()
    task_key = (task or "").strip().lower()
    return (
        user_id,
        entry_date,
        project_key,
        task_key,
        _hours_key(hours),
        is_billable,
        _notes_key(notes),
        entry_date.year,
        entry_date.month,
    )


def timesheet_duplicate_key_for_row(
    *,
    user_id: uuid.UUID,
    entry_date: date,
    tool_number: str | None,
    np_code: str | None,
    task_type: str | None,
    hours: Decimal,
    is_billable: bool,
    description: str | None,
    is_np_row: bool,
) -> tuple:
    return timesheet_duplicate_key(
        user_id=user_id,
        entry_date=entry_date,
        project_number=None if is_np_row else tool_number,
        np_code=np_code if is_np_row else None,
        task=None if is_np_row else task_type,
        hours=hours,
        is_billable=is_billable,
        notes=description,
    )


def load_timesheet_duplicate_keys(
    db: Session,
    user_id: uuid.UUID | None = None,
    *,
    year: int | None = None,
    month: int | None = None,
) -> set[tuple]:
    from sqlalchemy import extract, func, select

    query = (
        select(
            Timesheet.user_id,
            TimesheetEntry.entry_date,
            TimesheetEntry.hours,
            TimesheetEntry.is_billable,
            TimesheetEntry.description,
            TimesheetEntry.non_productive_code_id,
            Project.tool_number,
            NonProductiveCode.code,
            TaskType.name,
        )
        .join(Timesheet, TimesheetEntry.timesheet_id == Timesheet.id)
        .outerjoin(Project, TimesheetEntry.project_id == Project.id)
        .outerjoin(NonProductiveCode, TimesheetEntry.non_productive_code_id == NonProductiveCode.id)
        .outerjoin(TaskType, TimesheetEntry.task_type_id == TaskType.id)
    )
    if user_id is not None:
        query = query.where(Timesheet.user_id == user_id)
    if year is not None:
        query = query.where(extract("year", TimesheetEntry.entry_date) == year)
    if month is not None:
        query = query.where(extract("month", TimesheetEntry.entry_date) == month)

    keys: set[tuple] = set()
    for row in db.execute(query).all():
        is_np = row.non_productive_code_id is not None
        keys.add(
            timesheet_duplicate_key(
                user_id=row.user_id,
                entry_date=row.entry_date,
                project_number=None if is_np else row.tool_number,
                np_code=row.code if is_np else None,
                task=None if is_np else row.name,
                hours=row.hours,
                is_billable=bool(row.is_billable),
                notes=row.description,
            )
        )
    return keys


def designer_has_entries_for_month(
    db: Session,
    user_id: uuid.UUID,
    year: int,
    month: int,
) -> bool:
    from sqlalchemy import extract, func, select

    count = db.scalar(
        select(func.count())
        .select_from(TimesheetEntry)
        .join(Timesheet, TimesheetEntry.timesheet_id == Timesheet.id)
        .where(
            Timesheet.user_id == user_id,
            extract("year", TimesheetEntry.entry_date) == year,
            extract("month", TimesheetEntry.entry_date) == month,
        )
    )
    return int(count or 0) > 0


def _normalize_np_code(value: str | None) -> str | None:
    if not value:
        return None
    code = value.strip().upper()
    return code if code in KNOWN_NP_CODES else None


def _detect_timesheet_header(headers: list[str]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    normalized = [_normalize_header(h) for h in headers]
    for field_name, aliases in TIMESHEET_COLUMN_ALIASES.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if header in aliases or any(alias in header for alias in aliases):
                header_map[field_name] = idx
                break
    has_hours = "hours" in header_map or "entry_date" in header_map
    has_work = "tool_number" in header_map or "np_code" in header_map or "task_type" in header_map
    if has_hours and (has_work or "designer" in header_map):
        return header_map
    raise ValueError(
        "Could not locate a header row. Expected columns such as Date, Hours, "
        "Designer, Project/Tool Number, or NP Code."
    )


def _row_from_values(row_number: int, values: list[Any], header_map: dict[str, int]) -> ParsedTimesheetRow:
    def read(field: str) -> Any:
        idx = header_map.get(field)
        if idx is None or idx >= len(values):
            return None
        return values[idx]

    row = ParsedTimesheetRow(row_number=row_number)
    row.designer = _cell_text(read("designer"))
    row.tool_number = _cell_text(read("tool_number"))
    row.project_code = _cell_text(read("project_code"))
    row.customer = _cell_text(read("customer"))
    row.task_type = _cell_text(read("task_type"))
    row.np_code = _normalize_np_code(_cell_text(read("np_code")))
    row.description = _cell_text(read("description"))
    row.hours = _parse_decimal(read("hours"))
    row.entry_date = _parse_date(read("entry_date"))

    if row.np_code or not row.tool_number:
        row.is_np_row = True
    else:
        row.is_np_row = False

    _validate_row(row)
    return row


def _validate_row(row: ParsedTimesheetRow) -> None:
    if row.hours is None:
        row.errors.append("Hours are missing or invalid.")
    elif row.hours <= 0:
        row.errors.append("Hours must be greater than zero.")
    elif row.hours > 24:
        row.errors.append("Hours cannot exceed 24 per entry.")

    if row.entry_date is None:
        row.errors.append("Entry date is missing or invalid.")
    elif row.entry_date > date.today() + timedelta(days=7):
        row.warnings.append("Entry date is in the future.")

    if row.is_np_row:
        if row.np_code is None and not row.tool_number:
            row.errors.append("NP Code is required when project number is blank.")
        elif row.np_code is None and row.tool_number:
            row.np_code = _normalize_np_code(row.tool_number)
            if row.np_code is None:
                row.errors.append(f"Invalid NP Code: {row.tool_number}.")
        return

    if not row.tool_number and not row.project_code:
        row.errors.append("Project number or project code is required.")
    if not row.task_type:
        row.errors.append("Task type is required for productive entries.")
    elif _resolve_task_type_name(row.task_type) is None:
        row.warnings.append(f"Task type '{row.task_type}' may need manual mapping.")


def _resolve_task_type_name(value: str | None) -> str | None:
    if not value:
        return None
    normalized = _normalize_key(value)
    if normalized in TASK_TYPE_ALIASES:
        return TASK_TYPE_ALIASES[normalized]
    for canonical in TASK_TYPE_ALIASES.values():
        if _normalize_key(canonical) == normalized:
            return canonical
    return None


def parse_csv(file_path: Path) -> list[ParsedTimesheetRow]:
    content = file_path.read_text(encoding="utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows_list = list(reader)
    if not rows_list:
        return []

    header_map = None
    header_row_idx = 0
    for idx, row in enumerate(rows_list[:25]):
        try:
            header_map = _detect_timesheet_header(row)
            header_row_idx = idx
            break
        except ValueError:
            continue
    if header_map is None:
        raise ValueError("Could not locate a header row in CSV file.")

    parsed: list[ParsedTimesheetRow] = []
    for row_idx, values in enumerate(rows_list[header_row_idx + 1 :], start=header_row_idx + 2):
        if not any(str(v).strip() for v in values if v is not None):
            continue
        parsed.append(_row_from_values(row_idx, values, header_map))
    return parsed


def parse_workbook(file_path: Path) -> list[ParsedTimesheetRow]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Workbook has no active worksheet.")

        header_map: dict[str, int] | None = None
        header_row = 1
        for row_idx in range(1, min(sheet.max_row, 25) + 1):
            headers = [
                _normalize_header(sheet.cell(row_idx, col).value)
                for col in range(1, sheet.max_column + 1)
            ]
            try:
                csv_header_map = _detect_timesheet_header(
                    [sheet.cell(row_idx, col).value for col in range(1, sheet.max_column + 1)]
                )
                header_map = {k: v + 1 for k, v in csv_header_map.items()}
                header_row = row_idx
                break
            except ValueError:
                continue
        if header_map is None:
            raise ValueError("Could not locate a header row in workbook.")

        parsed: list[ParsedTimesheetRow] = []
        for row_idx in range(header_row + 1, sheet.max_row + 1):
            values = [sheet.cell(row_idx, col).value for col in range(1, sheet.max_column + 1)]
            if not any(v is not None and str(v).strip() for v in values):
                continue
            csv_map = {k: v - 1 for k, v in header_map.items()}
            parsed.append(_row_from_values(row_idx, values, csv_map))
        return parsed
    finally:
        workbook.close()


def parse_upload_file(file_path: Path) -> list[ParsedTimesheetRow]:
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return parse_csv(file_path)
    return parse_workbook(file_path)


def save_upload(upload_id: str, file_name: str, content: bytes) -> Path:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Only .xlsx, .xlsm, and .csv files are supported.")
    path = UPLOAD_DIR / f"{upload_id}{suffix}"
    path.write_bytes(content)
    meta_path = UPLOAD_DIR / f"{upload_id}.meta"
    meta_path.write_text(file_name, encoding="utf-8")
    return path


def get_upload_path(upload_id: str) -> Path:
    for suffix in SUPPORTED_SUFFIXES:
        path = UPLOAD_DIR / f"{upload_id}{suffix}"
        if path.exists():
            return path
    raise FileNotFoundError("Upload not found or expired.")


def _resolutions_path(upload_id: str) -> Path:
    return UPLOAD_DIR / f"{upload_id}.resolutions.json"


def save_resolutions(upload_id: str, context: ImportContext) -> None:
    payload = {
        "duplicate_week_action": context.duplicate_week_action.value,
        "designer": context.designer_resolution.model_dump(mode="json") if context.designer_resolution else None,
        "project_resolutions": [r.model_dump(mode="json") for r in context.project_resolutions.values()],
        "customer_resolutions": [r.model_dump(mode="json") for r in context.customer_resolutions.values()],
        "task_type_resolutions": [r.model_dump(mode="json") for r in context.task_type_resolutions.values()],
    }
    _resolutions_path(upload_id).write_text(json.dumps(payload), encoding="utf-8")


def load_resolutions(upload_id: str) -> ImportContext:
    path = _resolutions_path(upload_id)
    if not path.exists():
        return ImportContext()
    data = json.loads(path.read_text(encoding="utf-8"))
    context = ImportContext(
        duplicate_week_action=DuplicateWeekAction(data.get("duplicate_week_action", "skip")),
    )
    if data.get("designer"):
        context.designer_resolution = DesignerResolution.model_validate(data["designer"])
    for item in data.get("project_resolutions", []):
        resolution = ProjectRowResolution.model_validate(item)
        context.project_resolutions[resolution.row_number] = resolution
    for item in data.get("customer_resolutions", []):
        resolution = CustomerRowResolution.model_validate(item)
        context.customer_resolutions[resolution.row_number] = resolution
    for item in data.get("task_type_resolutions", []):
        resolution = TaskTypeRowResolution.model_validate(item)
        context.task_type_resolutions[resolution.row_number] = resolution
    return context


def _match_user_by_name(db: Session, name: str) -> User | None:
    normalized = _normalize_key(name)
    users = db.scalars(
        select(User)
        .join(Role, User.role_id == Role.id)
        .where(User.is_active.is_(True))
    ).all()
    for user in users:
        full = _normalize_key(f"{user.first_name} {user.last_name}")
        first = _normalize_key(user.first_name)
        last = _normalize_key(user.last_name)
        if normalized in {full, first, last, _normalize_key(user.email.split("@")[0])}:
            return user
    return None


def _detect_designer(rows: list[ParsedTimesheetRow]) -> str:
    names = [row.designer for row in rows if row.designer]
    if not names:
        return "Unknown"
    return Counter(names).most_common(1)[0][0]


def _find_project(db: Session, tool_number: str | None, project_code: str | None) -> Project | None:
    if tool_number:
        normalized = tool_number.strip().lower()
        for project in db.scalars(select(Project).where(Project.is_deleted.is_(False))).all():
            if project.tool_number and project.tool_number.strip().lower() == normalized:
                return project
    if project_code:
        normalized = project_code.strip().lower()
        project = db.scalar(
            select(Project).where(
                func.lower(Project.code) == normalized,
                Project.is_deleted.is_(False),
            )
        )
        if project:
            return project
    return None


def _find_task_type(db: Session, name: str, stream_id: uuid.UUID) -> TaskType | None:
    canonical = _resolve_task_type_name(name)
    if canonical is None:
        return None
    return db.scalar(
        select(TaskType).where(
            TaskType.name == canonical,
            TaskType.stream_id == stream_id,
            TaskType.is_active.is_(True),
        )
    )


def _week_label(week_start: date) -> str:
    return week_start.strftime("%d-%b-%Y")


def _detect_duplicate_weeks(
    db: Session,
    rows: list[ParsedTimesheetRow],
    designer_name: str,
    designer_user_id: uuid.UUID | None,
) -> list[DuplicateWeekInfo]:
    if designer_user_id is None:
        return []

    week_starts: set[date] = set()
    for row in rows:
        if row.entry_date and not row.errors:
            week_starts.add(_week_start(row.entry_date))

    duplicates: list[DuplicateWeekInfo] = []
    for week in sorted(week_starts):
        entry_count = db.scalar(
            select(func.count())
            .select_from(TimesheetEntry)
            .join(Timesheet, TimesheetEntry.timesheet_id == Timesheet.id)
            .where(
                Timesheet.user_id == designer_user_id,
                Timesheet.week_start == week,
            )
        ) or 0
        history = db.scalar(
            select(TimesheetImportHistory)
            .where(
                TimesheetImportHistory.designer_user_id == designer_user_id,
                TimesheetImportHistory.week_start == week,
                TimesheetImportHistory.status == "completed",
            )
            .order_by(TimesheetImportHistory.created_at.desc())
        )
        if entry_count > 0 or history is not None:
            duplicates.append(
                DuplicateWeekInfo(
                    week_start=week,
                    week_label=_week_label(week),
                    designer_name=designer_name,
                    existing_entry_count=int(entry_count),
                    previously_imported=history is not None,
                    import_history_id=history.id if history else None,
                )
            )
    return duplicates


def _preview_from_row(
    db: Session,
    row: ParsedTimesheetRow,
    *,
    designer_user: User | None = None,
) -> TimesheetImportRowPreview:
    if row.errors:
        status = TimesheetImportRowStatus.error
    elif row.warnings:
        status = TimesheetImportRowStatus.warning
    else:
        status = TimesheetImportRowStatus.ready

    matched_project = None
    if not row.is_np_row and not row.errors:
        matched_project = _find_project(db, row.tool_number, row.project_code)

    matched_task = None
    if row.task_type and matched_project:
        task = _find_task_type(db, row.task_type, matched_project.stream_id)
        if task:
            matched_task = task.id

    messages = list(row.errors) + list(row.warnings)
    if row.is_np_row and row.np_code:
        messages.insert(0, f"Non-productive ({row.np_code})")

    return TimesheetImportRowPreview(
        row_number=row.row_number,
        designer=row.designer or (designer_user and f"{designer_user.first_name} {designer_user.last_name}"),
        tool_number=row.tool_number,
        project_code=row.project_code,
        customer=row.customer,
        task_type=row.task_type,
        np_code=row.np_code,
        hours=str(row.hours) if row.hours is not None else None,
        entry_date=row.entry_date.isoformat() if row.entry_date else None,
        description=row.description,
        is_np_row=row.is_np_row,
        status_label=status,
        messages=messages,
        matched_project_id=matched_project.id if matched_project else None,
        matched_user_id=designer_user.id if designer_user else None,
        matched_task_type_id=matched_task,
    )


def analyze_upload(db: Session, upload_id: str) -> TimesheetImportUploadResponse:
    path = get_upload_path(upload_id)
    rows = parse_upload_file(path)
    designer_name = _detect_designer(rows)
    matched_user = _match_user_by_name(db, designer_name)

    designer_info = DesignerMatchInfo(
        detected_name=designer_name,
        matched_user_id=matched_user.id if matched_user else None,
        matched_user_name=(
            f"{matched_user.first_name} {matched_user.last_name}" if matched_user else None
        ),
        requires_resolution=matched_user is None,
    )

    preview = [_preview_from_row(db, row, designer_user=matched_user) for row in rows]
    ready_rows = sum(1 for p in preview if p.status_label == TimesheetImportRowStatus.ready)
    error_rows = sum(1 for p in preview if p.status_label == TimesheetImportRowStatus.error)
    warning_rows = sum(1 for p in preview if p.status_label == TimesheetImportRowStatus.warning)
    skipped_rows = 0

    duplicate_weeks = _detect_duplicate_weeks(
        db,
        rows,
        designer_name,
        matched_user.id if matched_user else None,
    )

    meta_path = UPLOAD_DIR / f"{upload_id}.meta"
    file_name = meta_path.read_text(encoding="utf-8") if meta_path.exists() else path.name

    return TimesheetImportUploadResponse(
        upload_id=upload_id,
        file_name=file_name,
        total_rows=len(rows),
        ready_rows=ready_rows,
        error_rows=error_rows,
        warning_rows=warning_rows,
        skipped_rows=skipped_rows,
        designer=designer_info,
        duplicate_weeks=duplicate_weeks,
        preview=preview,
    )


def validate_upload(db: Session, upload_id: str) -> TimesheetImportValidateResponse:
    path = get_upload_path(upload_id)
    rows = parse_upload_file(path)
    issues: list[TimesheetValidationIssue] = []
    seen_keys: set[tuple] = set()

    for row in rows:
        for message in row.errors:
            issues.append(
                TimesheetValidationIssue(
                    row_number=row.row_number,
                    severity="error",
                    code="row_error",
                    message=message,
                )
            )
        for message in row.warnings:
            issues.append(
                TimesheetValidationIssue(
                    row_number=row.row_number,
                    severity="warning",
                    code="row_warning",
                    message=message,
                )
            )

        if row.entry_date and row.hours is not None and not row.errors:
            is_billable = False if row.is_np_row else True
            dup_key = (
                _normalize_key(row.designer or _detect_designer(rows)),
                row.entry_date,
                (row.np_code or row.tool_number or "").strip().upper(),
                (row.task_type or "").strip().lower() if not row.is_np_row else "",
                _hours_key(row.hours),
                is_billable,
                _notes_key(row.description),
                row.entry_date.year,
                row.entry_date.month,
            )
            if dup_key in seen_keys:
                issues.append(
                    TimesheetValidationIssue(
                        row_number=row.row_number,
                        severity="error",
                        code="duplicate_entry",
                        message="Duplicate entry within the file.",
                    )
                )
            seen_keys.add(dup_key)

        if not row.is_np_row and row.tool_number and not row.errors:
            project = _find_project(db, row.tool_number, row.project_code)
            if project is None:
                issues.append(
                    TimesheetValidationIssue(
                        row_number=row.row_number,
                        severity="warning",
                        code="missing_project",
                        message=f"Project '{row.tool_number or row.project_code}' not found.",
                    )
                )

        if row.is_np_row and row.np_code and row.np_code not in KNOWN_NP_CODES:
            issues.append(
                TimesheetValidationIssue(
                    row_number=row.row_number,
                    severity="error",
                    code="invalid_np_code",
                    message=f"Invalid NP Code: {row.np_code}.",
                )
            )

    error_rows = sum(1 for issue in issues if issue.severity == "error")
    ready_rows = len(rows) - error_rows
    return TimesheetImportValidateResponse(
        upload_id=upload_id,
        is_valid=error_rows == 0,
        issues=issues,
        ready_rows=max(ready_rows, 0),
        error_rows=error_rows,
    )


def _resolve_designer_user(
    db: Session,
    designer_name: str,
    resolution: DesignerResolution | None,
    summary: TimesheetImportSummary,
) -> User:
    if resolution and resolution.action == DesignerResolutionAction.match_existing:
        if resolution.user_id is None:
            raise ValueError("user_id is required when matching an existing designer.")
        user = db.get(User, resolution.user_id)
        if user is None or not user.is_active:
            raise ValueError("Selected designer user was not found.")
        return user

    matched = _match_user_by_name(db, designer_name)
    if matched is not None and (
        resolution is None or resolution.action != DesignerResolutionAction.create_new
    ):
        return matched

    if resolution and resolution.action == DesignerResolutionAction.create_new:
        role = db.scalar(select(Role).where(Role.name == DESIGNER))
        if role is None:
            raise ValueError("Designer role is not configured.")
        first_name, last_name = _split_name(designer_name)
        user = User(
            role_id=role.id,
            email=_unique_email(db, first_name, last_name),
            password_hash=hash_password(DEFAULT_IMPORT_PASSWORD),
            first_name=first_name,
            last_name=last_name,
            is_active=True,
        )
        db.add(user)
        db.flush()
        summary.users_created += 1
        return user

    raise ValueError(f"Designer '{designer_name}' was not found. Provide a resolution.")


def _create_minimal_project(
    db: Session,
    row: ParsedTimesheetRow,
    summary: TimesheetImportSummary,
    customer_resolution: CustomerRowResolution | None,
) -> Project:
    import_summary = ImportSummary()
    if customer_resolution and customer_resolution.customer_id:
        customer = db.get(Customer, customer_resolution.customer_id)
        if customer is None:
            raise ValueError("Mapped customer was not found.")
    elif row.customer:
        customer = _resolve_or_create_customer(db, row.customer, import_summary)
    else:
        raise ValueError("Customer is required to create a project.")

    summary.customers_created += import_summary.customers_created

    stream = db.scalar(select(Stream).where(Stream.is_active.is_(True)).order_by(Stream.name))
    if stream is None:
        raise ValueError("No active stream is configured.")

    design_leader = db.scalar(
        select(User)
        .join(Role, User.role_id == Role.id)
        .where(Role.name == "Design Leader", User.is_active.is_(True))
        .order_by(User.last_name)
    )
    if design_leader is None:
        raise ValueError("No Design Leader is configured.")

    tool_number = (row.tool_number or row.project_code or "IMPORT").strip()[:50]
    code_base = (row.project_code or tool_number).strip()[:50]
    code = code_base
    suffix = 1
    while db.scalar(select(Project.id).where(Project.code == code)):
        code = f"{code_base}-{suffix}"[:50]
        suffix += 1

    template = resolve_template_for_import(db, customer=customer)
    project = Project(
        tool_number=tool_number,
        part_description=row.description or "Imported from timesheet",
        customer_id=customer.id,
        design_leader_id=design_leader.id,
        stream_id=stream.id,
        code=code,
        quoted_hours=Decimal("0"),
        actual_hours=Decimal("0"),
        due_date=row.entry_date or date.today(),
        execution_status=ExecutionStatus.currently_being_worked_on,
    )
    if template is not None:
        project.project_type_id = template.project_type_id
    db.add(project)
    db.flush()
    create_milestones_from_template(db, project=project, template=template)
    summary.projects_created += 1
    return project


def _resolve_project_for_row(
    db: Session,
    row: ParsedTimesheetRow,
    context: ImportContext,
    summary: TimesheetImportSummary,
) -> Project | None:
    resolution = context.project_resolutions.get(row.row_number)
    existing = _find_project(db, row.tool_number, row.project_code)

    if resolution:
        if resolution.action == ProjectResolutionAction.skip:
            return None
        if resolution.action == ProjectResolutionAction.map_existing:
            if resolution.project_id is None:
                raise ValueError(f"Row {row.row_number}: project_id required for mapping.")
            project = db.get(Project, resolution.project_id)
            if project is None:
                raise ValueError(f"Row {row.row_number}: mapped project not found.")
            summary.projects_matched += 1
            return project
        if resolution.action == ProjectResolutionAction.create:
            return _create_minimal_project(
                db,
                row,
                summary,
                context.customer_resolutions.get(row.row_number),
            )

    if existing is not None:
        summary.projects_matched += 1
        return existing
    return None


def _resolve_task_type_for_row(
    db: Session,
    row: ParsedTimesheetRow,
    project: Project,
    context: ImportContext,
) -> TaskType:
    resolution = context.task_type_resolutions.get(row.row_number)
    if resolution:
        task_type = db.get(TaskType, resolution.task_type_id)
        if task_type is None or not task_type.is_active:
            raise ValueError(f"Row {row.row_number}: task type not found.")
        return task_type

    if row.task_type:
        task_type = _find_task_type(db, row.task_type, project.stream_id)
        if task_type:
            return task_type
    raise ValueError(f"Row {row.row_number}: task type could not be resolved.")


def _clear_week_entries(db: Session, user_id: uuid.UUID, week: date) -> None:
    timesheet = db.scalar(
        select(Timesheet).where(Timesheet.user_id == user_id, Timesheet.week_start == week)
    )
    if timesheet is None:
        return
    db.execute(delete(TimesheetEntry).where(TimesheetEntry.timesheet_id == timesheet.id))


def run_timesheet_import(
    db: Session,
    upload_id: str,
    *,
    dry_run: bool = False,
    context: ImportContext | None = None,
    imported_by_id: uuid.UUID | None = None,
    progress_callback: Callable[[int, int], None] | None = None,
) -> tuple[TimesheetImportSummary, list[TimesheetImportRowPreview], uuid.UUID | None]:
    started = datetime.now(timezone.utc)
    path = get_upload_path(upload_id)
    rows = parse_upload_file(path)
    if context is None:
        context = load_resolutions(upload_id)

    summary = TimesheetImportSummary()
    error_log: list[TimesheetImportRowPreview] = []
    designer_name = _detect_designer(rows)

    try:
        designer = _resolve_designer_user(db, designer_name, context.designer_resolution, summary)
    except ValueError as exc:
        summary.errors = len(rows)
        for row in rows:
            preview = _preview_from_row(db, row)
            preview.status_label = TimesheetImportRowStatus.error
            preview.messages.append(str(exc))
            error_log.append(preview)
        return summary, error_log, None

    summary.designer = f"{designer.first_name} {designer.last_name}"
    summary.rows_read = len(rows)

    duplicate_weeks = _detect_duplicate_weeks(db, rows, designer_name, designer.id)
    skip_weeks: set[date] = set()
    if context.duplicate_week_action == DuplicateWeekAction.skip:
        skip_weeks = {item.week_start for item in duplicate_weeks}
    elif context.duplicate_week_action == DuplicateWeekAction.replace:
        for item in duplicate_weeks:
            _clear_week_entries(db, designer.id, item.week_start)

    entries_to_add: list[TimesheetEntry] = []
    affected_projects: set[uuid.UUID] = set()
    valid_dates = [row.entry_date for row in rows if row.entry_date and not row.errors]
    check_duplicates = False
    seen_duplicate_keys: set[tuple] = set()
    if valid_dates:
        month_counts = Counter((d.year, d.month) for d in valid_dates)
        workbook_year, workbook_month = month_counts.most_common(1)[0][0]
        if designer_has_entries_for_month(db, designer.id, workbook_year, workbook_month):
            check_duplicates = True
            seen_duplicate_keys = load_timesheet_duplicate_keys(
                db,
                designer.id,
                year=workbook_year,
                month=workbook_month,
            )

    total = len(rows)
    processed = 0

    try:
        for row in rows:
            processed += 1
            if progress_callback:
                progress_callback(processed, total)

            preview = _preview_from_row(db, row, designer_user=designer)
            if row.errors:
                summary.rows_failed += 1
                summary.errors += 1
                preview.status_label = TimesheetImportRowStatus.error
                error_log.append(preview)
                continue

            if row.entry_date and _week_start(row.entry_date) in skip_weeks:
                summary.rows_skipped += 1
                preview.status_label = TimesheetImportRowStatus.skipped
                preview.messages.append("Week already imported; skipped.")
                error_log.append(preview)
                continue

            if row.warnings:
                summary.warnings += 1

            assert row.entry_date is not None
            assert row.hours is not None
            week = _week_start(row.entry_date)
            timesheet = _get_or_create_approved_timesheet(db, designer.id, week)

            if row.is_np_row:
                assert row.np_code is not None
                np_code = _resolve_np_code(db, row.np_code)
                customer = None
                customer_resolution = context.customer_resolutions.get(row.row_number)
                if customer_resolution and customer_resolution.customer_id:
                    customer = db.get(Customer, customer_resolution.customer_id)
                elif row.customer:
                    import_summary = ImportSummary()
                    customer = _resolve_or_create_customer(db, row.customer, import_summary)
                    summary.customers_created += import_summary.customers_created

                entry = TimesheetEntry(
                    timesheet_id=timesheet.id,
                    work_category=WorkCategory.non_productive,
                    non_productive_code_id=np_code.id,
                    project_id=None,
                    customer_id=customer.id if customer else None,
                    task_type_id=None,
                    is_billable=False,
                    entry_date=row.entry_date,
                    hours=row.hours,
                    description=row.description,
                )
            else:
                project = _resolve_project_for_row(db, row, context, summary)
                if project is None:
                    summary.rows_skipped += 1
                    preview.status_label = TimesheetImportRowStatus.skipped
                    preview.messages.append("Project skipped.")
                    error_log.append(preview)
                    continue

                task_type = _resolve_task_type_for_row(db, row, project, context)
                entry = TimesheetEntry(
                    timesheet_id=timesheet.id,
                    work_category=WorkCategory.productive,
                    project_id=project.id,
                    customer_id=project.customer_id,
                    task_type_id=task_type.id,
                    is_billable=True,
                    entry_date=row.entry_date,
                    hours=row.hours,
                    description=row.description,
                )
                affected_projects.add(project.id)

            dup_key = timesheet_duplicate_key_for_row(
                user_id=designer.id,
                entry_date=row.entry_date,
                tool_number=row.tool_number,
                np_code=row.np_code,
                task_type=row.task_type,
                hours=row.hours,
                is_billable=entry.is_billable,
                description=row.description,
                is_np_row=row.is_np_row,
            )
            if check_duplicates and dup_key in seen_duplicate_keys:
                summary.duplicates_skipped += 1
                summary.rows_skipped += 1
                preview.status_label = TimesheetImportRowStatus.skipped
                preview.messages.append(
                    "Duplicate entry skipped: matching designer, date, project/NP code, "
                    "task, hours, billable, notes, month, and year."
                )
                error_log.append(preview)
                continue
            if check_duplicates:
                seen_duplicate_keys.add(dup_key)

            entries_to_add.append(entry)
            summary.rows_imported += 1
            if row.is_np_row:
                summary.np_entries += 1
            preview.status_label = TimesheetImportRowStatus.imported
            error_log.append(preview)

        if dry_run:
            db.rollback()
        else:
            db.add_all(entries_to_add)
            db.flush()
            db.commit()
            for project_id in affected_projects:
                recalculate_project(db, project_id)

            duration_ms = int((datetime.now(timezone.utc) - started).total_seconds() * 1000)
            week_starts = sorted(
                {_week_start(r.entry_date) for r in rows if r.entry_date and not r.errors}
            )
            meta_path = UPLOAD_DIR / f"{upload_id}.meta"
            file_name = meta_path.read_text(encoding="utf-8") if meta_path.exists() else path.name

            history = TimesheetImportHistory(
                filename=file_name,
                imported_by_id=imported_by_id or designer.id,
                designer_name=designer_name,
                designer_user_id=designer.id,
                week_start=week_starts[0] if week_starts else None,
                date_range_label=(
                    _week_label(week_starts[0]) if len(week_starts) == 1 else f"{len(week_starts)} weeks"
                ),
                rows_read=summary.rows_read,
                rows_imported=summary.rows_imported,
                rows_failed=summary.rows_failed,
                rows_skipped=summary.rows_skipped,
                duration_ms=duration_ms,
                status="completed",
                upload_id=upload_id,
                log_json=json.dumps(summary.model_dump(mode="json")),
            )
            db.add(history)
            db.commit()
            return summary, error_log, history.id

    except Exception:
        db.rollback()
        raise

    return summary, error_log, None


def error_log_to_csv(error_log: list[TimesheetImportRowPreview]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "row_number",
            "status",
            "designer",
            "tool_number",
            "task_type",
            "np_code",
            "hours",
            "entry_date",
            "messages",
        ]
    )
    for row in error_log:
        writer.writerow(
            [
                row.row_number,
                row.status_label.value,
                row.designer or "",
                row.tool_number or "",
                row.task_type or "",
                row.np_code or "",
                row.hours or "",
                row.entry_date or "",
                "; ".join(row.messages),
            ]
        )
    return buffer.getvalue()


def list_import_history(db: Session, limit: int = 50) -> list[TimesheetImportHistory]:
    return list(
        db.scalars(
            select(TimesheetImportHistory)
            .order_by(TimesheetImportHistory.created_at.desc())
            .limit(limit)
        ).all()
    )


def get_import_history(db: Session, history_id: uuid.UUID) -> TimesheetImportHistory | None:
    return db.get(TimesheetImportHistory, history_id)

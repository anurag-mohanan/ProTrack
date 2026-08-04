"""Excel export — one designer's timesheet entries for a date range."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.crud.foundation import get_or_create_company_settings
from app.crud.timesheet_entry_metrics import build_timesheet_entry_reads
from app.models.enums import TimesheetStatus, WorkCategory
from app.models.models import Timesheet, TimesheetEntry, User
from app.services.reporting.excel.letterhead import (
    set_print_layout,
    style_total_row,
    write_report_letterhead,
)
from app.services.reporting.excel.styles import (
    autofit_columns,
    freeze_and_filter,
    style_body_rows,
    style_header_row,
)


def _decimal(value) -> Decimal:
    return Decimal(str(value or 0))


def _category_label(row) -> str:
    work = row.work_category
    is_np = work == WorkCategory.non_productive or (
        getattr(work, "value", work) == "non_productive"
    )
    if is_np:
        if (row.leave_count or 0) > 0 or row.non_productive_category == "leave":
            return "Leave"
        return "Non-Productive"
    return "Billable" if row.is_billable else "Non-Billable"


def generate_designer_individual_timesheet_excel(
    db: Session,
    *,
    user_id: UUID,
    period_start: date,
    period_end: date,
    period_label: str | None = None,
) -> tuple[bytes, str]:
    """Return (xlsx_bytes, download_filename) for one designer's entries."""
    person = db.get(User, user_id)
    if person is None:
        raise ValueError("Designer not found")

    company = get_or_create_company_settings(db)
    company_name = company.company_name or "Prosohm"

    entries = list(
        db.scalars(
            select(TimesheetEntry)
            .join(Timesheet, TimesheetEntry.timesheet_id == Timesheet.id)
            .where(
                Timesheet.user_id == user_id,
                TimesheetEntry.is_deleted.is_(False),
                TimesheetEntry.entry_date >= period_start,
                TimesheetEntry.entry_date <= period_end,
                Timesheet.status.in_(
                    (
                        TimesheetStatus.draft,
                        TimesheetStatus.submitted,
                        TimesheetStatus.approved,
                    )
                ),
            )
            .order_by(TimesheetEntry.entry_date, TimesheetEntry.created_at)
        ).all()
    )
    reads = build_timesheet_entry_reads(db, entries)
    timesheet_ids = {entry.timesheet_id for entry in entries}
    status_by_id: dict[UUID, str] = {}
    if timesheet_ids:
        for row in db.scalars(select(Timesheet).where(Timesheet.id.in_(tuple(timesheet_ids)))).all():
            status_by_id[row.id] = (
                row.status.value if hasattr(row.status, "value") else str(row.status)
            )

    designer_name = f"{person.first_name} {person.last_name}".strip()
    label = period_label or f"{period_start.isoformat()} – {period_end.isoformat()}"
    total_hours = sum((_decimal(row.hours) for row in reads), Decimal("0"))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timesheet Entries"[:31]
    headers = [
        "Date",
        "Tool / NP",
        "Description",
        "Task",
        "Hours",
        "Billable",
        "Category",
        "Project team",
        "Notes",
        "Status",
        "Contribution reason",
    ]
    next_row = write_report_letterhead(
        sheet,
        company_name=company_name,
        report_title=f"Timesheet — {designer_name}",
        period_label=f"Period: {label}",
        extra_lines=[
            f"Email: {person.email}",
            f"Entries: {len(reads)} · Total hours: {float(total_hours):.2f}",
        ],
        col_span=len(headers),
    )
    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, len(headers))

    for offset, row in enumerate(reads):
        excel_row = header_row + 1 + offset
        work = row.work_category
        is_np = work == WorkCategory.non_productive or (
            getattr(work, "value", work) == "non_productive"
        )
        if is_np:
            tool = row.non_productive_code or "—"
            description = row.non_productive_description or "Non-Productive"
            task = row.non_productive_description or "—"
        else:
            tool = row.project_tool_number or "—"
            description = row.customer_name or row.project_code or "—"
            task = row.task_type_name or "—"
        reason = row.contribution_reason
        reason_text = (
            reason.value
            if reason is not None and hasattr(reason, "value")
            else (str(reason) if reason else "")
        )
        values = [
            row.entry_date.isoformat(),
            tool,
            description,
            task,
            float(row.hours),
            "Yes" if row.is_billable else "No",
            _category_label(row),
            row.project_team_name or "—",
            row.description or "",
            status_by_id.get(row.timesheet_id, "—"),
            reason_text,
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col, value=value)
            if col == 5:
                cell.alignment = Alignment(horizontal="right")

    if reads:
        style_body_rows(sheet, header_row + 1, header_row + len(reads), len(headers))
        total_row = header_row + len(reads) + 1
        sheet.cell(row=total_row, column=1, value="TOTAL")
        sheet.cell(row=total_row, column=5, value=float(total_hours))
        style_total_row(sheet, total_row, len(headers), emphasize_cols={1, 5})

    freeze_and_filter(sheet, header_row, len(headers))
    autofit_columns(sheet)
    set_print_layout(sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    safe_name = "".join(
        ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in designer_name
    ).strip("_") or "designer"
    filename = (
        f"timesheet-{safe_name}-{period_start.isoformat()}_to_{period_end.isoformat()}.xlsx"
    )
    return buffer.getvalue(), filename

"""IT Hardware Workbook import — preserve numbers, parents, Open, 3D mouse."""

from __future__ import annotations

from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from sqlalchemy import select

from app.models.it_operations import Asset, Computer
from app.models.models import User
from app.services import it_data_import_service as imp


HEADERS = [
    "Asset Number",
    "Category",
    "Asset Type",
    "Make / Brand",
    "Model",
    "PC Name",
    "Serial / Service Tag",
    "IP Address",
    "MAC Address",
    "Assigned User",
    "Asset Status",
    "Hardware Owner Type",
    "Hardware Owner",
    "Parent Asset Number",
    "Notes",
]


def _workbook_bytes(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "IT Assets Import"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    seed = wb.create_sheet("Master Data Seed")
    seed.append(["Category", "Asset Type", "Make", "Model"])
    seed.append(["IT Hardware", "Monitor", "Dell", "P2722H"])
    opts = wb.create_sheet("Ownership Options")
    opts.append(["Owner Type", "Notes"])
    opts.append(["Prosohm", "Company owned"])
    opts.append(["Customer", "Customer owned"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _admin(session) -> User:
    user = session.query(User).filter(User.email == "admin@prosohm.com").first()
    assert user is not None
    return user


def test_hardware_workbook_analyze_and_commit(session):
    admin = _admin(session)
    # Prefer a real first name if present in seed users
    assignee_name = f"{admin.first_name} {getattr(admin, 'last_name', '')}".strip()

    content = _workbook_bytes(
        [
            [
                "IT-155",
                "IT Hardware",
                "Workstation",
                "Dell",
                "Precision 3680",
                "Proeng103",
                "DP5LM24",
                "192.168.20.103",
                "4C-D7-17-A5-2E-A1",
                assignee_name,
                "Assigned",
                "Prosohm",
                "Prosohm",
                "",
                "Primary WS",
            ],
            [
                "IT164",  # no hyphen — must preserve exactly
                "IT Hardware",
                "Monitor",
                "Dell",
                "P2722H",
                "",
                "",
                "",
                "",
                assignee_name,
                "Assigned",
                "",
                "",
                "IT-155",
                "",
            ],
            [
                "IT-176",
                "IT Hardware",
                "3D Mouse",
                "",
                "",
                "",
                "",
                "",
                "",
                assignee_name,
                "Assigned",
                "Prosohm",
                "Prosohm",
                "IT-155",
                "",
            ],
            [
                "IT-180",
                "IT Hardware",
                "Workstation",
                "Dell",
                "Precision 5820 Tower",
                "ProengOpen",
                "ABC123",
                "192.168.20.180",
                "11-22-33-44-55-66",
                "Open",
                "Available",
                "Prosohm",
                "Prosohm",
                "",
                "Spare",
            ],
        ]
    )

    analysis = imp.analyze_upload(
        session,
        actor=admin,
        content=content,
        filename="ProTrack_Current_IT_Assets_Import.xlsx",
        import_type="hardware_workbook",
    )
    preview_rows = analysis.get("first_10_records") or []
    assert analysis["sheet_name"] == "IT Assets Import"
    assert analysis["stats"]["records_found"] == 4
    assert analysis["commit_allowed"] is True
    assert analysis["stats"].get("ownership_required", 0) >= 1

    mouse = next(r for r in preview_rows if r.get("asset_number") == "IT-176")
    assert mouse.get("make") == "3Dconnexion"
    assert mouse.get("model") == "SpaceMouse Compact"

    open_row = next(r for r in preview_rows if r.get("asset_number") == "IT-180")
    assert not any(
        i.get("code") == "USER_UNKNOWN" for i in (open_row.get("_issues") or [])
    )

    batch_id = UUID(analysis["batch_id"])
    result = imp.commit_import(
        session, actor=admin, batch_id=batch_id, confirm=True, skip_duplicates=True
    )
    session.commit()
    assert result["imported"] >= 3

    ws = session.scalar(select(Asset).where(Asset.asset_number == "IT-155"))
    assert ws is not None
    assert ws.make == "Dell"
    computer = session.scalar(select(Computer).where(Computer.asset_id == ws.id))
    assert computer is not None
    assert computer.computer_name == "Proeng103"

    exact = session.scalar(select(Asset).where(Asset.asset_number == "IT164"))
    assert exact is not None
    assert exact.asset_number == "IT164"
    assert exact.parent_asset_id == ws.id

    mouse_asset = session.scalar(select(Asset).where(Asset.asset_number == "IT-176"))
    assert mouse_asset is not None
    assert mouse_asset.make == "3Dconnexion"
    assert mouse_asset.model == "SpaceMouse Compact"
    assert mouse_asset.parent_asset_id == ws.id

    open_asset = session.scalar(select(Asset).where(Asset.asset_number == "IT-180"))
    assert open_asset is not None
    assert open_asset.status == "available"
    assert not open_asset.assignments or all(
        a.returned_date is not None for a in open_asset.assignments
    )

    # Second import must not create duplicates
    session.expire_all()
    analysis2 = imp.analyze_upload(
        session,
        actor=admin,
        content=content,
        filename="ProTrack_Current_IT_Assets_Import.xlsx",
        import_type="hardware_workbook",
    )
    assert analysis2["stats"]["duplicates"] >= 4
    result2 = imp.commit_import(
        session,
        actor=admin,
        batch_id=UUID(analysis2["batch_id"]),
        confirm=True,
        skip_duplicates=True,
    )
    session.commit()
    assert result2["imported"] == 0
    count = session.scalars(
        select(Asset).where(Asset.asset_number.in_(["IT-155", "IT164", "IT-176", "IT-180"]))
    ).all()
    assert len(count) == 4

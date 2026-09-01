"""Excel export for designer-by-team timesheet reports — ProTrack report design system."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.reporting import DesignerTeamTimesheetPayload, ReportContextMeta
from app.services.reporting.excel.styles import (
    BODY_FONT,
    HOURS_FORMAT,
    LEFT,
    RIGHT,
    apply_column_widths,
    freeze_and_filter,
    style_body_rows,
    style_header_row,
)
from app.services.reporting.excel.template import (
    KpiCard,
    MetadataCard,
    apply_utilization_conditional_format,
    format_hours,
    format_hours_display,
    format_percent_display,
    format_report_datetime,
    render_kpi_cards,
    render_metadata_cards,
    render_notes_section,
    render_report_header,
    render_section_title,
    set_print_layout,
    style_total_row,
)


def generate_designer_team_timesheet_excel(payload: DesignerTeamTimesheetPayload) -> bytes:
    workbook = Workbook()
    workbook.properties.title = f"{payload.company_name} — {payload.title}"
    workbook.properties.subject = payload.period.label
    workbook.properties.company = payload.company_name

    _write_designers_sheet(workbook, payload)
    _write_projects_sheet(workbook, payload)
    _write_cross_team_sheet(workbook, payload)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _context(payload: DesignerTeamTimesheetPayload) -> ReportContextMeta:
    if payload.context is not None:
        return payload.context
    # Backward-compatible fallback when payload was built without context.
    from app.services.reporting.excel.template import format_period_month_year

    month, year, display = format_period_month_year(
        payload.period.start_date, payload.period.end_date
    )
    return ReportContextMeta(
        period_month=month,
        period_year=year,
        period_display=display,
    )


def _generated_label(payload: DesignerTeamTimesheetPayload) -> str:
    ctx = _context(payload)
    return format_report_datetime(
        payload.generated_at, timezone_name=ctx.timezone_name
    )


def _metadata_cards(payload: DesignerTeamTimesheetPayload) -> list[MetadataCard]:
    ctx = _context(payload)
    period_details = [ctx.period_display]
    if payload.period.working_days:
        period_details.append(f"Working days: {payload.period.working_days}")

    contact_details: list[str] = []
    if ctx.customer_contact_name:
        contact_details.append(f"Contact: {ctx.customer_contact_name}")
    else:
        contact_details.append("Contact: Not Available")
    if ctx.customer_contact_email:
        contact_details.append(ctx.customer_contact_email)
    elif ctx.customer_contact_phone:
        contact_details.append(ctx.customer_contact_phone)

    team_label = "Not Available"
    team_details: list[str] = []
    if len(ctx.teams) == 1:
        team_label = ctx.teams[0].team_name
        if ctx.teams[0].department_name:
            team_details.append(ctx.teams[0].department_name)
    elif len(ctx.teams) > 1:
        team_label = f"{len(ctx.teams)} Teams"
        team_details.append(", ".join(row.team_name for row in ctx.teams[:3]))
        if len(ctx.teams) > 3:
            team_details[-1] += f" +{len(ctx.teams) - 3} more"
    elif ctx.department_summary:
        team_label = ctx.department_summary

    manager_primary = ctx.engineering_manager_summary.split("\n")[0] if ctx.engineering_manager_summary else "Not Assigned"
    manager_details: list[str] = []
    if "\n" in (ctx.engineering_manager_summary or ""):
        manager_details.append("Multiple teams — see notes")
    if ctx.department_summary and len(ctx.teams) != 1:
        manager_details.append(ctx.department_summary)

    cards = [
        MetadataCard(
            title="Reporting Period",
            primary=f"{ctx.period_month} {ctx.period_year}",
            details=period_details,
        ),
        MetadataCard(
            title="Customer",
            primary=ctx.customer_label,
            details=contact_details,
        ),
        MetadataCard(
            title="Engineering Management",
            primary=manager_primary,
            details=[*( [f"Team: {team_label}"] if team_label != "Not Available" else [] ), *team_details, *manager_details][:2],
        ),
        MetadataCard(
            title="Report Information",
            primary=payload.title,
            details=[
                f"Generated: {_generated_label(payload)}",
                f"By: {ctx.generated_by_name or 'System'}",
            ],
        ),
    ]
    return cards


def _kpi_cards(payload: DesignerTeamTimesheetPayload) -> list[KpiCard]:
    ctx = _context(payload)
    return [
        KpiCard("Designers", str(payload.designer_count)),
        KpiCard("Teams", str(payload.team_count)),
        KpiCard("Projects", str(payload.project_count)),
        KpiCard("Total Hours", format_hours_display(payload.total_designer_hours)),
        KpiCard("Productive Hrs", format_hours_display(ctx.total_productive_hours)),
        KpiCard("Non-Prod. Hrs", format_hours_display(ctx.total_non_productive_hours)),
        KpiCard("Leave Days", format_hours_display(ctx.total_leave_days)),
        KpiCard("Avg Utilization", format_percent_display(ctx.average_utilization_percent)),
    ]


def _write_branded_preamble(
    sheet,
    payload: DesignerTeamTimesheetPayload,
    *,
    report_title: str | None = None,
    col_span: int,
    include_kpis: bool = False,
    extra_kpis: list[KpiCard] | None = None,
) -> int:
    ctx = _context(payload)
    row = render_report_header(
        sheet,
        company_name=payload.company_name,
        report_title=report_title or payload.title,
        col_span=col_span,
    )
    row = render_metadata_cards(
        sheet, _metadata_cards(payload), start_row=row, col_span=col_span
    )
    if include_kpis:
        row = render_section_title(
            sheet, "Performance Summary", row=row, col_span=col_span
        )
        row = render_kpi_cards(
            sheet,
            extra_kpis or _kpi_cards(payload),
            start_row=row,
            col_span=col_span,
            per_row=4,
        )
    return row


def _write_designers_sheet(workbook: Workbook, payload: DesignerTeamTimesheetPayload) -> None:
    sheet = workbook.active
    sheet.title = "Designer Hours by Team"[:31]
    ctx = _context(payload)

    include_customer = payload.include_customer_columns
    headers = [
        "TEAM",
        "DESIGNER",
        "PRODUCTIVE HOURS",
        "NON-PRODUCTIVE HOURS",
        "LEAVE DAYS",
        "TOTAL HOURS",
        "UTILIZATION %",
        "PROJECTS",
    ]
    if include_customer:
        headers.append("CUSTOMERS")

    col_span = len(headers)
    next_row = _write_branded_preamble(
        sheet, payload, col_span=col_span, include_kpis=True
    )
    next_row = render_section_title(
        sheet, "Timesheet Summary", row=next_row, col_span=col_span
    )

    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, col_span)

    numeric_cols = {3, 4, 5, 6, 7, 8}
    if include_customer:
        numeric_cols.add(9)

    for offset, row in enumerate(payload.designers):
        excel_row = header_row + 1 + offset
        values = [
            row.team_name or "—",
            row.designer_name,
            format_hours(row.productive_hours),
            format_hours(row.non_productive_hours),
            format_hours(row.leave_days),
            format_hours(row.total_hours),
            format_hours(row.utilization_percent),
            row.project_count,
        ]
        if include_customer:
            values.append(row.customer_count)
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col, value=value)
            if col in (1, 2):
                cell.alignment = LEFT
            else:
                cell.alignment = RIGHT
                if col in (3, 4, 5, 6, 7):
                    cell.number_format = HOURS_FORMAT

    last_data_row = header_row + len(payload.designers)
    if payload.designers:
        style_body_rows(sheet, header_row + 1, last_data_row, col_span)
        apply_utilization_conditional_format(
            sheet, start_row=header_row + 1, end_row=last_data_row, col=7
        )
        total_row = last_data_row + 1
        sheet.cell(row=total_row, column=1, value="TOTAL").font = BODY_FONT
        sheet.cell(
            row=total_row, column=3, value=format_hours(ctx.total_productive_hours)
        ).number_format = HOURS_FORMAT
        sheet.cell(
            row=total_row, column=4, value=format_hours(ctx.total_non_productive_hours)
        ).number_format = HOURS_FORMAT
        sheet.cell(
            row=total_row, column=5, value=format_hours(ctx.total_leave_days)
        ).number_format = HOURS_FORMAT
        sheet.cell(
            row=total_row, column=6, value=format_hours(payload.total_designer_hours)
        ).number_format = HOURS_FORMAT
        for col in (3, 4, 5, 6):
            sheet.cell(row=total_row, column=col).alignment = RIGHT
        style_total_row(sheet, total_row, col_span, emphasize_cols={1, 6})
        filter_end = total_row
        notes_start = total_row + 2
    else:
        filter_end = header_row
        notes_start = header_row + 2

    # Multi-manager team summary when needed
    notes: list[tuple[str, str]] = [
        (
            "Productive Hours",
            "Hours recorded against productive engineering work on projects.",
        ),
        (
            "Non-productive Hours",
            "Training, meetings, internal activities, and other non-project time.",
        ),
        (
            "Utilization",
            "Productive hours as a percentage of applicable available hours for the period.",
        ),
    ]
    if ctx.teams and len([t for t in ctx.teams if t.engineering_manager]) > 1:
        team_mgr = "; ".join(
            f"{t.team_name}: {t.engineering_manager or 'Not Assigned'}" for t in ctx.teams
        )
        notes.insert(0, ("Engineering Managers", team_mgr))

    render_notes_section(sheet, notes, start_row=notes_start, col_span=col_span)

    freeze_and_filter(sheet, header_row, col_span, end_row=filter_end)
    apply_column_widths(
        sheet,
        {
            1: 24,
            2: 22,
            3: 16,
            4: 18,
            5: 12,
            6: 14,
            7: 14,
            8: 12,
            **({9: 12} if include_customer else {}),
        },
    )
    set_print_layout(
        sheet,
        landscape=True,
        header_rows=f"{header_row}:{header_row}",
        company_name=payload.company_name,
        audience=ctx.audience,
        generated_label=f"Generated: {_generated_label(payload)}",
    )


def _write_projects_sheet(workbook: Workbook, payload: DesignerTeamTimesheetPayload) -> None:
    sheet = workbook.create_sheet("Project Hours"[:31])
    ctx = _context(payload)

    include_customer = payload.include_customer_columns
    headers = ["TOOL #"]
    if include_customer:
        headers.append("CUSTOMER")
    headers.extend(
        [
            "PART DESCRIPTION",
            "QUOTED HOURS",
            "PERIOD HOURS",
            "VARIANCE HOURS",
            "VARIANCE %",
            "COMPLETION %",
            "DESIGNER",
            "STAGE",
            "STATUS",
        ]
    )
    col_span = len(headers)
    next_row = _write_branded_preamble(
        sheet,
        payload,
        report_title="Project Hours",
        col_span=col_span,
        include_kpis=True,
        extra_kpis=[
            KpiCard("Projects", str(payload.project_count)),
            KpiCard(
                "Actual Hours", format_hours_display(payload.total_project_actual_hours)
            ),
            KpiCard("Designers", str(payload.designer_count)),
            KpiCard("Teams", str(payload.team_count)),
        ],
    )
    next_row = render_section_title(
        sheet, "Project Detail", row=next_row, col_span=col_span
    )

    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, col_span)

    money_cols = {4, 5, 6, 7, 8} if include_customer else {3, 4, 5, 6, 7}
    variance_col = 6 if include_customer else 5

    for offset, row in enumerate(payload.projects):
        excel_row = header_row + 1 + offset
        values: list = [row.tool_number]
        if include_customer:
            values.append(row.customer_name)
        values.extend(
            [
                row.part_description or "",
                format_hours(row.quoted_hours),
                format_hours(row.actual_hours),
                format_hours(row.variance_hours),
                format_hours(row.variance_percent),
                format_hours(row.completion_percent),
                row.designer_name or "",
                str(
                    row.project_stage.value
                    if hasattr(row.project_stage, "value")
                    else row.project_stage
                ),
                str(
                    row.execution_status.value
                    if hasattr(row.execution_status, "value")
                    else row.execution_status
                ),
            ]
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col, value=value)
            if col in money_cols:
                cell.alignment = RIGHT
                cell.number_format = HOURS_FORMAT
            else:
                cell.alignment = LEFT
            if col == variance_col and isinstance(value, (int, float)) and value > 0:
                cell.font = Font(name="Calibri", size=10, color="B42318", bold=True)
            if col == variance_col and isinstance(value, (int, float)) and value < 0:
                cell.font = Font(name="Calibri", size=10, color="1B7F4B")

    if payload.projects:
        style_body_rows(
            sheet, header_row + 1, header_row + len(payload.projects), col_span
        )
        total_row = header_row + len(payload.projects) + 1
        sheet.cell(row=total_row, column=1, value="TOTAL")
        actual_col = 5 if include_customer else 4
        cell = sheet.cell(
            row=total_row,
            column=actual_col,
            value=format_hours(payload.total_project_actual_hours),
        )
        cell.alignment = RIGHT
        cell.number_format = HOURS_FORMAT
        style_total_row(sheet, total_row, col_span, emphasize_cols={1, actual_col})

    freeze_and_filter(sheet, header_row, col_span)
    apply_column_widths(
        sheet,
        {
            1: 14,
            2: 22,
            3: 28,
            4: 14,
            5: 16,
            6: 14,
            7: 12,
            8: 12,
            9: 18,
            10: 12,
            11: 12,
        },
    )
    set_print_layout(
        sheet,
        landscape=True,
        header_rows=f"{header_row}:{header_row}",
        company_name=payload.company_name,
        audience=ctx.audience,
        generated_label=f"Generated: {_generated_label(payload)}",
    )


def _write_cross_team_sheet(workbook: Workbook, payload: DesignerTeamTimesheetPayload) -> None:
    sheet = workbook.create_sheet("Cross-Team Hours"[:31])
    ctx = _context(payload)
    headers = [
        "DIRECTION",
        "DESIGNER",
        "HOME TEAM",
        "TOOL #",
        "PROJECT TEAM",
        "CUSTOMER",
        "HOURS",
        "CONTRIBUTION REASON",
    ]
    col_span = len(headers)
    next_row = _write_branded_preamble(
        sheet,
        payload,
        report_title="Cross-Team Hours (Extra Effort)",
        col_span=col_span,
        include_kpis=True,
        extra_kpis=[
            KpiCard(
                "Outbound Hours",
                format_hours_display(payload.cross_team_hours_outbound),
            ),
            KpiCard(
                "Inbound Hours",
                format_hours_display(payload.cross_team_hours_inbound),
            ),
            KpiCard("Rows", str(len(payload.cross_team_hours))),
            KpiCard("Teams", str(payload.team_count)),
        ],
    )
    next_row = render_section_title(
        sheet, "Cross-Team Detail", row=next_row, col_span=col_span
    )

    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, col_span)

    for offset, row in enumerate(payload.cross_team_hours):
        excel_row = header_row + 1 + offset
        values = [
            row.direction,
            row.designer_name,
            row.home_team_name or "—",
            row.tool_number,
            row.project_team_name or "—",
            row.customer_name or "—",
            format_hours(row.hours),
            row.contribution_reason or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col, value=value)
            if col == 7:
                cell.alignment = RIGHT
                cell.number_format = HOURS_FORMAT
            else:
                cell.alignment = LEFT

    if payload.cross_team_hours:
        style_body_rows(
            sheet,
            header_row + 1,
            header_row + len(payload.cross_team_hours),
            col_span,
        )

    freeze_and_filter(sheet, header_row, col_span)
    apply_column_widths(
        sheet,
        {1: 12, 2: 20, 3: 18, 4: 12, 5: 18, 6: 18, 7: 12, 8: 28},
    )
    set_print_layout(
        sheet,
        landscape=True,
        header_rows=f"{header_row}:{header_row}",
        company_name=payload.company_name,
        audience=ctx.audience,
        generated_label=f"Generated: {_generated_label(payload)}",
    )

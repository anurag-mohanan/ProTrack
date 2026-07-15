from io import BytesIO

from openpyxl import Workbook

from app.services.workorder_pdf_extract import (
    extract_workorder_fields_from_file,
    extract_workorder_fields_from_text,
)


def test_extract_workorder_label_value_pairs():
    text = """
    Customer Work Order: CMT-WO-7788
    Part Description: Frt Door Main Panel L/R
    Press Tonnage: 650 T
    Plastic Material: PP GF30
    Cavity: 2
    Tool Type: Injection mould
    Specifications: Gate from runner side
    """
    result = extract_workorder_fields_from_text(text)
    assert result.work_order_number == "CMT-WO-7788"
    assert "Door" in (result.part_description or "")
    assert result.press_tonnage == "650T"
    assert result.plastic_material == "PP GF30"
    assert result.cavity_count == 2
    assert result.tool_type and "Injection" in result.tool_type
    assert result.customer_specs and "Gate" in result.customer_specs


def test_extract_workorder_loose_patterns():
    text = "Material ABS Tonnage 400ton Cavities 1"
    result = extract_workorder_fields_from_text(text)
    assert result.plastic_material and "ABS" in result.plastic_material
    assert result.press_tonnage == "400T"
    assert result.cavity_count == 1


def test_extract_cmt_intermediate_design_review_style():
    """CMT Intermediate Approval: Part Description with 1+1 Cavity prefix + notes tonnage."""
    text = """
    Customer CMT Intermediate Approval
    Job Number 2649
    Tool Number CMT-2649
    Part Description 1+1 Cavity Frt Door Main Panel L/R
    General Notes (Intermediate)
    2200T Press 308
    Locating Ring 5.990 "
    """
    result = extract_workorder_fields_from_text(text)
    assert result.part_description == "Frt Door Main Panel L/R"
    assert result.cavity_count == 2
    assert result.press_tonnage == "2200T"


def test_extract_abc_shop_order_style():
    """ABC Shop Order: PART NAME, # Cav., resin line."""
    text = """
    SHOP ORDER
    S.O. # : 8348
    PART NAME: BO-05 Cover Assy
    # Cav.: 4
    1a) MAT'L/FAMILY/GRADE/COLOUR 1:R RESIN - S375AHW-600R BLACK
    PRIMARY PRESSES / NUMBERS: Press 1800 / Bay 12
    """
    result = extract_workorder_fields_from_text(text)
    assert result.part_description and "Cover" in result.part_description
    assert result.cavity_count == 4
    assert result.plastic_material and "S375AHW" in result.plastic_material


def test_extract_bb_kickoff_excel_label_rows():
    """B&B Outsource Kick Off: label | value rows."""
    workbook = Workbook()
    sheet = workbook.active
    rows = [
        ("B&B Job Number", "3175"),
        ("Part Name", "Bracket RH Outer"),
        ("Tool Type", "Conventional"),
        ("Cavitation", "One"),
        ("Plastic Type (Main)", "PA66 GF30 Black"),
        ("Press Tonnage Primary", "500T"),
        ("Customer Tool Number", "BB-TOOL-88"),
        ("Shrink", "0.6%"),
    ]
    for idx, (label, value) in enumerate(rows, start=1):
        sheet.cell(idx, 1, label)
        sheet.cell(idx, 2, value)

    buffer = BytesIO()
    workbook.save(buffer)
    result = extract_workorder_fields_from_file(buffer.getvalue(), filename="3175 Outsource Kick Off.xlsx")
    assert result.part_description == "Bracket RH Outer"
    assert result.tool_type == "Conventional"
    assert result.cavity_count == 1
    assert result.plastic_material and "PA66" in result.plastic_material
    assert result.press_tonnage == "500T"
    # WO optional — may be filled from B&B Job Number but not required for pass
    assert not result.warnings or "No tooling fields" not in " ".join(result.warnings)


def test_extract_press_tonnage_from_general_notes_line():
    """Real customer WOs put tonnage in notes, e.g. '2200T Press 308'."""
    text = """
    General Notes (Intermediate)
    2200T Press 308
    Locating Ring 5.990 "
    76V X 62H X 34.75 SH
    Change direct press bolts to Slots . press bolts are 1.25 " and use 1.5" slots
    """
    result = extract_workorder_fields_from_text(text)
    assert result.press_tonnage == "2200T"


def test_extract_press_tonnage_variants():
    assert extract_workorder_fields_from_text("Press 1800T required").press_tonnage == "1800T"
    assert extract_workorder_fields_from_text("Press Tonnage: 650 T").press_tonnage == "650T"
    assert extract_workorder_fields_from_text("650 T Press for tool").press_tonnage == "650T"


def test_stdlib_pdf_scrape_reads_literals():
    from app.services.workorder_pdf_extract import _extract_text_stdlib, extract_workorder_fields_from_pdf

    content = b"""%PDF-1.4
BT (Work Order: WO-7788) Tj (Press Tonnage: 650 T) Tj (Plastic Material: PP GF30) Tj (Cavity: 2) Tj ET
%%EOF
"""
    text = _extract_text_stdlib(content)
    assert "WO-7788" in text or "Work Order" in text

    result = extract_workorder_fields_from_pdf(content)
    assert result.work_order_number == "WO-7788"
    assert result.press_tonnage == "650T"
    assert result.plastic_material == "PP GF30"
    assert result.cavity_count == 2

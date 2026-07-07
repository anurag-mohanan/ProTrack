"""Bulk historical timesheet import from a folder of Prosohm Excel workbooks."""

from __future__ import annotations

import json
import os
import re
import shutil
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path, PurePosixPath
from tempfile import gettempdir
from typing import Callable

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.enums import WorkCategory
from app.models.models import (
    Project,
    TimesheetEntry,
    User,
)
from app.schemas.historical_timesheet_folder_import import (
    FolderDesignerScanRow,
    FolderImportLogRow,
    FolderImportSummary,
    FolderScanResponse,
)
from app.services.historical_import_service import (
    _get_or_create_approved_timesheet,
    _normalize_header,
    _resolve_np_code,
    _week_start,
)
from app.services.historical_timesheet_import_service import (
    _cell_text,
    _find_project,
    _find_task_type,
    _match_user_by_name,
    _normalize_np_code,
    _parse_decimal,
    _parse_date,
    _resolve_task_type_name,
    designer_has_entries_for_month,
    load_timesheet_duplicate_keys,
    timesheet_duplicate_key,
)
from app.services.project_calculation_service import recalculate_project
from app.services.non_productive_entry_service import build_np_timesheet_entry

FOLDER_BATCH_DIR = Path(gettempdir()) / "protrack_timesheet_folder_imports"
BATCH_COMMIT_SIZE = 500

PROSOHM_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "sl_no": ("sl no", "sl no.", "s.no", "s no", "#", "sl"),
    "entry_date": ("date",),
    "project": ("project", "project number", "tool no", "tool number", "tool no."),
    "customer": ("customer", "customer name"),
    "task": ("task", "task type", "activity"),
    "billable": ("billable",),
    "hours": ("hours", "hrs", "time"),
    "notes": ("notes", "description", "comments"),
}


@dataclass
class ProsohmRow:
    row_number: int
    entry_date: date | None = None
    project_number: str | None = None
    task: str | None = None
    billable: bool | None = None
    hours: Decimal | None = None
    notes: str | None = None
    is_np_row: bool = False
    np_code: str | None = None


@dataclass
class ProsohmWorkbook:
    file_path: Path
    relative_path: str
    folder_designer: str
    designer: str
    month_label: str | None
    workbook_year: int | None = None
    workbook_month: int | None = None
    rows: list[ProsohmRow] = field(default_factory=list)


@dataclass
class ResolvedFile:
    workbook: ProsohmWorkbook
    designer_user: User | None


def _is_temp_excel(name: str) -> bool:
    return name.startswith("~$")


def _is_completely_blank(values: list[object]) -> bool:
    return not any(v is not None and str(v).strip() for v in values)


def _parse_month_label(value: object) -> tuple[int, int] | None:
    if value is None:
        return None
    text = str(value).strip()
    match = re.match(r"^([A-Za-z]{3})-(\d{2})$", text)
    if not match:
        return None
    try:
        month = datetime.strptime(match.group(1).title(), "%b").month
    except ValueError:
        return None
    year = 2000 + int(match.group(2))
    return year, month


def _parse_entry_date(value: object, month_hint: tuple[int, int] | None) -> date | None:
    parsed = _parse_date(value)
    if parsed is not None:
        return parsed
    if month_hint is None or value is None:
        return None
    text = str(value).strip()
    if not text.isdigit():
        return None
    day = int(text)
    year, month = month_hint
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _detect_prosohm_header(headers: list[str]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    normalized = [_normalize_header(h) for h in headers]
    for field_name, aliases in PROSOHM_HEADER_ALIASES.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if header in aliases or any(alias in header for alias in aliases):
                header_map[field_name] = idx
                break
    required = {"entry_date", "project", "task", "hours"}
    if not required.issubset(header_map):
        raise ValueError(
            "Could not locate Prosohm header row. Expected Date, Project, Task, and Hours."
        )
    return header_map


def _folder_designer_from_path(relative_path: str) -> str:
    parts = PurePosixPath(relative_path.replace("\\", "/")).parts
    if len(parts) >= 2:
        return parts[-2]
    return "Unknown"


def parse_prosohm_workbook(file_path: Path, *, relative_path: str | None = None) -> ProsohmWorkbook:
    rel = relative_path or file_path.name
    folder_designer = _folder_designer_from_path(rel)
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Workbook has no active worksheet.")

        designer = _cell_text(sheet.cell(1, 2).value) or folder_designer
        month_label = _cell_text(sheet.cell(2, 2).value)
        month_hint = _parse_month_label(month_label) if month_label else None

        header_map: dict[str, int] | None = None
        header_row = 1
        max_col = max(sheet.max_column or 1, 8)
        for row_idx in range(1, min(sheet.max_row or 1, 40) + 1):
            headers = [sheet.cell(row_idx, col).value for col in range(1, max_col + 1)]
            try:
                header_map = _detect_prosohm_header(
                    [str(h) if h is not None else "" for h in headers]
                )
                header_row = row_idx
                break
            except ValueError:
                continue
        if header_map is None:
            raise ValueError("Could not locate Prosohm header row.")

        def read_col(values: list[object], field: str) -> object | None:
            idx = header_map.get(field)  # type: ignore[union-attr]
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        parsed_rows: list[ProsohmRow] = []
        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            values = [sheet.cell(row_idx, col).value for col in range(1, max_col + 1)]
            if _is_completely_blank(values):
                break

            project_number = _cell_text(read_col(values, "project"))
            task = _cell_text(read_col(values, "task"))
            hours = _parse_decimal(read_col(values, "hours"))
            entry_date = _parse_entry_date(read_col(values, "entry_date"), month_hint)
            billable_raw = read_col(values, "billable")
            billable = None
            if billable_raw is not None:
                text = str(billable_raw).strip().lower()
                if text in {"yes", "y", "true", "1"}:
                    billable = True
                elif text in {"no", "n", "false", "0"}:
                    billable = False
            notes = _cell_text(read_col(values, "notes"))

            np_code = _normalize_np_code(project_number)
            is_np = np_code is not None

            parsed_rows.append(
                ProsohmRow(
                    row_number=row_idx,
                    entry_date=entry_date,
                    project_number=project_number,
                    task=task,
                    billable=billable,
                    hours=hours,
                    notes=notes,
                    is_np_row=is_np,
                    np_code=np_code,
                )
            )

        return ProsohmWorkbook(
            file_path=file_path,
            relative_path=rel,
            folder_designer=folder_designer,
            designer=designer,
            month_label=month_label,
            workbook_year=month_hint[0] if month_hint else None,
            workbook_month=month_hint[1] if month_hint else None,
            rows=parsed_rows,
        )
    finally:
        workbook.close()


def collect_xlsx_files(root: Path) -> list[Path]:
    if not root.is_dir():
        raise ValueError(f"Source folder does not exist: {root}")
    files: list[Path] = []
    for path in sorted(root.rglob("*.xlsx")):
        if _is_temp_excel(path.name):
            continue
        files.append(path)
    return files


def save_folder_batch(batch_id: str, files: list[tuple[str, bytes]]) -> Path:
    batch_dir = FOLDER_BATCH_DIR / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)
    manifest: list[str] = []
    for rel_path, content in files:
        safe_rel = rel_path.replace("\\", "/").lstrip("/")
        dest = batch_dir / safe_rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(content)
        manifest.append(safe_rel)
    (batch_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return batch_dir


def get_batch_dir(batch_id: str) -> Path:
    batch_dir = FOLDER_BATCH_DIR / batch_id
    if not batch_dir.is_dir():
        raise FileNotFoundError("Upload batch not found or expired.")
    return batch_dir


def list_batch_files(batch_id: str) -> list[tuple[Path, str]]:
    batch_dir = get_batch_dir(batch_id)
    manifest_path = batch_dir / "manifest.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        return [(batch_dir / rel, rel) for rel in manifest]
    files: list[tuple[Path, str]] = []
    for path in collect_xlsx_files(batch_dir):
        rel = str(path.relative_to(batch_dir)).replace("\\", "/")
        files.append((path, rel))
    return files


def resolve_source_files(
    *,
    batch_id: str | None = None,
    source_path: str | None = None,
) -> tuple[list[tuple[Path, str]], str]:
    if batch_id:
        files = list_batch_files(batch_id)
        label = f"Uploaded batch ({len(files)} files)"
        return files, label
    if source_path:
        root = Path(source_path).expanduser().resolve()
        paths = collect_xlsx_files(root)
        files = [(path, str(path.relative_to(root)).replace("\\", "/")) for path in paths]
        return files, str(root)
    raise ValueError("Either batch_id or source_path is required.")


def scan_folder_source(
    *,
    batch_id: str | None = None,
    source_path: str | None = None,
) -> FolderScanResponse:
    files, label = resolve_source_files(batch_id=batch_id, source_path=source_path)
    by_designer: dict[str, dict[str, int]] = defaultdict(lambda: {"files": 0, "entries": 0})

    for path, rel in files:
        try:
            workbook = parse_prosohm_workbook(path, relative_path=rel)
            designer = workbook.designer or workbook.folder_designer
            by_designer[designer]["files"] += 1
            by_designer[designer]["entries"] += len(workbook.rows)
        except Exception:
            folder_designer = _folder_designer_from_path(rel)
            by_designer[folder_designer]["files"] += 1

    designers = [
        FolderDesignerScanRow(
            designer=name,
            files=stats["files"],
            entries=stats["entries"],
        )
        for name, stats in sorted(by_designer.items(), key=lambda item: item[0].lower())
    ]
    file_count = len(files)
    estimated_entries = sum(row.entries for row in designers)
    return FolderScanResponse(
        batch_id=batch_id,
        source_label=label,
        designer_count=len(designers),
        file_count=file_count,
        estimated_entries=estimated_entries,
        designers=designers,
    )


def create_pre_import_backup() -> Path | None:
    database_url = os.getenv("DATABASE_URL", "sqlite:///./protrack.db")
    if not database_url.startswith("sqlite"):
        return None
    db_path = Path(database_url.removeprefix("sqlite:///"))
    if not db_path.is_file():
        return None
    backup_root = Path("Backups") / "Before Historical Import"
    backup_root.mkdir(parents=True, exist_ok=True)
    dest = backup_root / f"{datetime.now().strftime('%Y-%m-%d_%H%M')}.db"
    shutil.copy2(db_path, dest)
    return dest


DUPLICATE_SKIP_REASON = (
    "Duplicate Entry: matching designer, date, project/NP code, task, hours, "
    "billable, notes, month, and year"
)


def _log_row(
    file_name: str,
    *,
    row_number: int | None = None,
    designer: str | None = None,
    entry_date: date | None = None,
    project: str | None = None,
    reason: str,
) -> FolderImportLogRow:
    return FolderImportLogRow(
        file=file_name,
        row=row_number,
        designer=designer,
        entry_date=entry_date.isoformat() if entry_date else None,
        project=project,
        reason=reason,
    )


def import_log_to_excel(log_rows: list[FolderImportLogRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import Log"
    sheet.append(["File", "Row", "Designer", "Date", "Project", "Reason"])
    for row in log_rows:
        sheet.append(
            [
                row.file,
                row.row or "",
                row.designer or "",
                row.entry_date or "",
                row.project or "",
                row.reason,
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def run_folder_import(
    db: Session,
    *,
    batch_id: str | None = None,
    source_path: str | None = None,
    imported_by_id: uuid.UUID | None = None,
    progress_callback: Callable[..., None] | None = None,
    cancel_check: Callable[[], bool] | None = None,
    backup_path: Path | None = None,
    after_database_reset: bool = False,
    ignore_duplicate_check: bool = True,
) -> tuple[FolderImportSummary, list[FolderImportLogRow]]:
    started = datetime.now(timezone.utc)
    files, _ = resolve_source_files(batch_id=batch_id, source_path=source_path)
    skip_duplicate_check = ignore_duplicate_check
    summary = FolderImportSummary(
        backup_path=str(backup_path) if backup_path else None,
        database_reset_performed=after_database_reset,
        duplicate_check_disabled=skip_duplicate_check,
    )
    log_rows: list[FolderImportLogRow] = []

    workbooks: list[ProsohmWorkbook] = []
    for path, rel in files:
        try:
            workbooks.append(parse_prosohm_workbook(path, relative_path=rel))
        except Exception as exc:
            summary.errors += 1
            log_rows.append(
                _log_row(path.name, designer=_folder_designer_from_path(rel), reason=str(exc))
            )

    rows_total = sum(len(wb.rows) for wb in workbooks)
    rows_processed = 0
    pending_entries: list[TimesheetEntry] = []
    affected_projects: set[uuid.UUID] = set()
    designers_seen: set[str] = set()
    files_imported = 0

    def flush_batch() -> None:
        nonlocal pending_entries
        if not pending_entries:
            return
        db.add_all(pending_entries)
        db.commit()
        pending_entries = []

    try:
        for file_index, workbook in enumerate(workbooks, start=1):
            if cancel_check and cancel_check():
                flush_batch()
                break

            designer_name = workbook.designer or workbook.folder_designer
            designer_user = _match_user_by_name(db, designer_name)
            if designer_user is None:
                for row in workbook.rows:
                    summary.errors += 1
                    rows_processed += 1
                    log_rows.append(
                        _log_row(
                            workbook.file_path.name,
                            row_number=row.row_number,
                            designer=designer_name,
                            entry_date=row.entry_date,
                            project=row.project_number,
                            reason="Unknown Designer",
                        )
                    )
                if progress_callback:
                    progress_callback(
                        percent_complete=int((rows_processed / rows_total) * 100) if rows_total else 0,
                        current_designer=designer_name,
                        current_file=workbook.file_path.name,
                        rows_imported=summary.rows_imported,
                        rows_total=rows_total,
                        files_processed=file_index,
                        files_total=len(workbooks),
                    )
                continue

            designers_seen.add(designer_name)
            file_had_import = False

            check_duplicates = False
            workbook_dup_keys: set[tuple] = set()
            if (
                not skip_duplicate_check
                and workbook.workbook_year is not None
                and workbook.workbook_month is not None
                and designer_has_entries_for_month(
                    db,
                    designer_user.id,
                    workbook.workbook_year,
                    workbook.workbook_month,
                )
            ):
                check_duplicates = True
                workbook_dup_keys = load_timesheet_duplicate_keys(
                    db,
                    designer_user.id,
                    year=workbook.workbook_year,
                    month=workbook.workbook_month,
                )

            for row in workbook.rows:
                if cancel_check and cancel_check():
                    break

                summary.rows_read += 1
                rows_processed += 1
                file_name = workbook.file_path.name

                if row.entry_date is None:
                    summary.errors += 1
                    log_rows.append(
                        _log_row(
                            file_name,
                            row_number=row.row_number,
                            designer=designer_name,
                            project=row.project_number,
                            reason="Invalid Date",
                        )
                    )
                    continue

                if row.hours is None or row.hours <= 0:
                    summary.errors += 1
                    log_rows.append(
                        _log_row(
                            file_name,
                            row_number=row.row_number,
                            designer=designer_name,
                            entry_date=row.entry_date,
                            project=row.project_number,
                            reason="Invalid Hours",
                        )
                    )
                    continue

                week = _week_start(row.entry_date)
                timesheet = _get_or_create_approved_timesheet(db, designer_user.id, week)

                if row.is_np_row:
                    if row.np_code is None:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason="Unknown Project",
                            )
                        )
                        continue
                    np_record = _resolve_np_code(db, row.np_code)
                    entry = build_np_timesheet_entry(
                        np_code=np_record,
                        timesheet_id=timesheet.id,
                        entry_date=row.entry_date,
                        hours=row.hours,
                        description=row.notes,
                    )
                else:
                    if not row.project_number:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                reason="Unknown Project",
                            )
                        )
                        continue
                    if not row.task:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason="Unknown Task",
                            )
                        )
                        continue

                    project = _find_project(db, row.project_number, None)
                    if project is None:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason="Unknown Project",
                            )
                        )
                        continue

                    if _resolve_task_type_name(row.task) is None:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason="Unknown Task",
                            )
                        )
                        continue

                    task_type = _find_task_type(db, row.task, project.stream_id)
                    if task_type is None:
                        summary.errors += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason="Unknown Task",
                            )
                        )
                        continue

                    is_billable = row.billable if row.billable is not None else True
                    entry = TimesheetEntry(
                        timesheet_id=timesheet.id,
                        work_category=WorkCategory.productive,
                        project_id=project.id,
                        customer_id=project.customer_id,
                        task_type_id=task_type.id,
                        is_billable=is_billable,
                        entry_date=row.entry_date,
                        hours=row.hours,
                        description=row.notes,
                    )
                    affected_projects.add(project.id)

                if check_duplicates:
                    dup_key = timesheet_duplicate_key(
                        user_id=designer_user.id,
                        entry_date=row.entry_date,
                        project_number=None if row.is_np_row else row.project_number,
                        np_code=row.np_code if row.is_np_row else None,
                        task=None if row.is_np_row else row.task,
                        hours=row.hours,
                        is_billable=entry.is_billable,
                        notes=row.notes,
                    )
                    if dup_key in workbook_dup_keys:
                        summary.duplicates_skipped += 1
                        log_rows.append(
                            _log_row(
                                file_name,
                                row_number=row.row_number,
                                designer=designer_name,
                                entry_date=row.entry_date,
                                project=row.project_number,
                                reason=DUPLICATE_SKIP_REASON,
                            )
                        )
                        continue
                    workbook_dup_keys.add(dup_key)

                pending_entries.append(entry)
                summary.rows_imported += 1
                file_had_import = True

                if len(pending_entries) >= BATCH_COMMIT_SIZE:
                    flush_batch()

                if progress_callback and rows_processed % 25 == 0:
                    progress_callback(
                        percent_complete=int((rows_processed / rows_total) * 100) if rows_total else 0,
                        current_designer=designer_name,
                        current_file=file_name,
                        rows_imported=summary.rows_imported,
                        rows_total=rows_total,
                        files_processed=file_index,
                        files_total=len(workbooks),
                    )

            if file_had_import:
                files_imported += 1

            if progress_callback:
                progress_callback(
                    percent_complete=int((rows_processed / rows_total) * 100) if rows_total else 100,
                    current_designer=designer_name,
                    current_file=workbook.file_path.name,
                    rows_imported=summary.rows_imported,
                    rows_total=rows_total,
                    files_processed=file_index,
                    files_total=len(workbooks),
                )

        flush_batch()
        for project_id in affected_projects:
            recalculate_project(db, project_id)

        summary.designers_imported = len(designers_seen)
        summary.files_imported = files_imported
        summary.duration_seconds = int((datetime.now(timezone.utc) - started).total_seconds())
        return summary, log_rows
    except Exception:
        db.rollback()
        raise

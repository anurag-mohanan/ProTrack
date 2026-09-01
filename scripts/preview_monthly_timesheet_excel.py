"""Generate a sample Monthly Timesheet Excel and print a layout audit."""

from __future__ import annotations

import sys
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from app.models.enums import ExecutionStatus, ProjectStage
from app.schemas.reporting import (
    DesignerProductivityRow,
    DesignerTeamTimesheetPayload,
    ReportContextMeta,
    ReportPeriod,
    ReportTeamManagerRow,
    ToolHoursRow,
)
from app.services.reporting.excel.designer_team_timesheet import (
    generate_designer_team_timesheet_excel,
)
from datetime import datetime, timezone
from uuid import uuid4


def main() -> None:
    period = ReportPeriod(
        period_type="monthly",
        label="Monthly — 01 Aug 2026 to 31 Aug 2026",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
        working_days=21,
    )
    payload = DesignerTeamTimesheetPayload(
        report_id="monthly-timesheet",
        title="Monthly Timesheet Report",
        company_name="Prosohm Projects Pvt Ltd",
        period=period,
        generated_at=datetime(2026, 8, 31, 3, 32, tzinfo=timezone.utc),
        designers=[
            DesignerProductivityRow(
                user_id=uuid4(),
                designer_name="Logesh Perumal",
                team_name="Eng 3 - Redoe",
                productive_hours=Decimal("84.0"),
                non_productive_hours=Decimal("0"),
                leave_days=Decimal("0"),
                total_hours=Decimal("84.0"),
                billable_percent=Decimal("100"),
                utilization_percent=Decimal("50.0"),
                project_count=3,
                customer_count=2,
            ),
            DesignerProductivityRow(
                user_id=uuid4(),
                designer_name="Sarath Babu K.",
                team_name="Eng 3 - Redoe",
                productive_hours=Decimal("99.5"),
                non_productive_hours=Decimal("0"),
                leave_days=Decimal("0"),
                total_hours=Decimal("99.5"),
                billable_percent=Decimal("100"),
                utilization_percent=Decimal("59.2"),
                project_count=4,
                customer_count=2,
            ),
        ],
        projects=[
            ToolHoursRow(
                project_id=uuid4(),
                tool_number="T-1001",
                customer_name="Example Automotive",
                part_description="Door Panel RH",
                quoted_hours=Decimal("120"),
                actual_hours=Decimal("95.5"),
                variance_hours=Decimal("-24.5"),
                variance_percent=Decimal("-20.4"),
                completion_percent=Decimal("80"),
                designer_name="Logesh Perumal",
                project_stage=ProjectStage.intermediate,
                execution_status=ExecutionStatus.currently_being_worked_on,
            )
        ],
        total_designer_hours=Decimal("183.5"),
        total_project_actual_hours=Decimal("95.5"),
        team_count=1,
        designer_count=2,
        project_count=1,
        include_customer_columns=True,
        context=ReportContextMeta(
            period_month="AUGUST",
            period_year=2026,
            period_display="01 Aug 2026 - 31 Aug 2026",
            customer_label="Multiple Customers",
            customer_contact_name=None,
            teams=[
                ReportTeamManagerRow(
                    team_name="Eng 3 - Redoe",
                    engineering_manager="Jane Doe",
                    department_name="Mold Engineering",
                )
            ],
            engineering_manager_summary="Jane Doe",
            department_summary="Mold Engineering",
            generated_by_name="System Admin",
            timezone_name="Asia/Kolkata",
            audience="internal",
            total_productive_hours=Decimal("183.5"),
            total_non_productive_hours=Decimal("0"),
            total_leave_days=Decimal("0"),
            average_utilization_percent=Decimal("54.6"),
        ),
    )

    content = generate_designer_team_timesheet_excel(payload)
    out = ROOT / "uploads" / "report-preview-monthly-timesheet.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(content)
    print(f"Wrote {out} ({len(content)} bytes)")

    wb = load_workbook(BytesIO(content))
    sheet = wb["Designer Hours by Team"]
    print(f"Images: {len(sheet._images)}")
    print(f"Freeze panes: {sheet.freeze_panes}")
    print(f"Auto filter: {sheet.auto_filter.ref}")
    print(f"Orientation: {sheet.page_setup.orientation}")
    print(f"FitToWidth: {sheet.page_setup.fitToWidth}")
    print(f"Print title rows: {sheet.print_title_rows}")
    print(f"Footer L: {sheet.oddFooter.left.text}")
    print(f"Footer C: {sheet.oddFooter.center.text}")
    print(f"Footer R: {sheet.oddFooter.right.text}")
    print("--- Layout dump (first 35 rows) ---")
    for row in range(1, 36):
        vals = []
        for col in range(1, 10):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                vals.append(f"{cell.coordinate}={cell.value!r}")
        if vals:
            print(f"R{row}: " + " | ".join(vals))
        elif sheet.row_dimensions[row].height:
            print(f"R{row}: (empty, height={sheet.row_dimensions[row].height})")


if __name__ == "__main__":
    main()

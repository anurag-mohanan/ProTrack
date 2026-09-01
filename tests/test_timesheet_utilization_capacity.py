"""Tests for timesheet utilization capacity and report section selection."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.services.reporting.timesheet_report_sections import (
    DEFAULT_TIMESHEET_SECTIONS,
    parse_timesheet_sections,
    TIMESHEET_SECTION_DESIGNERS,
    TIMESHEET_SECTION_PROJECTS,
)
from app.services.reporting.utilization_capacity import (
    calculate_utilization_percent,
    count_applicable_working_days,
    designer_available_hours,
)
from app.services.reporting.excel.designer_team_timesheet import generate_designer_team_timesheet_excel
from app.schemas.reporting import (
    DesignerProductivityRow,
    DesignerTeamTimesheetPayload,
    ReportContextMeta,
    ReportPeriod,
)


def test_count_applicable_working_days_mid_month_start():
    """Designer starting 17 Aug should only count working days from 17 Aug."""
    holidays: set[date] = set()
    intervals = [(date(2026, 8, 17), date(2026, 8, 31))]
    start, end, days = count_applicable_working_days(
        intervals,
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        holidays=holidays,
    )
    assert start == date(2026, 8, 17)
    assert end == date(2026, 8, 31)
    # Aug 2026: 17-31 has 11 Mon-Fri days
    assert days == 11


def test_designer_available_hours_and_utilization():
    holidays: set[date] = set()
    intervals = [(date(2026, 8, 17), date(2026, 8, 31))]
    start, end, working_days, available = designer_available_hours(
        intervals,
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        holidays=holidays,
        daily_hours=Decimal("8"),
    )
    assert working_days == 11
    assert available == Decimal("88")
    util = calculate_utilization_percent(Decimal("70"), available)
    assert util == Decimal("79.5")


def test_utilization_zero_available_returns_none():
    assert calculate_utilization_percent(Decimal("10"), Decimal("0")) is None


def test_parse_timesheet_sections_defaults_to_all():
    assert parse_timesheet_sections(None) == list(DEFAULT_TIMESHEET_SECTIONS)


def test_parse_timesheet_sections_subset():
    selected = parse_timesheet_sections("designers,utilization")
    assert selected == [TIMESHEET_SECTION_DESIGNERS, "utilization"]


def test_parse_timesheet_sections_empty_tokens_rejected():
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        parse_timesheet_sections(",")
    assert exc.value.status_code == 400


def _sample_payload(**kwargs) -> DesignerTeamTimesheetPayload:
    period = ReportPeriod(
        period_type="monthly",
        label="Monthly — 01 Aug 2026 to 31 Aug 2026",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 31),
        working_days=21,
    )
    designer = DesignerProductivityRow(
        user_id=kwargs.get("user_id", uuid4()),
        designer_name="Test Designer",
        team_name="Eng 3",
        productive_hours=Decimal("70"),
        non_productive_hours=Decimal("8"),
        leave_days=Decimal("0"),
        total_hours=Decimal("78"),
        billable_percent=Decimal("90"),
        utilization_percent=Decimal("79.5"),
        project_count=2,
        customer_count=1,
        applicable_start_date=date(2026, 8, 17),
        applicable_end_date=date(2026, 8, 31),
        applicable_working_days=11,
        available_hours=Decimal("88"),
    )
    defaults = {
        "report_id": "monthly-timesheet",
        "title": "Monthly Timesheet Report",
        "company_name": "Prosohm Projects Pvt Ltd",
        "period": period,
        "generated_at": datetime.now(timezone.utc),
        "designers": [designer],
        "projects": [],
        "total_designer_hours": Decimal("78"),
        "total_project_actual_hours": Decimal("0"),
        "team_count": 1,
        "designer_count": 1,
        "project_count": 0,
        "context": ReportContextMeta(
            period_month="AUGUST",
            period_year=2026,
            period_display="01 Aug 2026 - 31 Aug 2026",
            customer_label="All Customers",
            engineering_manager_summary="Manager",
            total_productive_hours=Decimal("70"),
            total_non_productive_hours=Decimal("8"),
            total_leave_days=Decimal("0"),
            average_utilization_percent=Decimal("79.5"),
        ),
        "selected_sections": list(DEFAULT_TIMESHEET_SECTIONS),
    }
    defaults.update(kwargs)
    return DesignerTeamTimesheetPayload(**defaults)


def test_excel_generates_only_selected_sections(tmp_path, monkeypatch):
    from app.services.reporting.excel import template as template_module

    logo = tmp_path / "company-logo.png"
    logo.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    monkeypatch.setattr(template_module, "COMPANY_LOGO_DIR", tmp_path)
    monkeypatch.setattr(
        template_module,
        "resolve_company_logo_path",
        lambda logo_url=None: logo,
    )

    payload = _sample_payload(
        selected_sections=[TIMESHEET_SECTION_DESIGNERS, TIMESHEET_SECTION_PROJECTS],
        company_logo_path=str(logo),
    )
    content = generate_designer_team_timesheet_excel(payload)
    workbook = load_workbook(BytesIO(content))
    titles = workbook.sheetnames
    assert "Designer Hours by Team" in titles
    assert "Project Hours" in titles
    assert "Cross-Team Hours" not in titles
    assert "Utilization Summary" not in titles
    assert len(workbook["Designer Hours by Team"]._images) == 1

"""IT migration analyze / import — secrets excluded, duplicates skipped."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select

from app.models.it_operations import Asset, ITUserAccount
from app.models.models import User


def _hardware_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "1.HARDWARES"
    ws.append(
        [
            "SL NO",
            "ID",
            "DESCRIPTION",
            "CUSTOMER USED FOR",
            "TYPE",
            "CURRENT USER",
            "MAKE",
            "MODEL",
            "CURRENT STATUS",
            "SERVICE TAG#",
            "PWD",
            "Teamviewer ID/pwd",
        ]
    )
    ws.append(
        [
            "1",
            "PP-MIG-01",
            "Test workstation",
            "Generic",
            "Workstation",
            "All",
            "DELL",
            "3680",
            "WORKING",
            "MIGTAG001",
            "should-never-appear",
            "tv-secret",
        ]
    )
    ws.append(
        [
            "2",
            "IT-136",
            "Unclear ownership",
            "Sybridge",
            "Workstation",
            "Nobody",
            "DELL",
            "5820",
            "WORKING",
            "MIGTAG002",
            "secret2",
            "",
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _accounts_workbook_bytes(email: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "2.USER_CREDENTIALS"
    ws.append(
        [
            "SL NO.",
            "EMPLOYEE NAME",
            "DEPARTMENT",
            "STATUS",
            "USERNAME",
            "PWD",
            "EMAIL",
            "PWD3",
        ]
    )
    ws.append(
        [
            "1",
            "Migration User",
            "Admin",
            "Current",
            "miguser",
            "plaintext-password",
            email,
            "another-secret",
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_migration_analyze_excludes_secrets_and_requires_confirm(client):
    content = _hardware_workbook_bytes()
    response = client.post(
        "/api/v1/it/migration/analyze",
        headers=client.auth_headers,
        data={"source_type": "hardware"},
        files={
            "file": (
                "hardware.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["records_found"] == 2
    assert body["confirm_required"] is True
    assert "PWD" in body["sensitive_columns_excluded"]
    assert any("Teamviewer" in c for c in body["sensitive_columns_excluded"])
    preview_blob = str(body["preview_rows"]).lower()
    assert "should-never-appear" not in preview_blob
    assert "tv-secret" not in preview_blob
    assert body["session_id"]

    # Import without confirm must fail
    denied = client.post(
        "/api/v1/it/migration/import",
        headers=client.auth_headers,
        data={"session_id": body["session_id"], "confirm": "false"},
    )
    assert denied.status_code == 422


def test_migration_import_hardware_and_skip_duplicate(client, session):
    content = _hardware_workbook_bytes()
    analyzed = client.post(
        "/api/v1/it/migration/analyze",
        headers=client.auth_headers,
        data={"source_type": "hardware"},
        files={
            "file": (
                "hardware.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert analyzed.status_code == 200, analyzed.text
    session_id = analyzed.json()["session_id"]

    imported = client.post(
        "/api/v1/it/migration/import",
        headers=client.auth_headers,
        data={
            "session_id": session_id,
            "confirm": "true",
            "skip_duplicates": "true",
        },
    )
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert result["imported"] >= 1

    assets = list(session.scalars(select(Asset).where(Asset.is_deleted.is_(False))).all())
    legacy_ids = {a.legacy_asset_number for a in assets}
    assert "PP-MIG-01" in legacy_ids
    by_legacy = {a.legacy_asset_number: a for a in assets if a.legacy_asset_number}
    mig = by_legacy["PP-MIG-01"]
    assert mig.asset_number == "PP-MIG-01"
    assert mig.purchased_by == "unknown"
    assert mig.serial_number is None or mig.serial_number == ""
    assert mig.service_tag == "MIGTAG001"
    assert mig.serial_number != "SERIAL_MISSING"
    assert "OWNERSHIP_UNCLEAR" not in (mig.asset_number or "")
    assert "OWNERSHIP_UNCLEAR" not in (mig.notes or "")
    # No password anywhere on notes
    for a in assets:
        assert a.notes is None or "should-never-appear" not in (a.notes or "")
        assert a.notes is None or "tv-secret" not in (a.notes or "")
        assert "EX-HAR" not in (a.asset_number or "")
        assert "SERIAL_MISSING" not in (a.serial_number or "")

    # Re-analyze same file → duplicates detected against DB
    analyzed2 = client.post(
        "/api/v1/it/migration/analyze",
        headers=client.auth_headers,
        data={"source_type": "hardware"},
        files={
            "file": (
                "hardware.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert analyzed2.status_code == 200
    assert analyzed2.json()["potential_duplicates"] >= 1


def _normalized_hardware_workbook_bytes() -> bytes:
    """Format B — ProTrack compatible sheet with snake_case headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "hardware"
    ws.append(
        [
            "external_id",
            "asset_number",
            "legacy_asset_number",
            "asset_type",
            "name",
            "description",
            "ownership_type",
            "owner_customer",
            "customer_used_for",
            "assigned_to",
            "team_or_department",
            "location",
            "service_tag",
            "serial_number",
            "current_status",
        ]
    )
    ws.append(
        [
            "IT - 001",
            "IT - 001",
            "IT - 001",
            "IT",
            "LG DISPLAY 1",
            "LARGE DISPLAY 60 Inch-Conference-2",
            "Prosohm",
            "",
            "PLATINUM",
            "",
            "Sybridge",
            "Everest at top floor",
            "",
            "207PLTV162817",
            "Working",
        ]
    )
    ws.append(
        [
            "IT - 002",
            "IT - 002",
            "IT - 002",
            "IT",
            "LG DISPLAY 2",
            "LARGE DISPLAY 60 Inch-Sybridge Job display",
            "Customer",
            "Sybridge",
            "SYBRIDGE",
            "",
            "Sybridge",
            "Ground floor",
            "",
            "206PLCD174419",
            "Working",
        ]
    )
    ws.append(
        [
            "IT - 003",
            "IT - 003",
            "IT - 003",
            "IT",
            "MI DISPLAY 3",
            "DISPLAY 43 Inch",
            "Prosohm",
            "",
            "",
            "",
            "Prosohm",
            "Trishul",
            "",
            "",
            "Working",
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_normalized_workbook_column_mapping(client, session):
    from app.models.models import Customer

    # Ensure Sybridge exists for owner match
    if session.scalar(select(Customer).where(Customer.name.ilike("%sybridge%"))) is None:
        session.add(Customer(name="Sybridge", code="SYB", is_active=True))
        session.commit()

    content = _normalized_hardware_workbook_bytes()
    response = client.post(
        "/api/v1/it/migration/analyze",
        headers=client.auth_headers,
        data={"source_type": "hardware"},
        files={
            "file": (
                "PP-ProTrack_IT_Import_Compatible.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["workbook_format"] == "normalized_protrack"
    assert body["column_bindings"]["source_id"] == "external_id"
    assert body["column_bindings"]["ownership_type"] == "ownership_type"
    assert body["column_bindings"]["service_tag"] == "service_tag"
    assert body["records_found"] == 3

    preview = body["preview_rows"]
    assert preview[0]["source_id"] == "IT - 001"
    assert preview[0]["owner"] == "Prosohm"
    assert preview[0]["serial"] == "207PLTV162817"
    assert preview[1]["source_id"] == "IT - 002"
    assert "Sybridge" in preview[1]["owner"] or preview[1]["owner"] == "customer"
    assert preview[2]["source_id"] == "IT - 003"

    # Must NOT claim blank Hardware ID
    details = " ".join(str(e.get("detail", "")) for e in body["exceptions"])
    assert "Hardware ID ''" not in details
    assert "Source ID is blank" not in details

    # Ownership known for Prosohm/Customer rows → not 3x OWNERSHIP_UNCLEAR
    ownership_unclear = [
        e for e in body["exceptions"] if e.get("code") == "OWNERSHIP_UNCLEAR"
    ]
    assert len(ownership_unclear) == 0

    # SERIAL_MISSING only for row 3 (no serial/service tag) as warning
    serial_missing = [e for e in body["exceptions"] if e.get("code") == "SERIAL_MISSING"]
    assert len(serial_missing) == 1
    assert serial_missing[0]["severity"] == "warning"
    assert body["requires_review"] == 0


def test_migration_accounts_never_store_passwords(client, session):
    admin = session.scalar(select(User).where(User.email == "admin@prosohm.com"))
    assert admin is not None
    content = _accounts_workbook_bytes(admin.email)

    analyzed = client.post(
        "/api/v1/it/migration/analyze",
        headers=client.auth_headers,
        data={"source_type": "accounts"},
        files={
            "file": (
                "accounts.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert analyzed.status_code == 200, analyzed.text
    body = analyzed.json()
    assert body["sensitive_data_excluded_count"] >= 1
    assert "plaintext-password" not in str(body).lower()

    imported = client.post(
        "/api/v1/it/migration/import",
        headers=client.auth_headers,
        data={"session_id": body["session_id"], "confirm": "true"},
    )
    assert imported.status_code == 200, imported.text

    accounts = list(
        session.scalars(
            select(ITUserAccount).where(ITUserAccount.user_id == admin.id)
        ).all()
    )
    assert accounts
    for acct in accounts:
        assert acct.credential_status == "migration_required"
        blob = f"{acct.username} {acct.notes} {acct.display_name}".lower()
        assert "plaintext" not in blob
        assert "password" not in blob or "passwords not imported" in blob

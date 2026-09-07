"""Sequential IT Data Import — analyze / preview / commit / rollback.

Imports only canonical sheets from the split workbook package under
``docs/migration-sources/split/``. Compat sheets (hardware, inventory_list, …)
are never auto-selected; they appear as override options only.
"""

from __future__ import annotations

import ipaddress
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from openpyxl import load_workbook
from sqlalchemy import delete, func, or_, select
from sqlalchemy.orm import Session

from app.core.access_control import MODULE_IT_OPERATIONS
from app.core.config import UPLOAD_DIR
from app.core.exceptions import ProTrackValidationError
from app.models.enums import ActivityAction, EntityType
from app.models.it_operations import (
    Asset,
    AssetAssignment,
    AssetCustomerReturn,
    AssetType,
    Computer,
    InventoryItem,
    IPAddress,
    IPAssignmentHistory,
    ITImportBatch,
    ITMigrationException,
    ITSupplier,
    ITUserAccount,
    SoftwareAssignment,
    SoftwareCatalog,
    SoftwareLicensePool,
    Network,
)
from app.models.models import Customer, User
from app.services import it_asset_service
from app.services.activity_service import log_activity

MODULE = MODULE_IT_OPERATIONS

_SESSIONS_DIR = UPLOAD_DIR / "it_data_import"

_SENSITIVE_HEADER_RE = re.compile(
    r"(pass|pwd|secret|credential|teamviewer|windows\s*license|license\s*#)",
    re.I,
)

# Canonical sheet names — NEVER auto-import compat sheets.
CANONICAL_SHEETS: dict[str, str] = {
    "hardware_workbook": "IT Assets Import",
    "assets": "Assets",
    "computers": "Computers",
    "ip_addresses": "IP_Addresses",
    "inventory": "Inventory_Items",
    "software": "Software_Licenses",
    "user_accounts": "User_Accounts",
    "suppliers": "Suppliers",
    "migration_exceptions": "Migration_Exceptions",
}

COMPAT_SHEETS: frozenset[str] = frozenset(
    {
        "hardware",
        "inventory_list",
        "consumables",
        "softwares_&_licenses",
        "user_credentials",
        "supplier_list",
        "master data seed",
        "ownership options",
    }
)

# Signature headers used to detect wrong-file uploads (normalized).
_SIGNATURE_HEADERS: dict[str, frozenset[str]] = {
    "hardware_workbook": frozenset({"asset_number", "asset_type"}),
    "assets": frozenset({"external_id", "asset_number", "ownership_type", "asset_type"}),
    "computers": frozenset({"asset_external_id", "asset_number", "computer_name", "computer_type"}),
    "ip_addresses": frozenset({"ip_address", "asset_external_id", "assignment_status"}),
    "inventory": frozenset({"item_name", "external_id", "quantity", "category"}),
    "software": frozenset({"software_name", "license_quantity", "purchased_by"}),
    "user_accounts": frozenset({"employee_name", "username", "company_email"}),
    "suppliers": frozenset({"supplier_name"}),
    "migration_exceptions": frozenset({"type", "issue", "identifier", "action"}),
}

FIELD_ALIASES: dict[str, dict[str, tuple[str, ...]]] = {
    "hardware_workbook": {
        "asset_number": (
            "asset_number",
            "asset number",
            "asset tag",
            "asset tag cpu#",
            "it asset no",
            "it asset number",
            "id_tag",
        ),
        "category": ("category", "asset category"),
        "asset_type": ("asset_type", "asset type", "type"),
        "make": ("make", "make / brand", "make_brand", "brand", "manufacturer"),
        "model_number": ("model", "model_number", "model number"),
        "computer_name": ("pc name", "pc_name", "computer_name", "hostname"),
        "serial_number": ("serial_number", "serial", "serial / service tag"),
        "service_tag": ("service_tag", "service tag", "serial / service tag"),
        "ip_address": ("ip address", "ip_address", "ip"),
        "mac_address": ("mac address", "mac_address", "mac"),
        "assigned_to": (
            "assigned user",
            "assigned_to",
            "assigned_user",
            "current_user",
        ),
        "current_status": ("asset status", "current_status", "status"),
        "ownership_type": (
            "hardware owner type",
            "ownership_type",
            "owner_type",
            "purchased_by",
        ),
        "owner_customer": (
            "hardware owner",
            "owner_customer",
            "owner",
            "cust_paid",
        ),
        "parent_asset_number": (
            "parent asset number",
            "parent_asset_number",
            "parent asset",
        ),
        "migration_note": ("notes", "migration_note", "remarks"),
        "description": ("description",),
        "location": ("location",),
    },
    "assets": {
        "external_id": ("external_id", "source_id", "id"),
        "asset_number": (
            "asset_number",
            "asset number",
            "id_tag",
            "asset tag",
            "it asset no",
        ),
        "legacy_asset_number": ("legacy_asset_number",),
        "category": ("category", "asset category"),
        "asset_type": ("asset_type", "asset type", "type"),
        "make": ("make", "make / brand", "brand", "manufacturer"),
        "name": ("name",),
        "description": ("description",),
        "ownership_type": (
            "ownership_type",
            "owner_type",
            "purchased_by",
            "hardware owner type",
        ),
        "owner_customer": (
            "owner_customer",
            "owner",
            "cust_paid",
            "hardware owner",
        ),
        "customer_used_for": ("customer_used_for", "customer_used"),
        "assigned_to": (
            "assigned_to",
            "assigned_user",
            "assigned user",
            "current_user",
        ),
        "team_or_department": ("team_or_department", "department", "team"),
        "location": ("location", "room", "storage_location"),
        "purchase_date": ("purchase_date", "date_of_purchase", "date"),
        "supplier": ("supplier",),
        "service_tag": ("service_tag", "service_tag_number", "serial / service tag"),
        "warranty_expiry": ("warranty_expiry", "warranty_valid_upto", "warranty_valid_until"),
        "condition": ("condition",),
        "quantity": ("quantity", "qty"),
        "purchase_value": ("purchase_value", "value"),
        "model_number": ("model_number", "model"),
        "serial_number": ("serial_number", "serial"),
        "invoice_number": ("invoice_number", "invoice_no", "invoice"),
        "current_status": ("current_status", "status", "asset status"),
        "is_current_asset": ("is_current_asset", "current_asset"),
        "returned_to_owner": ("returned_to_owner", "returned_to_customer"),
        "return_date": ("return_date",),
        "return_notes": ("return_notes",),
        "parent_asset_number": (
            "parent_asset_number",
            "parent asset number",
            "parent asset",
        ),
        "computer_name": ("pc name", "pc_name", "computer_name", "hostname"),
        "ip_address": ("ip address", "ip_address", "ip"),
        "mac_address": ("mac address", "mac_address", "mac"),
        "source_system": ("source_system",),
        "source_row": ("source_row",),
        "migration_note": ("migration_note", "remarks", "notes"),
    },
    "computers": {
        "asset_external_id": ("asset_external_id", "external_id"),
        "asset_number": ("asset_number",),
        "computer_name": ("computer_name", "hostname", "name"),
        "computer_type": ("computer_type", "type"),
        "manufacturer": ("manufacturer", "make"),
        "model": ("model", "model_number"),
        "serial_number": ("serial_number", "serial"),
        "service_tag": ("service_tag",),
        "cpu": ("cpu", "processor"),
        "ram_gb": ("ram", "ram_gb"),
        "gpu": ("gpu",),
        "operating_system": ("operating_system", "os"),
        "mac_address": ("mac_address", "mac"),
        "assigned_to": ("assigned_to",),
        "customer_used_for": ("customer_used_for",),
        "status": ("status",),
        "purchase_date": ("purchase_date",),
        "warranty_expiry": ("warranty_expiry",),
        "source_system": ("source_system",),
    },
    "ip_addresses": {
        "ip_address": ("ip_address", "address", "ip"),
        "asset_external_id": ("asset_external_id", "external_id"),
        "asset_number": ("asset_number",),
        "computer_name": ("computer_name", "hostname"),
        "assignment_status": ("assignment_status", "status"),
        "source_system": ("source_system",),
    },
    "inventory": {
        "external_id": ("external_id", "source_id"),
        "item_name": ("item_name", "name"),
        "description": ("description",),
        "category": ("category",),
        "department": ("department",),
        "storage_location": ("storage_location", "location"),
        "purchase_date": ("purchase_date",),
        "supplier": ("supplier",),
        "condition": ("condition",),
        "quantity": ("quantity", "qty", "total_qty"),
        "quantity_in_use": ("quantity_in_use", "in_use", "issued_qty"),
        "quantity_available": ("quantity_available", "available"),
        "unit_or_total_value": ("unit_or_total_value", "unit_cost", "value", "purchase_value"),
        "model_number": ("model_number", "model"),
        "serial_number": ("serial_number",),
        "ownership_type": ("ownership_type", "purchased_by"),
        "owner_customer": ("owner_customer",),
        "current_status": ("current_status", "status"),
        "source_system": ("source_system",),
        "source_row": ("source_row",),
    },
    "software": {
        "software_name": ("software_name", "software", "name"),
        "purchased_by": ("purchased_by",),
        "ownership_type": ("ownership_type",),
        "owner_customer": ("owner_customer",),
        "license_quantity": ("license_quantity", "seats", "no_of_user_license"),
        "license_valid_until": ("license_valid_until", "expiry", "license_valid_upto"),
        "renewal_mode": ("renewal_mode", "renewal"),
        "assigned_user": ("assigned_user", "assigned_to", "user"),
        "department": ("department",),
        "comments": ("comments", "notes"),
        "source_system": ("source_system",),
    },
    "user_accounts": {
        "employee_name": ("employee_name", "employee", "name"),
        "employee_status": ("employee_status", "status", "account_status"),
        "department": ("department",),
        "designation": ("designation",),
        "customer_assigned_to": ("customer_assigned_to",),
        "username": ("username",),
        "powerapps_or_pwa_id": ("powerapps_or_pwa_id", "pwa_id"),
        "data_complete_id": ("data_complete_id",),
        "m365_teams_id": ("m365_teams_id", "ms_teams_id", "teams_id"),
        "company_email": ("company_email", "email"),
        "customer_teams_id": ("customer_teams_id",),
        "credential_status": ("credential_status",),
        "source_system": ("source_system",),
    },
    "suppliers": {
        "supplier_name": ("supplier_name", "name"),
        "product": ("product", "products"),
        "lead_time": ("lead_time",),
        "product_link": ("product_link",),
        "description": ("description",),
        "website": ("website",),
        "phone": ("phone",),
        "address": ("address",),
        "email": ("email",),
        "remarks": ("remarks", "notes"),
        "source_system": ("source_system",),
    },
    "migration_exceptions": {
        "exception_type": ("type", "exception_type"),
        "source": ("source",),
        "source_row": ("source_row",),
        "identifier": ("identifier",),
        "issue": ("issue",),
        "action": ("action",),
    },
}

IMPORT_TYPES: list[dict[str, Any]] = [
    {
        "id": "hardware_workbook",
        "label": "IT Hardware Workbook",
        "expected_filename": "ProTrack_Current_IT_Assets_Import.xlsx",
        "recommended_order": 0,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "assets",
        "label": "Assets",
        "expected_filename": "01_Assets.xlsx",
        "recommended_order": 1,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "computers",
        "label": "Computers",
        "expected_filename": "02_Computers.xlsx",
        "recommended_order": 2,
        "depends_on": ["assets"],
        "review_only": False,
    },
    {
        "id": "ip_addresses",
        "label": "IP Addresses",
        "expected_filename": "03_IP_Addresses.xlsx",
        "recommended_order": 3,
        "depends_on": ["assets", "computers"],
        "review_only": False,
    },
    {
        "id": "inventory",
        "label": "Inventory Items",
        "expected_filename": "04_Inventory_Items.xlsx",
        "recommended_order": 4,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "software",
        "label": "Software Licenses",
        "expected_filename": "05_Software_Licenses.xlsx",
        "recommended_order": 5,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "user_accounts",
        "label": "User Accounts",
        "expected_filename": "06_User_Accounts.xlsx",
        "recommended_order": 6,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "suppliers",
        "label": "Suppliers",
        "expected_filename": "07_Suppliers.xlsx",
        "recommended_order": 7,
        "depends_on": [],
        "review_only": False,
    },
    {
        "id": "migration_exceptions",
        "label": "Migration Exceptions (review)",
        "expected_filename": "08_Migration_Exceptions.xlsx",
        "recommended_order": 8,
        "depends_on": [],
        "review_only": True,
    },
]

_IMPORT_TYPE_IDS = frozenset(t["id"] for t in IMPORT_TYPES)
_IMPORT_TYPE_BY_ID = {t["id"]: t for t in IMPORT_TYPES}


# ---------------------------------------------------------------------------
# Helpers (mirrored from it_migration_service patterns)
# ---------------------------------------------------------------------------


def _norm_header(header: str) -> str:
    s = (header or "").strip().lower()
    s = s.replace("&", " and ")
    s = re.sub(r"[?#]+", "", s)
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return re.sub(r"_+", "_", s).strip("_")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _is_sensitive_header(header: str) -> bool:
    """True for password/secret columns. Never treat credential_status as sensitive."""
    norm = _norm_header(header)
    if norm in {"credential_status", "employee_status", "account_status", "assignment_status"}:
        return False
    return bool(_SENSITIVE_HEADER_RE.search(header or ""))


def _norm_id(value: str) -> str:
    return re.sub(r"\s+", "", (value or "").upper())


def _norm_name(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def _parse_date(value: str) -> date | None:
    raw = (value or "").strip()
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:19], fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        return None


def _parse_int(value: str) -> int | None:
    raw = (value or "").strip()
    if not raw:
        return None
    try:
        return int(float(raw))
    except ValueError:
        return None


def _parse_decimal(value: str) -> Decimal | None:
    raw = (value or "").strip().replace(",", "")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def _sessions_dir() -> Path:
    _SESSIONS_DIR.mkdir(parents=True, exist_ok=True)
    return _SESSIONS_DIR


def _session_path(session_id: str) -> Path:
    safe = re.sub(r"[^a-zA-Z0-9_-]", "", session_id)
    if not safe:
        raise ProTrackValidationError("Invalid import session id.")
    return _sessions_dir() / f"{safe}.json"


def _save_session(payload: dict[str, Any], session_id: str | None = None) -> str:
    sid = session_id or str(uuid4())
    _session_path(sid).write_text(json.dumps(payload, default=str), encoding="utf-8")
    return sid


def _load_session(session_id: str) -> dict[str, Any]:
    path = _session_path(session_id)
    if not path.exists():
        raise ProTrackValidationError("Import session not found or expired.")
    return json.loads(path.read_text(encoding="utf-8"))


def _delete_session(session_id: str) -> None:
    path = _session_path(session_id)
    if path.exists():
        path.unlink()


def _customer_index(db: Session) -> dict[str, Customer]:
    customers = list(db.scalars(select(Customer).where(Customer.is_active.is_(True))).all())
    index: dict[str, Customer] = {}
    for c in customers:
        index[_norm_name(c.name)] = c
        if c.code:
            index[_norm_name(c.code)] = c
    return index


def _user_index(db: Session) -> dict[str, User]:
    users = list(db.scalars(select(User).where(User.is_deleted.is_(False))).all())
    index: dict[str, User] = {}
    for u in users:
        full = f"{u.first_name} {u.last_name}".strip()
        if full:
            index[_norm_name(full)] = u
        if u.email:
            index[_norm_name(u.email)] = u
            index[_norm_name(u.email.split("@")[0])] = u
    return index


def _match_customer(index: dict[str, Customer], label: str) -> Customer | None:
    key = _norm_name(label)
    if not key:
        return None
    if key in {"prosohm", "prosohm eng", "prosohm admin", "generic", "admin", "organization"}:
        return None
    if key in index:
        return index[key]
    for k, cust in index.items():
        if key in k or k in key:
            return cust
    return None


def _match_user(index: dict[str, User], label: str) -> User | None:
    key = _norm_name(label)
    if not key:
        return None
    if key in {
        "all",
        "open",
        "n/a",
        "servprosohm",
        "conference room",
        "planning board",
        "eng planning board",
        "prosohm eng team",
    }:
        return None
    if key in index:
        return index[key]
    base = re.sub(r"\(.*?\)", "", key).strip()
    if base in index:
        return index[base]
    for k, user in index.items():
        if base and (base in k or k in base):
            return user
    return None


def _existing_asset_keys(db: Session) -> dict[str, set[str]]:
    assets = list(db.scalars(select(Asset).where(Asset.is_deleted.is_(False))).all())
    return {
        "legacy": {_norm_id(a.legacy_asset_number or "") for a in assets if a.legacy_asset_number},
        "serial": {_norm_id(a.serial_number or "") for a in assets if a.serial_number},
        "service_tag": {_norm_id(a.service_tag or "") for a in assets if a.service_tag},
        "asset_number": {_norm_id(a.asset_number or "") for a in assets if a.asset_number},
    }


def _find_asset_by_number(db: Session, *labels: str) -> Asset | None:
    for label in labels:
        key = (label or "").strip()
        if not key:
            continue
        asset = db.scalar(
            select(Asset).where(
                Asset.is_deleted.is_(False),
                Asset.asset_number == key,
            )
        )
        if asset:
            return asset
        norm = _norm_id(key)
        assets = list(db.scalars(select(Asset).where(Asset.is_deleted.is_(False))).all())
        for a in assets:
            if _norm_id(a.asset_number) == norm or _norm_id(a.legacy_asset_number or "") == norm:
                return a
    return None


def _resolve_ownership(
    ownership_raw: str,
    owner_raw: str,
    customers: dict[str, Customer],
) -> tuple[str, UUID | None, list[dict[str, str]]]:
    """Map ownership_type / owner_customer → purchased_by + owner id + issues."""
    issues: list[dict[str, str]] = []
    ownership = (ownership_raw or "").strip()
    owner = (owner_raw or "").strip()

    def _org() -> tuple[str, UUID | None, list[dict[str, str]]]:
        return "organization", None, issues

    def _cust(label: str) -> tuple[str, UUID | None, list[dict[str, str]]]:
        matched = _match_customer(customers, label)
        if matched:
            return "customer", matched.id, issues
        issues.append(
            {
                "code": "CUSTOMER_UNKNOWN",
                "severity": "review",
                "detail": f"Owner customer '{label}' not matched; ownership left unknown.",
            }
        )
        return "unknown", None, issues

    if ownership:
        n = _norm_name(ownership)
        if n in {"prosohm", "organization", "company", "prosohm eng", "ps"}:
            return _org()
        if n in {"customer", "client"}:
            if owner:
                return _cust(owner)
            issues.append(
                {
                    "code": "OWNERSHIP_UNCLEAR",
                    "severity": "review",
                    "detail": "ownership_type=Customer but owner_customer is blank.",
                }
            )
            return "unknown", None, issues
        if n not in {"unknown", "n/a", "na", "none", "null"}:
            return _cust(ownership)

    if owner:
        n = _norm_name(owner)
        if n in {"prosohm", "organization", "company"}:
            return _org()
        return _cust(owner)

    return "unknown", None, issues


def _map_asset_status(
    current_status: str,
    is_current_asset: str,
    returned_to_owner: str,
) -> str:
    ret = _norm_name(returned_to_owner)
    if ret in {"yes", "y", "true", "1"}:
        return "returned_to_customer"
    curr = (current_status or "").strip().upper()
    is_curr = _norm_name(is_current_asset)
    if "DISPOSE" in curr or "SCRAP" in curr:
        return "disposed"
    if "RETIRE" in curr:
        return "retired"
    if "DAMAGE" in curr or "NOT WORKING" in curr:
        return "damaged"
    if "RETURN" in curr:
        return "returned_to_customer"
    if "MAINTENANCE" in curr or "REPAIR" in curr:
        return "maintenance"
    if is_curr in {"no", "n", "false", "0"}:
        if "DISPOSE" in curr:
            return "disposed"
        return "retired"
    if "ASSIGN" in curr:
        return "assigned"
    if curr in {"CURRENT", "WORKING", "AVAILABLE", "ACTIVE", ""}:
        return "available"
    return "available"


def _resolve_or_create_asset_type(db: Session, type_label: str) -> AssetType:
    raw = (type_label or "OTHER").strip() or "OTHER"
    code = re.sub(r"[^A-Za-z0-9]+", "_", raw).strip("_").upper()[:20] or "OTHER"
    at = db.scalar(select(AssetType).where(func.upper(AssetType.code) == code))
    if at is not None:
        return at
    at = db.scalar(select(AssetType).where(func.lower(AssetType.name) == raw.lower()))
    if at is not None:
        return at
    # Common aliases
    alias_map = {
        "LAPTOP": "LAPTOP",
        "DESKTOP": "DESKTOP",
        "WORKSTATION": "DESKTOP",
        "MONITOR": "MONITOR",
        "SERVER": "SERVER",
        "KEYBOARD": "KEYBOARD",
        "MOUSE": "MOUSE",
    }
    mapped = alias_map.get(code)
    if mapped:
        at = db.scalar(select(AssetType).where(AssetType.code == mapped))
        if at is not None:
            return at
    at = AssetType(
        id=uuid4(),
        code=code,
        name=raw,
        category="other",
        is_active=True,
        numbering_prefix=code[:6] if len(code) <= 6 else None,
    )
    db.add(at)
    db.flush()
    return at


def _get_or_create_supplier(
    db: Session,
    name: str,
    *,
    import_batch_id: UUID | None = None,
    source_system: str | None = None,
    extra: dict[str, Any] | None = None,
) -> ITSupplier | None:
    norm = _norm_name(name)
    if not norm:
        return None
    existing = list(db.scalars(select(ITSupplier)).all())
    for s in existing:
        if _norm_name(s.name) == norm:
            return s
    for s in existing:
        if norm in _norm_name(s.name) or _norm_name(s.name) in norm:
            return s
    fields = extra or {}
    supplier = ITSupplier(
        id=uuid4(),
        name=name.strip(),
        website=(fields.get("website") or None),
        phone=(fields.get("phone") or None),
        email=(fields.get("email") or None),
        address=(fields.get("address") or None),
        products=(fields.get("product") or fields.get("products") or None),
        notes=(fields.get("remarks") or fields.get("notes") or None),
        is_active=True,
        import_batch_id=import_batch_id,
        source_system=source_system,
    )
    db.add(supplier)
    db.flush()
    return supplier


def _next_batch_code(db: Session) -> str:
    today = date.today().strftime("%Y%m%d")
    prefix = f"IT-IMPORT-{today}-"
    codes = list(
        db.scalars(
            select(ITImportBatch.batch_code).where(ITImportBatch.batch_code.like(f"{prefix}%"))
        ).all()
    )
    nums: list[int] = []
    for code in codes:
        try:
            nums.append(int(str(code).rsplit("-", 1)[-1]))
        except ValueError:
            continue
    n = (max(nums) if nums else 0) + 1
    return f"{prefix}{n:03d}"


def _latest_committed_batch(db: Session, import_type: str) -> ITImportBatch | None:
    return db.scalar(
        select(ITImportBatch)
        .where(
            ITImportBatch.import_type == import_type,
            ITImportBatch.status == "committed",
        )
        .order_by(ITImportBatch.created_at.desc())
        .limit(1)
    )


def _dependency_warnings(db: Session, import_type: str) -> list[str]:
    meta = _IMPORT_TYPE_BY_ID.get(import_type) or {}
    warnings: list[str] = []
    for dep in meta.get("depends_on") or []:
        if _latest_committed_batch(db, dep) is None:
            dep_label = (_IMPORT_TYPE_BY_ID.get(dep) or {}).get("label", dep)
            warnings.append(
                f"Recommended dependency '{dep_label}' has no committed import yet. "
                "You can still proceed; linking may be incomplete."
            )
    return warnings


# ---------------------------------------------------------------------------
# Workbook / column binding
# ---------------------------------------------------------------------------


def _inspect_workbook(content: bytes) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for name in wb.sheetnames:
            ws = wb[name]
            raw = list(ws.iter_rows(values_only=True))
            headers: list[str] = []
            data_rows = 0
            if raw:
                headers_raw = [_cell_str(c) for c in (raw[0] or ())]
                while headers_raw and not headers_raw[-1]:
                    headers_raw.pop()
                headers = headers_raw
                for row in raw[1:]:
                    vals = [_cell_str(c) for c in (row or ())]
                    if any(vals):
                        data_rows += 1
            sheets.append(
                {
                    "name": name,
                    "headers": headers,
                    "row_count": data_rows,
                    "is_compat": _norm_header(name) in {_norm_header(c) for c in COMPAT_SHEETS},
                }
            )
        return {"sheetnames": list(wb.sheetnames), "sheets": sheets}
    finally:
        wb.close()


def _pick_sheet(
    sheetnames: list[str],
    import_type: str,
    override: str | None,
) -> tuple[str, list[str]]:
    """Return (selected_sheet, other_sheets)."""
    if override:
        if override not in sheetnames:
            raise ProTrackValidationError(
                f"Sheet '{override}' not found. Available: {', '.join(sheetnames)}"
            )
        others = [s for s in sheetnames if s != override]
        return override, others

    canonical = CANONICAL_SHEETS[import_type]
    # Exact match first, then case-insensitive.
    if canonical in sheetnames:
        selected = canonical
    else:
        lowered = {s.lower(): s for s in sheetnames}
        selected = lowered.get(canonical.lower())
        if selected is None:
            raise ProTrackValidationError(
                f"Canonical sheet '{canonical}' not found for import type '{import_type}'. "
                f"Available: {', '.join(sheetnames)}. "
                "Pass sheet_name to override (compat sheets are never auto-selected)."
            )
    others = [s for s in sheetnames if s != selected]
    return selected, others


def _read_sheet(
    content: bytes,
    sheet_name: str,
) -> tuple[list[str], list[dict[str, str]], list[str]]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ProTrackValidationError(f"Sheet '{sheet_name}' not found.")
        ws = wb[sheet_name]
        raw_rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not raw_rows:
        return [], [], []

    headers_raw = [_cell_str(c) for c in (raw_rows[0] or ())]
    while headers_raw and not headers_raw[-1]:
        headers_raw.pop()

    headers: list[str] = []
    seen: dict[str, int] = {}
    for h in headers_raw:
        key = h or f"col_{len(headers)}"
        if key in seen:
            seen[key] += 1
            key = f"{key}_{seen[key]}"
        else:
            seen[key] = 1
        headers.append(key)

    sensitive = [h for h in headers if _is_sensitive_header(h)]
    safe_headers = [h for h in headers if h not in sensitive]

    rows: list[dict[str, str]] = []
    for row in raw_rows[1:]:
        vals = [_cell_str(c) for c in (row or ())]
        if len(vals) < len(headers):
            vals += [""] * (len(headers) - len(vals))
        record = {headers[i]: vals[i] for i in range(len(headers))}
        meaningful = [v for k, v in record.items() if v and k not in sensitive]
        if not meaningful:
            continue
        rows.append({k: record.get(k, "") for k in safe_headers})

    return safe_headers, rows, sensitive


def _build_bindings(
    headers: list[str],
    import_type: str,
) -> tuple[dict[str, str | None], list[str]]:
    """Map logical fields → source header; return unmapped non-sensitive headers."""
    index = {_norm_header(h): h for h in headers if h}
    aliases = FIELD_ALIASES.get(import_type, {})
    bindings: dict[str, str | None] = {}
    used_headers: set[str] = set()
    for logical, alias_list in aliases.items():
        matched = None
        for alias in alias_list:
            key = index.get(_norm_header(alias))
            if key is not None:
                matched = key
                used_headers.add(key)
                break
        bindings[logical] = matched
    unmapped = [h for h in headers if h not in used_headers and not _is_sensitive_header(h)]
    return bindings, unmapped


def _map_row(row: dict[str, str], bindings: dict[str, str | None]) -> dict[str, str]:
    out: dict[str, str] = {}
    for logical, header in bindings.items():
        if header is None:
            out[logical] = ""
        else:
            out[logical] = (row.get(header) or "").strip()
    return out


def _validate_file_signature(import_type: str, headers: list[str], filename: str) -> None:
    sig = _SIGNATURE_HEADERS.get(import_type, frozenset())
    if not sig:
        return
    norms = {_norm_header(h) for h in headers}
    hits = len(sig & norms)
    if hits < max(1, len(sig) // 2):
        expected = (_IMPORT_TYPE_BY_ID.get(import_type) or {}).get("expected_filename", "")
        raise ProTrackValidationError(
            f"File does not look like a '{import_type}' import. "
            f"Expected headers near {sorted(sig)}; found {headers[:12]}. "
            f"Expected filename hint: {expected or 'n/a'} (got '{filename}')."
        )


# ---------------------------------------------------------------------------
# Public API — list / get
# ---------------------------------------------------------------------------


def list_import_types(db: Session) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for meta in IMPORT_TYPES:
        latest = _latest_committed_batch(db, meta["id"])
        status = "imported" if latest else "not_imported"
        if not latest and meta.get("depends_on"):
            missing = [
                dep
                for dep in meta["depends_on"]
                if _latest_committed_batch(db, dep) is None
            ]
            if missing:
                status = f"waiting_for_{missing[0]}"
        if meta.get("review_only") and not latest:
            status = "review_only"
        result.append(
            {
                **meta,
                "canonical_sheet": CANONICAL_SHEETS[meta["id"]],
                "status": status,
                "latest_batch_id": str(latest.id) if latest else None,
                "latest_batch_code": latest.batch_code if latest else None,
                "latest_committed_at": latest.updated_at.isoformat()
                if latest and latest.updated_at
                else None,
                "success_count": latest.success_count if latest else 0,
            }
        )
    return result


def list_batches(
    db: Session,
    *,
    import_type: str | None = None,
    status: str | None = None,
    limit: int = 50,
) -> list[ITImportBatch]:
    stmt = select(ITImportBatch).order_by(ITImportBatch.created_at.desc())
    if import_type:
        stmt = stmt.where(ITImportBatch.import_type == import_type)
    if status:
        stmt = stmt.where(ITImportBatch.status == status)
    stmt = stmt.limit(max(1, min(limit, 200)))
    return list(db.scalars(stmt).all())


def get_batch(db: Session, batch_id: UUID) -> ITImportBatch:
    batch = db.get(ITImportBatch, batch_id)
    if batch is None:
        raise ProTrackValidationError("Import batch not found.")
    return batch


def _batch_to_dict(batch: ITImportBatch) -> dict[str, Any]:
    return {
        "id": str(batch.id),
        "batch_code": batch.batch_code,
        "import_type": batch.import_type,
        "filename": batch.filename,
        "sheet_name": batch.sheet_name,
        "uploaded_by_user_id": str(batch.uploaded_by_user_id),
        "status": batch.status,
        "record_count": batch.record_count,
        "success_count": batch.success_count,
        "skipped_count": batch.skipped_count,
        "error_count": batch.error_count,
        "warning_count": batch.warning_count,
        "session_id": batch.session_id,
        "summary_json": batch.summary_json,
        "notes": batch.notes,
        "created_at": batch.created_at.isoformat() if batch.created_at else None,
        "updated_at": batch.updated_at.isoformat() if getattr(batch, "updated_at", None) else None,
    }


# ---------------------------------------------------------------------------
# Analyze
# ---------------------------------------------------------------------------


def analyze_upload(
    db: Session,
    *,
    actor: User,
    content: bytes,
    filename: str,
    import_type: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    import_type = (import_type or "").strip().lower()
    if import_type not in _IMPORT_TYPE_IDS:
        raise ProTrackValidationError(
            f"Invalid import_type. Allowed: {', '.join(sorted(_IMPORT_TYPE_IDS))}"
        )
    if not content:
        raise ProTrackValidationError("Uploaded file is empty.")

    inspection = _inspect_workbook(content)
    selected_sheet, other_sheets = _pick_sheet(
        inspection["sheetnames"], import_type, sheet_name
    )
    headers, raw_rows, sensitive = _read_sheet(content, selected_sheet)
    _validate_file_signature(import_type, headers, filename)
    bindings, unmapped = _build_bindings(headers, import_type)

    customers = _customer_index(db)
    users = _user_index(db)
    existing = _existing_asset_keys(db)
    existing_computers = {
        _norm_id(c.computer_name)
        for c in db.scalars(select(Computer)).all()
        if c.computer_name
    }
    existing_ips = {
        (ip.address or "").strip()
        for ip in db.scalars(select(IPAddress)).all()
        if ip.address
    }
    existing_inventory = {
        _norm_name(i.name)
        for i in db.scalars(select(InventoryItem).where(InventoryItem.is_deleted.is_(False))).all()
    }
    existing_suppliers = {
        _norm_name(s.name) for s in db.scalars(select(ITSupplier)).all()
    }

    mapped_rows: list[dict[str, Any]] = []
    first_10: list[dict[str, Any]] = []
    error_count = 0
    warning_count = 0
    review_count = 0
    duplicate_count = 0
    importable_count = 0
    skip_count = 0
    id_column_present = False
    id_blank_count = 0
    id_total = 0

    # Detect presence of ID-like columns for fidelity gate.
    id_binding_keys = {
        "hardware_workbook": ("asset_number",),
        "assets": ("external_id", "asset_number"),
        "computers": ("asset_external_id", "asset_number"),
        "inventory": ("external_id",),
        "ip_addresses": ("ip_address",),
    }
    id_keys = id_binding_keys.get(import_type, ())
    if import_type == "hardware_workbook":
        # Blank Asset Numbers are allowed (system-generated); do not fidelity-block.
        id_column_present = False
    else:
        id_column_present = any(bindings.get(k) for k in id_keys)
    existing_macs = {
        (c.mac_address or "").strip().upper()
        for c in db.scalars(select(Computer)).all()
        if c.mac_address
    }
    ownership_required_count = 0
    user_match_required_count = 0

    for idx, raw in enumerate(raw_rows, start=2):  # Excel-ish (header = row 1)
        fields = _map_row(raw, bindings)
        issues: list[dict[str, str]] = []
        action = "import"

        if import_type == "hardware_workbook":
            from app.services import it_hardware_workbook_import as hw

            if id_column_present:
                id_total += 1
                if not (fields.get("asset_number") or "").strip():
                    id_blank_count += 1
            action, issues, extras = hw.analyze_hardware_issues(
                fields,
                existing=existing,
                existing_ips=existing_ips,
                existing_macs=existing_macs,
                users=users,
                customers=customers,
                resolve_ownership=_resolve_ownership,
                match_user=_match_user,
                match_customer=_match_customer,
                preview_next_asset_number=it_asset_service.preview_next_asset_number,
                resolve_or_create_asset_type=_resolve_or_create_asset_type,
                db=db,
            )
            fields = extras.get("fields") or fields
            preview = extras.get("preview") or {}
            if action == "skip_duplicate":
                duplicate_count += 1
            if any(i.get("code") == "OWNERSHIP_REQUIRED" for i in issues):
                ownership_required_count += 1
            if any(i.get("code") == "USER_UNKNOWN" for i in issues):
                user_match_required_count += 1

        elif import_type == "assets":
            ext = fields.get("external_id") or ""
            anum = fields.get("asset_number") or ext
            legacy = fields.get("legacy_asset_number") or anum
            if id_column_present:
                id_total += 1
                if not (ext or anum):
                    id_blank_count += 1
            if not anum and not ext:
                issues.append(
                    {
                        "code": "MISSING_SOURCE_ID",
                        "severity": "error",
                        "detail": "asset_number / external_id is blank.",
                    }
                )
                action = "error"
            else:
                norm = _norm_id(anum or ext)
                strong = None
                st = fields.get("service_tag") or ""
                sn = fields.get("serial_number") or ""
                if st and _norm_id(st) in existing["service_tag"]:
                    strong = "service_tag"
                elif sn and _norm_id(sn) in existing["serial"]:
                    strong = "serial_number"
                elif norm in existing["asset_number"] or norm in existing["legacy"]:
                    strong = "asset_number"
                if strong:
                    action = "skip_duplicate"
                    issues.append(
                        {
                            "code": "DUPLICATE",
                            "severity": "warning",
                            "detail": f"Matches existing asset on {strong}.",
                        }
                    )
                    duplicate_count += 1
                purchased_by, _oid, own_issues = _resolve_ownership(
                    fields.get("ownership_type") or "",
                    fields.get("owner_customer") or "",
                    customers,
                )
                fields["_purchased_by"] = purchased_by
                issues.extend(own_issues)
                assignee = fields.get("assigned_to") or ""
                from app.services.it_hardware_workbook_import import is_placeholder_assignee

                if assignee and is_placeholder_assignee(assignee):
                    fields["assigned_to"] = ""
                    if not (fields.get("current_status") or "").strip():
                        fields["current_status"] = "Available"
                elif assignee and _match_user(users, assignee) is None:
                    issues.append(
                        {
                            "code": "USER_UNKNOWN",
                            "severity": "warning",
                            "detail": f"Assigned user '{assignee}' not matched.",
                        }
                    )
                cust_used = fields.get("customer_used_for") or ""
                if cust_used and _match_customer(customers, cust_used) is None:
                    if _norm_name(cust_used) not in {
                        "generic",
                        "prosohm eng",
                        "prosohm admin",
                        "prosohm",
                    }:
                        issues.append(
                            {
                                "code": "CUSTOMER_UNKNOWN",
                                "severity": "warning",
                                "detail": f"customer_used_for '{cust_used}' not matched.",
                            }
                        )
            preview = {
                "asset_number": anum or ext or "",
                "external_id": ext or anum or "",
                "name": fields.get("name") or "",
                "description": fields.get("description") or "",
                "ownership_type": fields.get("ownership_type") or "",
                "owner_customer": fields.get("owner_customer") or "",
                "asset_type": fields.get("asset_type") or "",
                "current_status": fields.get("current_status") or "",
                "import_action": action,
            }

        elif import_type == "computers":
            link = fields.get("asset_number") or fields.get("asset_external_id") or ""
            cname = fields.get("computer_name") or ""
            if not link:
                issues.append(
                    {
                        "code": "MISSING_ASSET_LINK",
                        "severity": "error",
                        "detail": "asset_number / asset_external_id is required.",
                    }
                )
                action = "error"
            else:
                asset = _find_asset_by_number(
                    db, fields.get("asset_number") or "", fields.get("asset_external_id") or ""
                )
                if asset is None:
                    issues.append(
                        {
                            "code": "ASSET_NOT_FOUND",
                            "severity": "review",
                            "detail": f"No asset found for '{link}'. Import assets first.",
                        }
                    )
                display_name = cname or (asset.asset_number if asset else link)
                if display_name and _norm_id(display_name) in existing_computers:
                    action = "skip_duplicate"
                    duplicate_count += 1
                    issues.append(
                        {
                            "code": "DUPLICATE",
                            "severity": "warning",
                            "detail": f"Computer '{display_name}' already exists.",
                        }
                    )
            preview = {
                "asset_number": link,
                "computer_name": cname,
                "computer_type": fields.get("computer_type") or "",
                "manufacturer": fields.get("manufacturer") or "",
                "model": fields.get("model") or "",
                "import_action": action,
            }

        elif import_type == "ip_addresses":
            ip = fields.get("ip_address") or ""
            if id_column_present:
                id_total += 1
                if not ip:
                    id_blank_count += 1
            if not ip:
                issues.append(
                    {
                        "code": "MISSING_IP",
                        "severity": "error",
                        "detail": "ip_address is blank.",
                    }
                )
                action = "error"
            else:
                try:
                    ipaddress.IPv4Address(ip)
                except ValueError:
                    issues.append(
                        {
                            "code": "INVALID_IP",
                            "severity": "error",
                            "detail": f"Invalid IPv4 address '{ip}'.",
                        }
                    )
                    action = "error"
                if action != "error" and ip in existing_ips:
                    action = "skip_duplicate"
                    duplicate_count += 1
                    issues.append(
                        {
                            "code": "DUPLICATE",
                            "severity": "warning",
                            "detail": f"IP {ip} already exists.",
                        }
                    )
            preview = {
                "ip_address": ip,
                "asset_number": fields.get("asset_number") or fields.get("asset_external_id") or "",
                "computer_name": fields.get("computer_name") or "",
                "assignment_status": fields.get("assignment_status") or "",
                "import_action": action,
            }

        elif import_type == "inventory":
            name = fields.get("item_name") or ""
            if not name:
                issues.append(
                    {
                        "code": "MISSING_NAME",
                        "severity": "error",
                        "detail": "item_name is blank.",
                    }
                )
                action = "error"
            elif _norm_name(name) in existing_inventory:
                action = "skip_duplicate"
                duplicate_count += 1
            purchased_by, _oid, own_issues = _resolve_ownership(
                fields.get("ownership_type") or "",
                fields.get("owner_customer") or "",
                customers,
            )
            fields["_purchased_by"] = purchased_by
            issues.extend(own_issues)
            preview = {
                "item_name": name,
                "external_id": fields.get("external_id") or "",
                "category": fields.get("category") or "",
                "quantity": fields.get("quantity") or "",
                "import_action": action,
            }

        elif import_type == "software":
            sw = fields.get("software_name") or ""
            if not sw:
                issues.append(
                    {
                        "code": "MISSING_SOFTWARE",
                        "severity": "error",
                        "detail": "software_name is blank.",
                    }
                )
                action = "error"
            else:
                purchased_by, _oid, own_issues = _resolve_ownership(
                    fields.get("ownership_type") or fields.get("purchased_by") or "",
                    fields.get("owner_customer") or fields.get("purchased_by") or "",
                    customers,
                )
                fields["_purchased_by"] = purchased_by
                issues.extend(own_issues)
                assignee = fields.get("assigned_user") or ""
                if assignee and _match_user(users, assignee) is None:
                    issues.append(
                        {
                            "code": "USER_UNKNOWN",
                            "severity": "warning",
                            "detail": f"Assigned user '{assignee}' not matched.",
                        }
                    )
            preview = {
                "software_name": sw,
                "purchased_by": fields.get("purchased_by") or fields.get("ownership_type") or "",
                "license_quantity": fields.get("license_quantity") or "",
                "assigned_user": fields.get("assigned_user") or "",
                "import_action": action,
            }

        elif import_type == "user_accounts":
            name = fields.get("employee_name") or ""
            email = fields.get("company_email") or ""
            username = fields.get("username") or ""
            user = (
                _match_user(users, email)
                or _match_user(users, name)
                or _match_user(users, username)
            )
            if user is None:
                issues.append(
                    {
                        "code": "USER_UNKNOWN",
                        "severity": "error",
                        "detail": f"Employee '{name or username or email}' not matched.",
                    }
                )
                action = "error"
            else:
                issues.append(
                    {
                        "code": "SENSITIVE_EXCLUDED",
                        "severity": "info",
                        "detail": "Passwords never imported; credential_status=migration_required.",
                    }
                )
            preview = {
                "employee_name": name,
                "username": username,
                "company_email": email,
                "matched_user": bool(user),
                "import_action": action,
            }

        elif import_type == "suppliers":
            name = fields.get("supplier_name") or ""
            if not name:
                issues.append(
                    {
                        "code": "MISSING_NAME",
                        "severity": "error",
                        "detail": "supplier_name is blank.",
                    }
                )
                action = "error"
            elif _norm_name(name) in existing_suppliers:
                action = "skip_duplicate"
                duplicate_count += 1
            preview = {
                "supplier_name": name,
                "product": fields.get("product") or "",
                "import_action": action,
            }

        else:  # migration_exceptions
            etype = fields.get("exception_type") or ""
            if not etype and not (fields.get("issue") or ""):
                issues.append(
                    {
                        "code": "EMPTY_EXCEPTION",
                        "severity": "warning",
                        "detail": "Exception row has no type or issue.",
                    }
                )
                action = "skip"
            preview = {
                "type": etype,
                "identifier": fields.get("identifier") or "",
                "issue": (fields.get("issue") or "")[:120],
                "import_action": action,
            }

        # Tally severities
        has_error = any(i["severity"] == "error" for i in issues)
        has_review = any(i["severity"] == "review" for i in issues)
        has_warning = any(i["severity"] == "warning" for i in issues)
        if has_error:
            error_count += 1
            if action == "import":
                action = "error"
        if has_review:
            review_count += 1
            if action == "import":
                action = "review"
        if has_warning:
            warning_count += 1

        if action in {"import", "review"}:
            importable_count += 1
        elif action in {"skip", "skip_duplicate", "error"}:
            skip_count += 1

        record = {
            **fields,
            "_source_row": idx,
            "_action": action,
            "_issues": issues,
        }
        mapped_rows.append(record)
        if len(first_10) < 10:
            first_10.append({**preview, "_source_row": idx, "_action": action, "_issues": issues})

    block_commit = False
    fidelity_note = None
    if id_column_present and id_total > 0 and (id_blank_count / id_total) > 0.5:
        block_commit = True
        fidelity_note = (
            f"Fidelity gate: {id_blank_count}/{id_total} mapped ID values are blank "
            "(>50%). Commit blocked until column mapping is corrected."
        )

    if import_type == "suppliers" and importable_count == 0 and not mapped_rows:
        warning_count += 1

    dep_warnings = _dependency_warnings(db, import_type)
    commit_allowed = not block_commit and not (
        import_type != "migration_exceptions" and importable_count == 0 and error_count > 0 and not mapped_rows
    )
    # Still allow empty suppliers sheet commit (0 rows).
    if block_commit:
        commit_allowed = False
    if fidelity_note is None and not mapped_rows and import_type != "suppliers":
        # Empty non-supplier sheets: allow preview but warn.
        pass

    stats = {
        "records_found": len(mapped_rows),
        "importable": importable_count,
        "duplicates": duplicate_count,
        "errors": error_count,
        "warnings": warning_count,
        "reviews": review_count,
        "skipped": skip_count,
        "sensitive_columns_excluded": len(sensitive),
        "id_blank_ratio": (id_blank_count / id_total) if id_total else 0.0,
        "ownership_required": ownership_required_count,
        "user_match_required": user_match_required_count,
    }

    batch = ITImportBatch(
        id=uuid4(),
        batch_code=_next_batch_code(db),
        import_type=import_type,
        filename=filename,
        sheet_name=selected_sheet,
        uploaded_by_user_id=actor.id,
        status="preview",
        record_count=len(mapped_rows),
        success_count=0,
        skipped_count=skip_count,
        error_count=error_count,
        warning_count=warning_count + review_count,
        session_id=None,
        summary_json=json.dumps({"stats": stats, "bindings": bindings}),
        notes=fidelity_note,
    )
    db.add(batch)
    db.flush()

    session_id = str(uuid4())
    session_payload = {
        "filename": filename,
        "import_type": import_type,
        "sheet_name": selected_sheet,
        "headers": headers,
        "bindings": bindings,
        "unmapped_columns": unmapped,
        "sensitive_columns_excluded": sensitive,
        "rows": mapped_rows,
        "batch_id": str(batch.id),
        "stats": stats,
        "created_at": datetime.utcnow().isoformat() + "Z",
        "actor_id": str(actor.id),
    }
    _save_session(session_payload, session_id=session_id)
    batch.session_id = session_id
    db.add(batch)

    log_activity(
        db,
        user=actor,
        entity_type=EntityType.it_settings,
        entity_id=batch.id,
        action=ActivityAction.it_data_import_analyzed,
        new_value={
            "import_type": import_type,
            "filename": filename,
            "sheet_name": selected_sheet,
            "batch_code": batch.batch_code,
            "records": len(mapped_rows),
            "session_id": session_id,
            "commit_allowed": commit_allowed,
        },
        outcome="success",
        module=MODULE,
        commit=True,
    )

    meta = _IMPORT_TYPE_BY_ID[import_type]
    return {
        "batch_id": str(batch.id),
        "batch_code": batch.batch_code,
        "session_id": session_id,
        "import_type": import_type,
        "filename": filename,
        "sheet_name": selected_sheet,
        "other_sheets": other_sheets,
        "canonical_sheet": CANONICAL_SHEETS[import_type],
        "review_only": bool(meta.get("review_only")),
        "headers": headers,
        "column_bindings": bindings,
        "unmapped_columns": unmapped,
        "sensitive_columns_excluded": sensitive,
        "stats": stats,
        "first_10_records": first_10,
        "dependency_warnings": dep_warnings,
        "fidelity_note": fidelity_note,
        "block_commit": block_commit,
        "commit_allowed": commit_allowed,
        "confirm_required": True,
        "message": (
            f"Preview only — {len(mapped_rows)} rows from sheet '{selected_sheet}'. "
            + (fidelity_note or "Confirm to commit non-error rows.")
        ),
    }


# ---------------------------------------------------------------------------
# Commit
# ---------------------------------------------------------------------------


def commit_import(
    db: Session,
    *,
    actor: User,
    batch_id: UUID,
    confirm: bool,
    skip_duplicates: bool = True,
) -> dict[str, Any]:
    if not confirm:
        raise ProTrackValidationError(
            "Import requires confirm=true after reviewing the analyze preview."
        )
    batch = get_batch(db, batch_id)
    if batch.status != "preview":
        raise ProTrackValidationError(
            f"Batch {batch.batch_code} is status '{batch.status}', expected 'preview'."
        )
    if not batch.session_id:
        raise ProTrackValidationError("Batch has no analyze session to commit.")

    session = _load_session(batch.session_id)
    if session.get("batch_id") != str(batch.id):
        raise ProTrackValidationError("Session does not match this batch.")
    if session.get("stats", {}).get("id_blank_ratio", 0) > 0.5:
        # Re-check fidelity gate
        id_keys = {
            "assets": ("external_id", "asset_number"),
            "computers": ("asset_external_id", "asset_number"),
            "inventory": ("external_id",),
            "ip_addresses": ("ip_address",),
        }.get(batch.import_type, ())
        if id_keys:
            raise ProTrackValidationError(
                "Commit blocked by fidelity gate (>50% blank IDs). Re-analyze after fixing mapping."
            )

    import_type = batch.import_type
    rows: list[dict[str, Any]] = session.get("rows") or []
    customers = _customer_index(db)
    users = _user_index(db)

    imported = 0
    skipped = 0
    duplicated = 0
    errors: list[str] = []
    source_system_default = f"it_data_import:{session.get('filename') or batch.filename}"

    try:
        if import_type == "hardware_workbook":
            from app.services import it_hardware_workbook_import as hw

            existing = _existing_asset_keys(db)
            parent_links: list[tuple[Asset, str | None]] = []
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                # Re-apply defaults in case session stored pre-default fields
                row_fields = hw.apply_hardware_row_defaults(
                    {k: ("" if v is None else str(v)) for k, v in row.items() if not str(k).startswith("_")}
                )
                anum = (row_fields.get("asset_number") or row.get("asset_number") or "").strip()
                generated_number = bool(row.get("_generated_asset_number") or row_fields.get("_generated_asset_number"))
                if generated_number:
                    anum = ""  # allocate under lock
                norm = _norm_id(anum) if anum else ""
                st = (row_fields.get("service_tag") or "").strip()
                sn = (row_fields.get("serial_number") or "").strip()
                if skip_duplicates and anum and (
                    (st and _norm_id(st) in existing["service_tag"])
                    or (sn and _norm_id(sn) in existing["serial"])
                    or (norm in existing["asset_number"])
                    or (norm in existing["legacy"])
                ):
                    duplicated += 1
                    skipped += 1
                    continue

                purchased_by, owner_id, _ = _resolve_ownership(
                    row_fields.get("ownership_type") or "",
                    row_fields.get("owner_customer") or "",
                    customers,
                )
                if not (row_fields.get("ownership_type") or "").strip() and not (
                    row_fields.get("owner_customer") or ""
                ).strip():
                    purchased_by, owner_id = "unknown", None
                if purchased_by == "customer" and owner_id is None:
                    purchased_by = "unknown"

                status = _map_asset_status(
                    row_fields.get("current_status") or "",
                    "",
                    "",
                )
                if row_fields.get("_assignee_placeholder") == "1" and status == "available":
                    status = "available"
                type_label = row_fields.get("asset_type") or "OTHER"
                category_label = (row_fields.get("category") or "").strip()
                try:
                    at = _resolve_or_create_asset_type(db, type_label)
                    if category_label:
                        # Keep type category aligned with workbook Category master label
                        cat_code = re.sub(r"[^A-Za-z0-9]+", "_", category_label).strip("_").lower()[:40]
                        if cat_code:
                            at.category = cat_code if "hardware" in cat_code else (at.category or cat_code)
                            if "hardware" in category_label.casefold():
                                at.category = "peripheral" if "monitor" in type_label.casefold() or "mouse" in type_label.casefold() else "computer"
                            db.add(at)
                    make = (row_fields.get("make") or "").strip() or None
                    model = (row_fields.get("model_number") or row_fields.get("model") or "").strip() or None
                    asset = it_asset_service.create_asset(
                        db,
                        actor=actor,
                        asset_type_id=at.id,
                        asset_number=None if generated_number else anum,
                        legacy_asset_number=None if generated_number else (anum or None),
                        serial_number=sn or None,
                        make=make,
                        model=model,
                        location=(row_fields.get("location") or "").strip() or None,
                        notes=(row_fields.get("migration_note") or "").strip() or None,
                        description=(row_fields.get("description") or "").strip() or None,
                        service_tag=st or None,
                        purchased_by=purchased_by,
                        owner_customer_id=owner_id,
                        commit=False,
                    )
                    asset.status = status
                    asset.import_batch_id = batch.id
                    asset.source_system = source_system_default
                    asset.source_record_id = (anum or asset.asset_number or "")[:80]
                    asset.source_row = _parse_int(str(row.get("_source_row") or ""))
                    db.add(asset)
                    db.flush()

                    assignee_label = (row_fields.get("assigned_to") or "").strip()
                    assignee = _match_user(users, assignee_label) if assignee_label else None
                    if assignee and asset.status in {"available", "assigned"}:
                        try:
                            it_asset_service.assign_asset(
                                db,
                                asset,
                                user_id=assignee.id,
                                by_user=actor,
                                notes="IT Hardware Workbook import",
                                commit=False,
                            )
                        except Exception:  # noqa: BLE001
                            pass

                    hw.ensure_computer_and_ip(
                        db,
                        asset=asset,
                        fields=row_fields,
                        batch_id=batch.id,
                        source_system=source_system_default,
                        actor=actor,
                        ensure_slash24_network=_ensure_slash24_network,
                    )
                    parent_links.append(
                        (asset, (row_fields.get("parent_asset_number") or "").strip() or None)
                    )
                    imported += 1
                    if asset.asset_number:
                        existing["asset_number"].add(_norm_id(asset.asset_number))
                        existing["legacy"].add(_norm_id(asset.asset_number))
                    if sn:
                        existing["serial"].add(_norm_id(sn))
                    if st:
                        existing["service_tag"].add(_norm_id(st))
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Row {row.get('_source_row')}: {exc}")
                    skipped += 1

            for warn in hw.link_parent_assets(db, parent_links):
                errors.append(warn)

        elif import_type == "assets":
            existing = _existing_asset_keys(db)
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                anum = (row.get("asset_number") or row.get("external_id") or "").strip()
                generated_number = False
                if not anum:
                    # Missing number → allocate during create_asset (locked highest+1)
                    generated_number = True
                    at_preview = _resolve_or_create_asset_type(
                        db, row.get("asset_type") or "OTHER"
                    )
                    suggested = it_asset_service.preview_next_asset_number(db, at_preview)[
                        "asset_number"
                    ]
                    row["asset_number"] = suggested
                    row["_generated_asset_number"] = True
                    anum = ""  # force create_asset to allocate under lock
                norm = _norm_id(anum) if anum else ""
                st = (row.get("service_tag") or "").strip()
                sn = (row.get("serial_number") or "").strip()
                if skip_duplicates and (
                    (st and _norm_id(st) in existing["service_tag"])
                    or (sn and _norm_id(sn) in existing["serial"])
                    or (anum and norm in existing["asset_number"])
                    or (anum and norm in existing["legacy"])
                ):
                    duplicated += 1
                    skipped += 1
                    continue

                purchased_by, owner_id, _ = _resolve_ownership(
                    row.get("ownership_type") or "",
                    row.get("owner_customer") or "",
                    customers,
                )
                if purchased_by == "customer" and owner_id is None:
                    purchased_by = "unknown"

                cust_used = row.get("customer_used_for") or ""
                used_for = _match_customer(customers, cust_used) if cust_used else None
                supplier_name = (row.get("supplier") or "").strip()
                supplier = (
                    _get_or_create_supplier(
                        db,
                        supplier_name,
                        import_batch_id=batch.id,
                        source_system=row.get("source_system") or source_system_default,
                    )
                    if supplier_name
                    else None
                )

                status = _map_asset_status(
                    row.get("current_status") or "",
                    row.get("is_current_asset") or "",
                    row.get("returned_to_owner") or "",
                )
                type_label = row.get("asset_type") or "OTHER"
                desc = (row.get("description") or row.get("name") or "").strip() or None
                notes_parts = [
                    p
                    for p in [
                        (row.get("migration_note") or "").strip(),
                        (row.get("return_notes") or "").strip(),
                    ]
                    if p
                ]
                try:
                    at = _resolve_or_create_asset_type(db, type_label)
                    warranty = _parse_date(row.get("warranty_expiry") or "")
                    if warranty and warranty.year >= 2090:
                        warranty = None
                    asset = it_asset_service.create_asset(
                        db,
                        actor=actor,
                        asset_type_id=at.id,
                        asset_number=None if generated_number else anum,
                        legacy_asset_number=(
                            None
                            if generated_number
                            else ((row.get("legacy_asset_number") or anum).strip() or None)
                        ),
                        serial_number=sn or None,
                        make=(row.get("make") or row.get("manufacturer") or "").strip() or None,
                        model=(row.get("model_number") or row.get("model") or "").strip() or None,
                        purchase_date=_parse_date(row.get("purchase_date") or ""),
                        purchase_cost=_parse_decimal(row.get("purchase_value") or ""),
                        warranty_expiry=warranty,
                        location=(row.get("location") or row.get("team_or_department") or "").strip()
                        or None,
                        notes="; ".join(notes_parts) or None,
                        description=desc,
                        service_tag=st or None,
                        purchased_by=purchased_by,
                        owner_customer_id=owner_id,
                        customer_used_for_id=used_for.id if used_for else None,
                        supplier_id=supplier.id if supplier else None,
                        supplier_name=supplier_name or None,
                        invoice_number=(row.get("invoice_number") or "").strip() or None,
                        condition=(row.get("condition") or row.get("current_status") or "").strip()
                        or None,
                        commit=False,
                    )
                    asset.status = status
                    asset.import_batch_id = batch.id
                    asset.source_system = (
                        (row.get("source_system") or "").strip() or source_system_default
                    )
                    asset.source_record_id = (row.get("external_id") or anum)[:80]
                    src_row = _parse_int(str(row.get("source_row") or row.get("_source_row") or ""))
                    asset.source_row = src_row
                    db.add(asset)

                    assignee_label = (row.get("assigned_to") or "").strip()
                    assignee = _match_user(users, assignee_label) if assignee_label else None
                    if assignee and asset.status in {"available", "assigned"}:
                        try:
                            it_asset_service.assign_asset(
                                db,
                                asset,
                                user_id=assignee.id,
                                by_user=actor,
                                notes="IT Data Import assignment",
                                commit=False,
                            )
                        except Exception:  # noqa: BLE001
                            pass

                    imported += 1
                    existing["asset_number"].add(norm)
                    existing["legacy"].add(norm)
                    if sn:
                        existing["serial"].add(_norm_id(sn))
                    if st:
                        existing["service_tag"].add(_norm_id(st))
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Row {row.get('_source_row')}: {exc}")
                    skipped += 1

        elif import_type == "computers":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                asset = _find_asset_by_number(
                    db,
                    row.get("asset_number") or "",
                    row.get("asset_external_id") or "",
                )
                if asset is None:
                    errors.append(
                        f"Row {row.get('_source_row')}: asset not found for "
                        f"{row.get('asset_number') or row.get('asset_external_id')}"
                    )
                    skipped += 1
                    continue
                existing_c = db.scalar(select(Computer).where(Computer.asset_id == asset.id))
                if existing_c is not None:
                    duplicated += 1
                    skipped += 1
                    continue
                cname = (
                    (row.get("computer_name") or "").strip()
                    or asset.asset_number
                    or asset.legacy_asset_number
                    or "UNKNOWN"
                )[:40]
                name_clash = db.scalar(select(Computer).where(Computer.computer_name == cname))
                if name_clash is not None:
                    cname = f"{cname[:36]}-{str(uuid4())[:3]}"
                try:
                    db.add(
                        Computer(
                            id=uuid4(),
                            asset_id=asset.id,
                            computer_name=cname,
                            os=(row.get("operating_system") or "").strip() or None,
                            processor=(row.get("cpu") or "").strip() or None,
                            ram_gb=_parse_int(row.get("ram_gb") or ""),
                            mac_address=((row.get("mac_address") or "").strip() or None),
                            import_batch_id=batch.id,
                            source_system=(
                                (row.get("source_system") or "").strip() or source_system_default
                            ),
                        )
                    )
                    # Enrich asset make/model/serial when blank
                    if not asset.make and (row.get("manufacturer") or "").strip():
                        asset.make = (row.get("manufacturer") or "").strip()
                    if not asset.model and (row.get("model") or "").strip():
                        asset.model = (row.get("model") or "").strip()
                    if not asset.serial_number and (row.get("serial_number") or "").strip():
                        asset.serial_number = (row.get("serial_number") or "").strip()
                    if not asset.service_tag and (row.get("service_tag") or "").strip():
                        asset.service_tag = (row.get("service_tag") or "").strip()
                    db.add(asset)
                    imported += 1
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Row {row.get('_source_row')}: {exc}")
                    skipped += 1

        elif import_type == "ip_addresses":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                ip_raw = (row.get("ip_address") or "").strip()
                if not ip_raw:
                    skipped += 1
                    continue
                try:
                    ip_obj = ipaddress.IPv4Address(ip_raw)
                except ValueError:
                    errors.append(f"Row {row.get('_source_row')}: invalid IP {ip_raw}")
                    skipped += 1
                    continue
                existing_ip = db.scalar(select(IPAddress).where(IPAddress.address == str(ip_obj)))
                if existing_ip is not None:
                    duplicated += 1
                    skipped += 1
                    continue
                try:
                    network = _ensure_slash24_network(db, str(ip_obj))
                    ip_row = IPAddress(
                        id=uuid4(),
                        network_id=network.id,
                        address=str(ip_obj),
                        status="allocated",
                        allocation_type="static",
                        import_batch_id=batch.id,
                    )
                    db.add(ip_row)
                    db.flush()
                    asset = _find_asset_by_number(
                        db,
                        row.get("asset_number") or "",
                        row.get("asset_external_id") or "",
                    )
                    db.add(
                        IPAssignmentHistory(
                            id=uuid4(),
                            ip_address_id=ip_row.id,
                            assigned_to_asset_id=asset.id if asset else None,
                            assigned_by_user_id=actor.id,
                            hostname=(row.get("computer_name") or "").strip() or None,
                            assigned_date=date.today(),
                            notes=(
                                f"IT Data Import ({row.get('assignment_status') or 'Assigned'})"
                            ),
                        )
                    )
                    imported += 1
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Row {row.get('_source_row')}: {exc}")
                    skipped += 1

        elif import_type == "inventory":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                name = (row.get("item_name") or "").strip()
                if not name:
                    skipped += 1
                    continue
                existing_item = db.scalar(
                    select(InventoryItem).where(
                        func.lower(InventoryItem.name) == name.lower(),
                        InventoryItem.is_deleted.is_(False),
                    )
                )
                if existing_item is not None:
                    duplicated += 1
                    skipped += 1
                    continue
                purchased_by, owner_id, _ = _resolve_ownership(
                    row.get("ownership_type") or "",
                    row.get("owner_customer") or "",
                    customers,
                )
                if purchased_by == "customer" and owner_id is None:
                    purchased_by = "unknown"
                supplier_name = (row.get("supplier") or "").strip()
                supplier = (
                    _get_or_create_supplier(
                        db,
                        supplier_name,
                        import_batch_id=batch.id,
                        source_system=row.get("source_system") or source_system_default,
                    )
                    if supplier_name
                    else None
                )
                total = _parse_int(row.get("quantity") or "") or 0
                issued = _parse_int(row.get("quantity_in_use") or "") or 0
                db.add(
                    InventoryItem(
                        id=uuid4(),
                        name=name,
                        description=(row.get("description") or "").strip() or None,
                        category=(row.get("category") or "").strip() or None,
                        purchased_by=purchased_by,
                        owner_customer_id=owner_id,
                        supplier_id=supplier.id if supplier else None,
                        total_qty=total,
                        issued_qty=issued,
                        unit_cost=_parse_decimal(row.get("unit_or_total_value") or ""),
                        location=(row.get("storage_location") or "").strip() or None,
                        condition=(row.get("condition") or "").strip() or None,
                        status=(row.get("current_status") or "current").strip() or "current",
                        notes=None,
                        import_batch_id=batch.id,
                        source_system=(
                            (row.get("source_system") or "").strip() or source_system_default
                        ),
                        source_record_id=(row.get("external_id") or "")[:80] or None,
                    )
                )
                imported += 1

        elif import_type == "software":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                software = (row.get("software_name") or "").strip()
                if not software:
                    skipped += 1
                    continue
                catalog = db.scalar(
                    select(SoftwareCatalog).where(
                        func.lower(SoftwareCatalog.name) == software.lower()
                    )
                )
                if catalog is None:
                    catalog = SoftwareCatalog(id=uuid4(), name=software, is_active=True)
                    db.add(catalog)
                    db.flush()
                purchased_by, owner_id, _ = _resolve_ownership(
                    row.get("ownership_type") or row.get("purchased_by") or "",
                    row.get("owner_customer") or row.get("purchased_by") or "",
                    customers,
                )
                if purchased_by == "customer" and owner_id is None:
                    purchased_by = "unknown"
                seats = _parse_int(row.get("license_quantity") or "") or 1
                pool = SoftwareLicensePool(
                    id=uuid4(),
                    software_id=catalog.id,
                    purchased_by=purchased_by,
                    owner_customer_id=owner_id,
                    seat_count=seats,
                    expiry_date=_parse_date(row.get("license_valid_until") or ""),
                    renewal_mode=(row.get("renewal_mode") or "").strip() or None,
                    notes=(row.get("comments") or "").strip() or None,
                    import_batch_id=batch.id,
                    source_system=(
                        (row.get("source_system") or "").strip() or source_system_default
                    ),
                )
                db.add(pool)
                db.flush()
                assignee_label = (row.get("assigned_user") or "").strip()
                assignee = _match_user(users, assignee_label) if assignee_label else None
                if assignee:
                    db.add(
                        SoftwareAssignment(
                            id=uuid4(),
                            license_pool_id=pool.id,
                            user_id=assignee.id,
                            assigned_date=date.today(),
                            department=(row.get("department") or "").strip() or None,
                            import_batch_id=batch.id,
                        )
                    )
                imported += 1

        elif import_type == "user_accounts":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                name = (row.get("employee_name") or "").strip()
                email = (row.get("company_email") or "").strip()
                username = (row.get("username") or "").strip()
                user = (
                    _match_user(users, email)
                    or _match_user(users, name)
                    or _match_user(users, username)
                )
                if user is None:
                    skipped += 1
                    continue
                status_raw = (row.get("employee_status") or "").strip().lower()
                acct_status = "active"
                if any(x in status_raw for x in ("ex-", "inactive", "left", "terminated")):
                    acct_status = "deactivated"
                src = (row.get("source_system") or "").strip() or source_system_default
                created_any = False
                if username:
                    existing_acct = db.scalar(
                        select(ITUserAccount).where(
                            ITUserAccount.user_id == user.id,
                            ITUserAccount.account_type == "domain",
                            ITUserAccount.username == username,
                        )
                    )
                    if existing_acct:
                        duplicated += 1
                    else:
                        db.add(
                            ITUserAccount(
                                id=uuid4(),
                                user_id=user.id,
                                account_type="domain",
                                username=username,
                                display_name=name or None,
                                status=acct_status,
                                credential_status="migration_required",
                                notes=(
                                    "Legacy credential — Migration Required "
                                    "(passwords not imported)"
                                ),
                                import_batch_id=batch.id,
                                source_system=src,
                            )
                        )
                        created_any = True
                teams = (row.get("m365_teams_id") or "").strip()
                if teams:
                    existing_t = db.scalar(
                        select(ITUserAccount).where(
                            ITUserAccount.user_id == user.id,
                            ITUserAccount.account_type == "email",
                            ITUserAccount.username == teams,
                        )
                    )
                    if not existing_t:
                        db.add(
                            ITUserAccount(
                                id=uuid4(),
                                user_id=user.id,
                                account_type="email",
                                username=teams,
                                display_name=name or None,
                                status=acct_status,
                                credential_status="migration_required",
                                notes="Imported Teams/email identity — Migration Required",
                                import_batch_id=batch.id,
                                source_system=src,
                            )
                        )
                        created_any = True
                if created_any:
                    imported += 1
                else:
                    skipped += 1

        elif import_type == "suppliers":
            for row in rows:
                if row.get("_action") == "error":
                    skipped += 1
                    continue
                if row.get("_action") == "skip_duplicate" and skip_duplicates:
                    duplicated += 1
                    skipped += 1
                    continue
                name = (row.get("supplier_name") or "").strip()
                if not name:
                    skipped += 1
                    continue
                before = db.scalar(
                    select(func.count())
                    .select_from(ITSupplier)
                    .where(func.lower(ITSupplier.name) == name.lower())
                )
                _get_or_create_supplier(
                    db,
                    name,
                    import_batch_id=batch.id,
                    source_system=(row.get("source_system") or "").strip() or source_system_default,
                    extra=row,
                )
                if before and before > 0:
                    duplicated += 1
                    skipped += 1
                else:
                    imported += 1

        elif import_type == "migration_exceptions":
            for row in rows:
                if row.get("_action") in {"error", "skip"} and not (
                    row.get("exception_type") or row.get("issue")
                ):
                    skipped += 1
                    continue
                db.add(
                    ITMigrationException(
                        id=uuid4(),
                        import_batch_id=batch.id,
                        exception_type=(row.get("exception_type") or "unspecified")[:80],
                        source=(row.get("source") or "")[:200] or None,
                        source_row=str(row.get("source_row") or row.get("_source_row") or "")[:40]
                        or None,
                        identifier=(row.get("identifier") or "")[:120] or None,
                        issue=(row.get("issue") or None),
                        action=(row.get("action") or None),
                        status="open",
                    )
                )
                imported += 1

        batch.status = "committed"
        batch.success_count = imported
        batch.skipped_count = skipped
        batch.error_count = len(errors)
        batch.warning_count = duplicated
        batch.summary_json = json.dumps(
            {
                "imported": imported,
                "skipped": skipped,
                "duplicated": duplicated,
                "errors": errors[:100],
            }
        )
        db.add(batch)
        db.flush()

        log_activity(
            db,
            user=actor,
            entity_type=EntityType.it_settings,
            entity_id=batch.id,
            action=ActivityAction.it_data_import_committed,
            new_value={
                "batch_code": batch.batch_code,
                "import_type": import_type,
                "imported": imported,
                "skipped": skipped,
                "duplicated": duplicated,
            },
            outcome="success",
            module=MODULE,
            commit=False,
        )
        db.commit()
        _delete_session(batch.session_id)
        batch.session_id = None
        db.add(batch)
        db.commit()

    except Exception:
        db.rollback()
        batch = get_batch(db, batch_id)
        batch.status = "failed"
        db.add(batch)
        db.commit()
        raise

    return {
        "batch_id": str(batch.id),
        "batch_code": batch.batch_code,
        "import_type": import_type,
        "filename": batch.filename,
        "status": batch.status,
        "imported": imported,
        "skipped": skipped,
        "duplicated": duplicated,
        "errors": errors[:100],
        "confirm_required": False,
        "message": (
            f"Committed {imported} row(s) for {import_type} "
            f"(skipped {skipped}, duplicates {duplicated}). "
            "Next import type is not started automatically."
        ),
    }


def _ensure_slash24_network(db: Session, ip_str: str) -> Network:
    """Find a network containing the IP, or create a /24 without auto-filling hosts."""
    ip = ipaddress.IPv4Address(ip_str)
    for net in db.scalars(select(Network)).all():
        try:
            if ip in ipaddress.IPv4Network(net.cidr, strict=False):
                return net
        except ValueError:
            continue
    cidr_net = ipaddress.IPv4Network(f"{ip}/24", strict=False)
    cidr = str(cidr_net)
    existing = db.scalar(select(Network).where(Network.cidr == cidr))
    if existing is not None:
        return existing
    network = Network(
        id=uuid4(),
        name=f"Import {cidr}",
        cidr=cidr,
        description="Auto-created during IT Data Import (hosts not pre-generated).",
        is_active=True,
    )
    db.add(network)
    db.flush()
    return network


# ---------------------------------------------------------------------------
# Rollback
# ---------------------------------------------------------------------------


def rollback_batch(
    db: Session,
    *,
    actor: User,
    batch_id: UUID,
    confirm: bool,
) -> dict[str, Any]:
    if not confirm:
        raise ProTrackValidationError("Rollback requires confirm=true.")
    batch = get_batch(db, batch_id)
    if batch.status not in {"committed", "failed"}:
        raise ProTrackValidationError(
            f"Only committed/failed batches can be rolled back (status={batch.status})."
        )

    deleted: dict[str, int] = {}

    # Software children first
    n = db.execute(
        delete(SoftwareAssignment).where(SoftwareAssignment.import_batch_id == batch.id)
    ).rowcount
    deleted["software_assignments"] = int(n or 0)

    n = db.execute(
        delete(SoftwareLicensePool).where(SoftwareLicensePool.import_batch_id == batch.id)
    ).rowcount
    deleted["software_license_pools"] = int(n or 0)

    # IP history for IPs in this batch
    ip_ids = list(
        db.scalars(select(IPAddress.id).where(IPAddress.import_batch_id == batch.id)).all()
    )
    if ip_ids:
        n = db.execute(
            delete(IPAssignmentHistory).where(IPAssignmentHistory.ip_address_id.in_(ip_ids))
        ).rowcount
        deleted["ip_assignment_history"] = int(n or 0)
    else:
        deleted["ip_assignment_history"] = 0

    n = db.execute(delete(IPAddress).where(IPAddress.import_batch_id == batch.id)).rowcount
    deleted["ip_addresses"] = int(n or 0)

    n = db.execute(
        delete(ITUserAccount).where(ITUserAccount.import_batch_id == batch.id)
    ).rowcount
    deleted["it_user_accounts"] = int(n or 0)

    n = db.execute(delete(Computer).where(Computer.import_batch_id == batch.id)).rowcount
    deleted["computers"] = int(n or 0)

    n = db.execute(
        delete(InventoryItem).where(InventoryItem.import_batch_id == batch.id)
    ).rowcount
    deleted["inventory_items"] = int(n or 0)

    n = db.execute(
        delete(ITMigrationException).where(ITMigrationException.import_batch_id == batch.id)
    ).rowcount
    deleted["migration_exceptions"] = int(n or 0)

    # Assets: remove FK children for assets in this batch, then assets
    asset_ids = list(
        db.scalars(select(Asset.id).where(Asset.import_batch_id == batch.id)).all()
    )
    if asset_ids:
        # Block if computers from other batches still reference these assets
        other_computers = db.scalar(
            select(func.count())
            .select_from(Computer)
            .where(
                Computer.asset_id.in_(asset_ids),
                or_(
                    Computer.import_batch_id.is_(None),
                    Computer.import_batch_id != batch.id,
                ),
            )
        )
        if other_computers and other_computers > 0:
            raise ProTrackValidationError(
                f"Cannot rollback assets: {other_computers} computer(s) from other batches "
                "still reference these assets. Rollback computers first."
            )
        db.execute(
            delete(IPAssignmentHistory).where(
                IPAssignmentHistory.assigned_to_asset_id.in_(asset_ids)
            )
        )
        db.execute(delete(AssetAssignment).where(AssetAssignment.asset_id.in_(asset_ids)))
        db.execute(
            delete(AssetCustomerReturn).where(AssetCustomerReturn.asset_id.in_(asset_ids))
        )
        db.execute(delete(Computer).where(Computer.asset_id.in_(asset_ids)))
        n = db.execute(delete(Asset).where(Asset.import_batch_id == batch.id)).rowcount
        deleted["assets"] = int(n or 0)
    else:
        deleted["assets"] = 0

    n = db.execute(delete(ITSupplier).where(ITSupplier.import_batch_id == batch.id)).rowcount
    deleted["it_suppliers"] = int(n or 0)

    # Prefer leave networks (even Import-named empties).

    batch.status = "rolled_back"
    batch.summary_json = json.dumps({"rollback_deleted": deleted})
    db.add(batch)

    log_activity(
        db,
        user=actor,
        entity_type=EntityType.it_settings,
        entity_id=batch.id,
        action=ActivityAction.it_data_import_rolled_back,
        new_value={"batch_code": batch.batch_code, "deleted": deleted},
        outcome="success",
        module=MODULE,
        commit=False,
    )
    db.commit()

    if batch.session_id:
        _delete_session(batch.session_id)

    return {
        "batch_id": str(batch.id),
        "batch_code": batch.batch_code,
        "import_type": batch.import_type,
        "status": batch.status,
        "deleted": deleted,
        "message": f"Rolled back batch {batch.batch_code}.",
    }

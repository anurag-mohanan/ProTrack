"""Excel workbook generator for engineering management reports."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

from app.schemas.reporting import EngineeringReportPayload
from app.services.reporting.excel.letterhead import write_report_letterhead, set_print_layout
from app.services.reporting.excel.styles import (
    DANGER,
    KPI_FILL,
    KPI_LABEL_FONT,
    KPI_VALUE_FONT,
    SUCCESS,
    WARNING,
    autofit_columns,
    freeze_and_filter,
    style_body_rows,
    style_header_row,
)


def generate_engineering_excel(payload: EngineeringReportPayload) -> bytes:
    workbook = Workbook()
    workbook.properties.title = f"{payload.company_name} Engineering Report"
    workbook.properties.subject = payload.period.label
    workbook.properties.company = payload.company_name
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)

    _write_executive_sheet(workbook, payload)
    _write_designer_productivity(workbook, payload)
    _write_designer_tool_breakdown(workbook, payload)
    _write_tool_hours(workbook, payload)
    _write_customer_summary(workbook, payload)
    _write_team_summary(workbook, payload)
    _write_function_hours(workbook, payload)
    _write_np_analysis(workbook, payload)
    _write_leave_analysis(workbook, payload)
    _write_quoted_vs_actual(workbook, payload)
    _write_project_performance(workbook, payload)
    _write_detailed_entries(workbook, payload)
    _write_charts_sheet(workbook, payload)
    _write_ai_insights(workbook, payload)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _write_executive_sheet(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    sheet = workbook.active
    sheet.title = "Executive Summary"
    next_row = write_report_letterhead(
        sheet,
        company_name=payload.company_name,
        report_title="Engineering Management Report",
        period_label=payload.period.label,
        extra_lines=["Executive KPI overview"],
        col_span=8,
    )

    row = next_row
    col = 1
    for index, kpi in enumerate(payload.executive.kpis):
        if index and index % 3 == 0:
            row += 3
            col = 1
        cell_label = sheet.cell(row=row, column=col, value=kpi.label)
        cell_value = sheet.cell(row=row + 1, column=col, value=kpi.value)
        cell_label.font = KPI_LABEL_FONT
        cell_value.font = KPI_VALUE_FONT
        cell_label.fill = KPI_FILL
        cell_value.fill = KPI_FILL
        cell_label.alignment = Alignment(wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        sheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        col += 3

    autofit_columns(sheet, min_width=14)
    set_print_layout(sheet, company_name=payload.company_name)


def _write_table_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    *,
    company_name: str = "",
    period_label: str | None = None,
    variance_col: int | None = None,
) -> None:
    sheet = workbook.create_sheet(title[:31])
    next_row = write_report_letterhead(
        sheet,
        company_name=company_name,
        report_title=title,
        period_label=period_label,
        col_span=max(len(headers), 6),
    )
    header_row = next_row
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col, value=header)
    style_header_row(sheet, header_row, len(headers))

    data_start = header_row + 1
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=data_start + offset, column=col, value=value)
    if rows:
        style_body_rows(sheet, data_start, data_start + len(rows) - 1, len(headers))
        freeze_and_filter(sheet, header_row, len(headers))

    if variance_col and rows:
        from openpyxl.utils import get_column_letter

        letter = get_column_letter(variance_col)
        rng = f"{letter}{data_start}:{letter}{data_start + len(rows) - 1}"
        sheet.conditional_formatting.add(
            rng,
            CellIsRule(operator="lessThan", formula=["0"], fill=_fill(SUCCESS)),
        )
        sheet.conditional_formatting.add(
            rng,
            CellIsRule(operator="between", formula=["0", "10"], fill=_fill(WARNING)),
        )
        sheet.conditional_formatting.add(
            rng,
            CellIsRule(operator="greaterThan", formula=["10"], fill=_fill(DANGER)),
        )

    autofit_columns(sheet)


def _write_designer_productivity(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Designer",
        "Team",
        "Productive",
        "Non-Productive",
        "Leave Days",
        "Total Hours",
        "Billable %",
        "Utilization %",
        "Projects",
        "Customers",
    ]
    rows = [
        [
            row.designer_name,
            row.team_name or "",
            float(row.productive_hours),
            float(row.non_productive_hours),
            float(row.leave_days),
            float(row.total_hours),
            float(row.billable_percent),
            float(row.utilization_percent or 0),
            row.project_count,
            row.customer_count,
        ]
        for row in payload.designer_productivity
    ]
    _write_table_sheet(workbook, "Designer Productivity", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_designer_tool_breakdown(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Designer",
        "Tool Number",
        "Customer",
        "Design",
        "Surfacing",
        "Review",
        "BOM",
        "Meetings",
        "NP",
        "Other",
        "Total",
    ]
    rows = [
        [
            row.designer_name,
            row.tool_number,
            row.customer_name,
            float(row.design_hours),
            float(row.surfacing_hours),
            float(row.review_hours),
            float(row.bom_hours),
            float(row.meeting_hours),
            float(row.np_hours),
            float(row.other_hours),
            float(row.total_hours),
        ]
        for row in payload.designer_tool_breakdown
    ]
    _write_table_sheet(workbook, "Designer Tool Breakdown", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_tool_hours(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Tool Number",
        "Customer",
        "Part Description",
        "Design Leader",
        "Designer",
        "Surfacer",
        "Quoted",
        "Actual",
        "Variance",
        "Variance %",
        "Completion %",
        "Stage",
        "Status",
    ]
    rows = [
        [
            row.tool_number,
            row.customer_name,
            row.part_description,
            row.design_leader_name or "",
            row.designer_name or "",
            row.surfacer_name or "",
            float(row.quoted_hours),
            float(row.actual_hours),
            float(row.variance_hours),
            float(row.variance_percent),
            float(row.completion_percent),
            row.project_stage.value,
            row.execution_status.value,
        ]
        for row in payload.tool_hours
    ]
    _write_table_sheet(workbook, "Tool Hours", headers, rows, variance_col=9, company_name=payload.company_name, period_label=payload.period.label)


def _write_customer_summary(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Customer",
        "Projects",
        "Productive Hours",
        "NP Hours",
        "Total Hours",
        "Designers",
        "Avg Hours / Project",
    ]
    rows = [
        [
            row.customer_name,
            row.project_count,
            float(row.productive_hours),
            float(row.np_hours),
            float(row.total_hours),
            row.designer_count,
            float(row.avg_hours_per_project),
        ]
        for row in payload.customer_summary
    ]
    _write_table_sheet(workbook, "Customer Summary", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_team_summary(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Team",
        "Designers",
        "Projects",
        "Productive",
        "NP",
        "Leave Days",
        "Total",
        "Utilization %",
    ]
    rows = [
        [
            row.team_name,
            row.designer_count,
            row.project_count,
            float(row.productive_hours),
            float(row.np_hours),
            float(row.leave_days),
            float(row.total_hours),
            float(row.utilization_percent or 0),
        ]
        for row in payload.team_summary
    ]
    _write_table_sheet(workbook, "Team Summary", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_function_hours(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = ["Function", "Hours", "%"]
    rows = [[row.function_group, float(row.hours), float(row.percent)] for row in payload.function_hours]
    _write_table_sheet(workbook, "Function Hours", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_np_analysis(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = ["Code", "Description", "Hours", "%"]
    rows = [[row.code, row.description, float(row.hours), float(row.percent)] for row in payload.np_analysis]
    _write_table_sheet(workbook, "Non-Productive Analysis", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_leave_analysis(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = ["Designer", "Leave Days", "Leave Hours"]
    rows = [
        [row.designer_name, float(row.leave_days), float(row.leave_hours)]
        for row in payload.leave_analysis
    ]
    _write_table_sheet(workbook, "Leave Analysis", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_quoted_vs_actual(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Tool",
        "Customer",
        "Quoted",
        "Original Hours",
        "Additional Work",
        "Rework",
        "Customer Change",
        "Internal Correction",
        "Post-Completion Total",
        "Variance",
        "Variance %",
        "Completion %",
        "Health",
        "Late Milestones",
    ]
    rows = [
        [
            row.tool_number,
            row.customer_name,
            float(row.quoted_hours),
            float(row.original_hours or row.actual_hours),
            float(row.additional_work_hours),
            float(row.rework_hours),
            float(row.customer_change_hours),
            float(row.internal_correction_hours),
            float(row.post_completion_hours),
            float(row.variance_hours),
            float(row.variance_percent),
            float(row.completion_percent),
            row.health.value if row.health else "",
            row.late_milestones,
        ]
        for row in payload.quoted_vs_actual
    ]
    _write_table_sheet(workbook, "Quoted vs Actual", headers, rows, variance_col=11, company_name=payload.company_name, period_label=payload.period.label)


def _write_project_performance(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Tool",
        "Customer",
        "Designer",
        "Surfacer",
        "Stage",
        "Quoted",
        "Actual",
        "Milestone %",
        "Health",
        "Predicted Finish",
    ]
    rows = [
        [
            row.tool_number,
            row.customer_name,
            row.designer_name or "",
            row.surfacer_name or "",
            row.project_stage.value,
            float(row.quoted_hours),
            float(row.actual_hours),
            float(row.milestone_completion_percent),
            row.health.value if row.health else "",
            row.predicted_finish.isoformat() if row.predicted_finish else "",
        ]
        for row in payload.project_performance
    ]
    _write_table_sheet(workbook, "Project Performance", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_detailed_entries(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    headers = [
        "Date",
        "Designer",
        "Team",
        "Customer",
        "Tool",
        "Task",
        "Hours",
        "Billable",
        "Category",
        "Post-Completion Type",
        "Notes",
    ]
    rows = [
        [
            row.entry_date.isoformat(),
            row.designer_name,
            row.team_name or "",
            row.customer_name or "",
            row.tool_number or "",
            row.task_name or "",
            float(row.hours),
            "Yes" if row.is_billable else "No",
            row.category,
            row.post_completion_type or "",
            row.notes or "",
        ]
        for row in payload.detailed_entries
    ]
    _write_table_sheet(workbook, "Detailed Entries", headers, rows, company_name=payload.company_name, period_label=payload.period.label)


def _write_charts_sheet(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    sheet = workbook.create_sheet("Charts")
    next_row = write_report_letterhead(
        sheet,
        company_name=payload.company_name,
        report_title="Report Charts",
        period_label=payload.period.label,
        col_span=8,
    )
    row = next_row
    for chart_data in payload.charts:
        sheet.cell(row=row, column=1, value=chart_data.title)
        sheet.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        sheet.cell(row=row, column=1, value="Label")
        sheet.cell(row=row, column=2, value="Value")
        row += 1
        start = row
        for label, value in zip(chart_data.labels, chart_data.values, strict=False):
            sheet.cell(row=row, column=1, value=label)
            sheet.cell(row=row, column=2, value=value)
            row += 1
        end = row - 1

        if len(chart_data.labels) <= 6:
            chart = PieChart()
            chart.title = chart_data.title
            data = Reference(sheet, min_col=2, min_row=start, max_row=end)
            labels = Reference(sheet, min_col=1, min_row=start, max_row=end)
            chart.add_data(data, titles_from_data=False)
            chart.set_categories(labels)
            sheet.add_chart(chart, f"D{start}")
        else:
            chart = BarChart()
            chart.type = "col"
            chart.title = chart_data.title
            data = Reference(sheet, min_col=2, min_row=start, max_row=end)
            labels = Reference(sheet, min_col=1, min_row=start, max_row=end)
            chart.add_data(data, titles_from_data=False)
            chart.set_categories(labels)
            sheet.add_chart(chart, f"D{start}")
        row += 16
    set_print_layout(sheet, company_name=payload.company_name)


def _write_ai_insights(workbook: Workbook, payload: EngineeringReportPayload) -> None:
    sheet = workbook.create_sheet("AI Insights")
    next_row = write_report_letterhead(
        sheet,
        company_name=payload.company_name,
        report_title="AI Engineering Insights",
        period_label=payload.period.label,
        extra_lines=["Generated from live database analytics"],
        col_span=6,
    )
    row = next_row
    for insight in payload.ai_insights:
        sheet.cell(row=row, column=1, value=f"• {insight}")
        sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
        row += 1
    sheet.column_dimensions["A"].width = 100
    set_print_layout(sheet, landscape=False, company_name=payload.company_name)

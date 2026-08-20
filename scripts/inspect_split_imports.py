"""Inspect split IT import workbooks: sheets, headers, first rows."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1] / "docs" / "migration-sources" / "split"


def main() -> None:
    for path in sorted(ROOT.glob("*.xlsx")):
        print("=" * 72)
        print(path.name)
        wb = load_workbook(path, read_only=True, data_only=True)
        print(f"sheets: {wb.sheetnames}")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                print(f"  [{sheet_name}] EMPTY")
                continue
            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            data_rows = [r for r in rows[1:] if any(c is not None and str(c).strip() for c in r)]
            print(f"  [{sheet_name}] headers={len(headers)} data_rows={len(data_rows)}")
            print(f"    headers: {headers}")
            for i, r in enumerate(data_rows[:3], 1):
                sample = {
                    headers[j]: (r[j] if j < len(r) else None)
                    for j in range(len(headers))
                    if headers[j]
                }
                # truncate long values
                compact = {
                    k: (str(v)[:60] if v is not None else None)
                    for k, v in list(sample.items())[:12]
                }
                print(f"    row{i}: {compact}")
        wb.close()


if __name__ == "__main__":
    main()

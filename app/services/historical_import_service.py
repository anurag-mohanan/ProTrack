"""Historical project import from Prosohm Tool Tracking Excel workbooks."""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import gettempdir
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.permissions import DESIGNER, PROJECT_STAFF_ROLES
from app.core.security import hash_password
from app.crud.project import DEFAULT_PROJECT_MILESTONES
from app.models.enums import MilestoneStatus, ProjectHealth, ProjectStatus
from app.models.models import Contact, Customer, Milestone, Project, ProjectType, Role, Stream, User
from app.services.project_template_service import (
    create_milestones_from_template,
    resolve_template_for_import,
)
from app.schemas.historical_import import (
    DuplicateAction,
    ImportRowPreview,
    ImportRowStatus,
    ImportSummary,
    ImportUploadResponse,
)
from app.services.project_calculation_service import calculate_project_health

UPLOAD_DIR = Path(gettempdir()) / "protrack_imports"
DEFAULT_IMPORT_PASSWORD = "Password@123"

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "tool_number": ("tool no.", "tool no", "tool number", "tool #", "tool"),
    "customer": ("customer", "customer name"),
    "designer": ("designer",),
    "surfacer": ("surfacer",),
    "quoted_hours": ("quoted hours", "quoted hrs", "quote hours"),
    "actual_hours": (
        "actual design hours",
        "actual hours",
        "actual hrs",
        "design hours",
    ),
    "design_phase": ("design phase", "phase", "current milestone", "milestone"),
    "progress": ("progress", "progress %", "progress percent"),
    "status": ("eng status", "engineering status", "status", "project status"),
    "part_description": ("part description", "part desc", "description", "part"),
    "due_date": ("due date", "due", "target date"),
    "design_leader": ("design leader", "design lead", "leader"),
    "code": ("project code", "code", "job code"),
}

STATUS_MAP: dict[str, ProjectStatus] = {
    "not started": ProjectStatus.not_started,
    "not_started": ProjectStatus.not_started,
    "new": ProjectStatus.not_started,
    "in progress": ProjectStatus.in_progress,
    "in_progress": ProjectStatus.in_progress,
    "active": ProjectStatus.in_progress,
    "wip": ProjectStatus.in_progress,
    "working": ProjectStatus.in_progress,
    "waiting for customer": ProjectStatus.waiting_for_customer,
    "waiting_for_customer": ProjectStatus.waiting_for_customer,
    "waiting": ProjectStatus.waiting_for_customer,
    "on hold": ProjectStatus.waiting_for_customer,
    "hold": ProjectStatus.waiting_for_customer,
    "completed": ProjectStatus.completed,
    "complete": ProjectStatus.completed,
    "done": ProjectStatus.completed,
    "closed": ProjectStatus.completed,
}

MILESTONE_ALIASES: dict[str, str] = {
    "feasibility": "Feasibility",
    "blockout": "Blockout",
    "block out": "Blockout",
    "roughing": "Roughing",
    "intermediate review": "Intermediate Review",
    "intermediate": "Intermediate Review",
    "final review": "Final Review",
    "final": "Final Review",
    "file release": "File Release",
    "file": "File Release",
    "bom release": "BOM Release",
    "bom": "BOM Release",
}


@dataclass
class ParsedImportRow:
    row_number: int
    tool_number: str | None = None
    customer: str | None = None
    designer: str | None = None
    surfacer: str | None = None
    quoted_hours: Decimal | None = None
    actual_hours: Decimal | None = None
    design_phase: str | None = None
    progress: Decimal | None = None
    status: ProjectStatus | None = None
    part_description: str | None = None
    due_date: date | None = None
    design_leader: str | None = None
    code: str | None = None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def _normalize_key(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", value.strip().lower())


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1].strip()
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_status(value: Any) -> ProjectStatus | None:
    text = _normalize_key(_cell_text(value))
    if not text:
        return None
    return STATUS_MAP.get(text)


def _resolve_milestone_name(value: str | None) -> str | None:
    if not value:
        return None
    normalized = _normalize_key(value)
    if normalized in MILESTONE_ALIASES:
        return MILESTONE_ALIASES[normalized]
    for canonical in DEFAULT_PROJECT_MILESTONES:
        if _normalize_key(canonical) == normalized:
            return canonical
    for alias, canonical in MILESTONE_ALIASES.items():
        if alias in normalized or normalized in alias:
            return canonical
    return None


def _detect_header_map(sheet) -> tuple[int, dict[str, int]]:
    for row_idx in range(1, min(sheet.max_row, 25) + 1):
        row_values = {
            col_idx: _normalize_header(sheet.cell(row_idx, col_idx).value)
            for col_idx in range(1, sheet.max_column + 1)
        }
        header_map: dict[str, int] = {}
        for field_name, aliases in COLUMN_ALIASES.items():
            for col_idx, header in row_values.items():
                if not header:
                    continue
                if header in aliases or any(alias in header for alias in aliases):
                    header_map[field_name] = col_idx
                    break
        if "tool_number" in header_map and "customer" in header_map:
            return row_idx, header_map
    raise ValueError(
        "Could not locate a header row. Expected columns such as Tool No. and Customer."
    )


def _read_cell(sheet, row_idx: int, header_map: dict[str, int], field_name: str) -> Any:
    col_idx = header_map.get(field_name)
    if col_idx is None:
        return None
    return sheet.cell(row_idx, col_idx).value


def _is_blank_row(sheet, row_idx: int, header_map: dict[str, int]) -> bool:
    for col_idx in header_map.values():
        value = sheet.cell(row_idx, col_idx).value
        if value is not None and str(value).strip():
            return False
    return True


def parse_workbook(file_path: Path) -> list[ParsedImportRow]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Workbook has no active worksheet.")

        header_row, header_map = _detect_header_map(sheet)
        parsed_rows: list[ParsedImportRow] = []

        for row_idx in range(header_row + 1, sheet.max_row + 1):
            if _is_blank_row(sheet, row_idx, header_map):
                continue

            row = ParsedImportRow(row_number=row_idx)
            row.tool_number = _cell_text(_read_cell(sheet, row_idx, header_map, "tool_number"))
            row.customer = _cell_text(_read_cell(sheet, row_idx, header_map, "customer"))
            row.designer = _cell_text(_read_cell(sheet, row_idx, header_map, "designer"))
            row.surfacer = _cell_text(_read_cell(sheet, row_idx, header_map, "surfacer"))
            row.design_phase = _cell_text(
                _read_cell(sheet, row_idx, header_map, "design_phase")
            )
            row.part_description = _cell_text(
                _read_cell(sheet, row_idx, header_map, "part_description")
            )
            row.design_leader = _cell_text(
                _read_cell(sheet, row_idx, header_map, "design_leader")
            )
            row.code = _cell_text(_read_cell(sheet, row_idx, header_map, "code"))

            quoted = _parse_decimal(_read_cell(sheet, row_idx, header_map, "quoted_hours"))
            actual = _parse_decimal(_read_cell(sheet, row_idx, header_map, "actual_hours"))
            progress = _parse_decimal(_read_cell(sheet, row_idx, header_map, "progress"))
            due_date = _parse_date(_read_cell(sheet, row_idx, header_map, "due_date"))
            status = _parse_status(_read_cell(sheet, row_idx, header_map, "status"))

            if quoted is None:
                row.errors.append("Quoted Hours is missing or invalid.")
            else:
                row.quoted_hours = quoted

            if actual is not None:
                row.actual_hours = actual

            if progress is not None:
                row.progress = progress

            if due_date is not None:
                row.due_date = due_date
            elif due_date is False:
                row.errors.append("Due Date is invalid.")

            invalid_date = _read_cell(sheet, row_idx, header_map, "due_date")
            if invalid_date not in (None, "") and row.due_date is None:
                row.errors.append("Due Date is invalid.")

            row.status = status

            if not row.tool_number:
                row.errors.append("Tool Number is required.")
            if not row.customer:
                row.errors.append("Customer is required.")

            parsed_rows.append(row)
        return parsed_rows
    finally:
        workbook.close()


def save_upload(upload_id: str, file_name: str, content: bytes) -> Path:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(file_name).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Only .xlsx and .xlsm files are supported.")
    path = UPLOAD_DIR / f"{upload_id}{suffix}"
    path.write_bytes(content)
    meta_path = UPLOAD_DIR / f"{upload_id}.meta"
    meta_path.write_text(file_name, encoding="utf-8")
    return path


def get_upload_path(upload_id: str) -> Path:
    for suffix in (".xlsx", ".xlsm"):
        path = UPLOAD_DIR / f"{upload_id}{suffix}"
        if path.exists():
            return path
    raise FileNotFoundError("Upload not found or expired.")


def _preview_from_row(
    row: ParsedImportRow,
    *,
    existing_project_id: uuid.UUID | None = None,
) -> ImportRowPreview:
    if row.errors:
        status = ImportRowStatus.error
    elif existing_project_id is not None:
        status = ImportRowStatus.duplicate
    else:
        status = ImportRowStatus.ready

    messages = list(row.errors)
    messages.extend(row.warnings)

    return ImportRowPreview(
        row_number=row.row_number,
        tool_number=row.tool_number,
        customer=row.customer,
        designer=row.designer,
        surfacer=row.surfacer,
        quoted_hours=str(row.quoted_hours) if row.quoted_hours is not None else None,
        actual_hours=str(row.actual_hours) if row.actual_hours is not None else None,
        design_phase=row.design_phase,
        progress=str(row.progress) if row.progress is not None else None,
        status=row.status.value if row.status is not None else None,
        status_label=status,
        messages=messages,
        existing_project_id=existing_project_id,
    )


def _find_project_by_tool_number(db: Session, tool_number: str) -> Project | None:
    normalized = tool_number.strip().lower()
    projects = db.scalars(select(Project)).all()
    for project in projects:
        if project.tool_number.strip().lower() == normalized:
            return project
    return None


def analyze_upload(db: Session, upload_id: str) -> ImportUploadResponse:
    path = get_upload_path(upload_id)
    rows = parse_workbook(path)
    preview: list[ImportRowPreview] = []

    ready_rows = error_rows = duplicate_rows = 0
    missing_customer_rows = missing_designer_rows = 0

    for row in rows:
        existing = (
            _find_project_by_tool_number(db, row.tool_number)
            if row.tool_number and not row.errors
            else None
        )
        item = _preview_from_row(row, existing_project_id=existing.id if existing else None)
        preview.append(item)

        if item.status_label == ImportRowStatus.error:
            error_rows += 1
        elif item.status_label == ImportRowStatus.duplicate:
            duplicate_rows += 1
        else:
            ready_rows += 1

        if not row.customer and "Customer is required." not in row.errors:
            missing_customer_rows += 1
        if row.customer and not row.designer:
            missing_designer_rows += 1

    meta_path = UPLOAD_DIR / f"{upload_id}.meta"
    file_name = meta_path.read_text(encoding="utf-8") if meta_path.exists() else path.name

    return ImportUploadResponse(
        upload_id=upload_id,
        file_name=file_name,
        total_rows=len(rows),
        ready_rows=ready_rows,
        error_rows=error_rows,
        duplicate_rows=duplicate_rows,
        missing_customer_rows=missing_customer_rows,
        missing_designer_rows=missing_designer_rows,
        preview=preview,
    )


def _customer_code(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]", "", name.upper())
    return (cleaned[:8] or "CUST")[:20]


def _resolve_or_create_customer(db: Session, name: str, summary: ImportSummary) -> Customer:
    normalized = _normalize_key(name)
    customers = db.scalars(select(Customer).where(Customer.is_active.is_(True))).all()
    for customer in customers:
        if _normalize_key(customer.name) == normalized:
            return customer

    code_base = _customer_code(name)
    code = code_base
    suffix = 1
    while db.scalar(select(Customer.id).where(Customer.code == code)):
        code = f"{code_base}{suffix}"[:20]
        suffix += 1

    customer = Customer(name=name.strip(), code=code, is_active=True)
    db.add(customer)
    db.flush()

    contact = Contact(
        customer_id=customer.id,
        first_name="Import",
        last_name="Contact",
        email=None,
        is_primary=True,
    )
    db.add(contact)
    db.flush()
    summary.customers_created += 1
    return customer


def _split_name(full_name: str) -> tuple[str, str]:
    parts = full_name.strip().split()
    if not parts:
        return "Imported", "User"
    if len(parts) == 1:
        return parts[0], parts[0]
    return parts[0], " ".join(parts[1:])


def _unique_email(db: Session, first_name: str, last_name: str) -> str:
    base = re.sub(r"[^a-z0-9]", "", f"{first_name}{last_name}".lower()) or "importuser"
    email = f"{base}@prosohm.com"
    suffix = 1
    while db.scalar(select(User.id).where(User.email == email)):
        email = f"{base}{suffix}@prosohm.com"
        suffix += 1
    return email


def _resolve_or_create_user(
    db: Session,
    name: str | None,
    *,
    summary: ImportSummary,
    default_role_name: str = DESIGNER,
) -> User | None:
    if not name or not name.strip():
        return None

    normalized = _normalize_key(name)
    users = db.scalars(
        select(User)
        .join(Role, User.role_id == Role.id)
        .where(User.is_active.is_(True), Role.name.in_(PROJECT_STAFF_ROLES))
    ).all()

    for user in users:
        full = _normalize_key(f"{user.first_name} {user.last_name}")
        first = _normalize_key(user.first_name)
        last = _normalize_key(user.last_name)
        if normalized in {full, first, last, _normalize_key(user.email.split("@")[0])}:
            return user

    role = db.scalar(select(Role).where(Role.name == default_role_name))
    if role is None:
        role = db.scalar(select(Role).where(Role.name == DESIGNER))
    if role is None:
        raise ValueError(f"Role {default_role_name} is not configured.")

    first_name, last_name = _split_name(name.strip())
    email = _unique_email(db, first_name, last_name)
    user = User(
        role_id=role.id,
        email=email,
        password_hash=hash_password(DEFAULT_IMPORT_PASSWORD),
        first_name=first_name,
        last_name=last_name,
        is_active=True,
    )
    db.add(user)
    db.flush()
    summary.users_created += 1
    return user


def _resolve_design_leader(db: Session, name: str | None) -> User:
    if name and name.strip():
        leader = db.scalars(
            select(User)
            .join(Role, User.role_id == Role.id)
            .where(Role.name == "Design Leader", User.is_active.is_(True))
        ).all()
        normalized = _normalize_key(name)
        for user in leader:
            full = _normalize_key(f"{user.first_name} {user.last_name}")
            if normalized in {full, _normalize_key(user.first_name)}:
                return user

    default = db.scalar(
        select(User)
        .join(Role, User.role_id == Role.id)
        .where(Role.name == "Design Leader", User.is_active.is_(True))
        .order_by(User.last_name, User.first_name)
    )
    if default is None:
        raise ValueError("No Design Leader is configured in the system.")
    return default


def _default_stream(db: Session) -> Stream:
    stream = db.scalar(
        select(Stream).where(Stream.is_active.is_(True)).order_by(Stream.name)
    )
    if stream is None:
        raise ValueError("No active stream is configured in the system.")
    return stream


def _unique_project_code(db: Session, base_code: str) -> str:
    code = base_code.strip()[:50]
    if db.scalar(select(Project.id).where(Project.code == code)) is None:
        return code
    suffix = 1
    while True:
        candidate = f"{code}-{suffix}"[:50]
        if db.scalar(select(Project.id).where(Project.code == candidate)) is None:
            return candidate
        suffix += 1


def _apply_milestone_history(
    db: Session,
    project: Project,
    design_phase: str | None,
    progress: Decimal | None,
    summary: ImportSummary,
) -> None:
    milestones = db.scalars(
        select(Milestone)
        .where(Milestone.project_id == project.id)
        .order_by(Milestone.sort_order)
    ).all()
    summary.milestones_created += len(milestones)

    target_name = _resolve_milestone_name(design_phase)
    target_index = None
    if target_name:
        for idx, milestone in enumerate(milestones):
            if milestone.name == target_name:
                target_index = idx
                break
        if target_index is None and design_phase:
            pass  # unknown phase leaves defaults

    completed_through = target_index
    if progress is not None and progress >= Decimal("100"):
        completed_through = len(milestones) - 1

    now = datetime.now(timezone.utc)
    for idx, milestone in enumerate(milestones):
        if completed_through is not None and idx < completed_through:
            milestone.status = MilestoneStatus.completed
            milestone.completed_at = now
        elif completed_through is not None and idx == completed_through:
            if progress is not None and progress >= Decimal("100"):
                milestone.status = MilestoneStatus.completed
                milestone.completed_at = now
            else:
                milestone.status = MilestoneStatus.in_progress
        else:
            milestone.status = MilestoneStatus.not_started
            milestone.completed_at = None
        db.add(milestone)


def _create_project_from_row(
    db: Session,
    row: ParsedImportRow,
    *,
    summary: ImportSummary,
    duplicate_action: DuplicateAction,
    existing: Project | None,
) -> tuple[str, ImportRowPreview]:
    if row.errors:
        return "error", _preview_from_row(row)

    assert row.tool_number is not None
    assert row.customer is not None
    assert row.quoted_hours is not None

    if existing is not None:
        if duplicate_action == DuplicateAction.skip:
            preview = _preview_from_row(row, existing_project_id=existing.id)
            preview.status_label = ImportRowStatus.skipped
            preview.messages.append("Duplicate tool number skipped.")
            return "skipped", preview
        if duplicate_action == DuplicateAction.create:
            existing = None

    customer = _resolve_or_create_customer(db, row.customer, summary)
    contact = db.scalar(
        select(Contact)
        .where(Contact.customer_id == customer.id, Contact.is_primary.is_(True))
        .order_by(Contact.created_at)
    )
    if contact is None:
        contact = db.scalar(
            select(Contact).where(Contact.customer_id == customer.id).order_by(Contact.created_at)
        )
    assert contact is not None

    designer = _resolve_or_create_user(db, row.designer, summary=summary)
    surfacer = _resolve_or_create_user(db, row.surfacer, summary=summary)
    design_leader = _resolve_design_leader(db, row.design_leader)
    stream = _default_stream(db)

    part_description = row.part_description or f"Imported project {row.tool_number}"
    due_date = row.due_date or (date.today() + timedelta(days=90))
    code = row.code or row.tool_number
    status = row.status or ProjectStatus.not_started

    if existing is not None and duplicate_action == DuplicateAction.update:
        existing.customer_id = customer.id
        existing.customer_contact_id = contact.id
        existing.design_leader_id = design_leader.id
        existing.designer_id = designer.id if designer else None
        existing.surfacer_id = surfacer.id if surfacer else None
        existing.stream_id = stream.id
        existing.part_description = part_description
        existing.quoted_hours = row.quoted_hours
        existing.due_date = due_date
        existing.status = status
        if row.actual_hours is not None:
            existing.actual_hours = row.actual_hours
        db.add(existing)
        db.flush()
        _apply_milestone_history(
            db, existing, row.design_phase, row.progress, summary
        )
        existing.health = calculate_project_health(existing)
        db.add(existing)
        db.commit()
        preview = _preview_from_row(row, existing_project_id=existing.id)
        preview.status_label = ImportRowStatus.updated
        return "updated", preview

    code = _unique_project_code(db, code)
    mold_design_type = db.scalar(
        select(ProjectType).where(ProjectType.name == "Mold Design")
    )
    template = resolve_template_for_import(db, customer=customer)
    project = Project(
        tool_number=row.tool_number.strip(),
        part_description=part_description,
        customer_id=customer.id,
        customer_contact_id=contact.id,
        design_leader_id=design_leader.id,
        designer_id=designer.id if designer else None,
        surfacer_id=surfacer.id if surfacer else None,
        stream_id=stream.id,
        project_type_id=mold_design_type.id if mold_design_type else None,
        project_template_id=template.id,
        code=code,
        quoted_hours=row.quoted_hours,
        actual_hours=row.actual_hours or Decimal("0"),
        due_date=due_date,
        status=status,
        health=ProjectHealth.green,
        notes="Imported from historical workbook",
    )
    db.add(project)
    db.flush()

    create_milestones_from_template(db, project=project, template=template)
    db.flush()

    _apply_milestone_history(db, project, row.design_phase, row.progress, summary)
    if row.actual_hours is not None:
        project.actual_hours = row.actual_hours
    if row.status is not None:
        project.status = row.status
    project.health = calculate_project_health(project)
    db.add(project)
    db.commit()

    preview = _preview_from_row(row)
    preview.status_label = ImportRowStatus.imported
    return "imported", preview


def run_import(
    db: Session,
    upload_id: str,
    *,
    dry_run: bool,
    duplicate_action: DuplicateAction,
    progress_callback=None,
) -> tuple[ImportSummary, list[ImportRowPreview]]:
    path = get_upload_path(upload_id)
    rows = parse_workbook(path)
    summary = ImportSummary()
    error_log: list[ImportRowPreview] = []

    for index, row in enumerate(rows, start=1):
        existing = (
            _find_project_by_tool_number(db, row.tool_number)
            if row.tool_number and not row.errors
            else None
        )

        if dry_run:
            preview = _preview_from_row(row, existing_project_id=existing.id if existing else None)
            if preview.status_label == ImportRowStatus.error:
                summary.errors += 1
                error_log.append(preview)
            elif existing is not None:
                if duplicate_action == DuplicateAction.skip:
                    summary.projects_skipped += 1
                    preview.status_label = ImportRowStatus.skipped
                    preview.messages.append("Duplicate tool number would be skipped.")
                    error_log.append(preview)
                elif duplicate_action == DuplicateAction.update:
                    summary.projects_updated += 1
                    preview.status_label = ImportRowStatus.updated
                else:
                    summary.projects_imported += 1
                    preview.status_label = ImportRowStatus.imported
            else:
                summary.projects_imported += 1
                if not row.designer:
                    preview.messages.append("Designer missing; will be left empty unless provided.")
                preview.messages.append("Customer and users will be created if missing.")
            if progress_callback:
                progress_callback(index, len(rows))
            continue

        try:
            result, preview = _create_project_from_row(
                db,
                row,
                summary=summary,
                duplicate_action=duplicate_action,
                existing=existing,
            )
            if result == "imported":
                summary.projects_imported += 1
            elif result == "updated":
                summary.projects_updated += 1
            elif result == "skipped":
                summary.projects_skipped += 1
            elif result == "error":
                summary.errors += 1
                error_log.append(preview)
        except Exception as exc:
            db.rollback()
            summary.errors += 1
            preview = _preview_from_row(row, existing_project_id=existing.id if existing else None)
            preview.status_label = ImportRowStatus.error
            preview.messages.append(str(exc))
            error_log.append(preview)

        if progress_callback:
            progress_callback(index, len(rows))

    if dry_run:
        db.rollback()

    return summary, error_log


def error_log_to_csv(error_log: list[ImportRowPreview]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "row_number",
            "tool_number",
            "customer",
            "designer",
            "surfacer",
            "status_label",
            "messages",
        ]
    )
    for row in error_log:
        writer.writerow(
            [
                row.row_number,
                row.tool_number or "",
                row.customer or "",
                row.designer or "",
                row.surfacer or "",
                row.status_label.value,
                "; ".join(row.messages),
            ]
        )
    return buffer.getvalue()

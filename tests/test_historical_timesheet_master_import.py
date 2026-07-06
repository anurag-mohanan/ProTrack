import io
import time
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.models.models import TimesheetEntry
from app.services.historical_timesheet_master_import_service import (
    parse_master_workbook_bytes,
    save_master_upload,
    scan_master_workbook,
)
from tests.conftest import login


def build_master_workbook(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master"
    headers = [
        "Designer",
        "SL No.",
        "Date",
        "Project",
        "Customer",
        "Task",
        "Billable",
        "Hours",
        "Notes",
        "Month",
        "MonthNo",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    for offset, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            sheet.cell(offset, col, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def master_workbook_bytes():
    return build_master_workbook(
        [
            ["Binil JR", 1, "06-04-2026", "T-100", "TI", "Design", "Yes", 8, "Master import"],
            ["Unknown Person", 2, "06-04-2026", "T-100", "TI", "Design", "Yes", 4, "Skip me"],
        ]
    )


def test_parse_master_workbook(master_workbook_bytes, test_session_factory):
    with test_session_factory() as db:
        rows = parse_master_workbook_bytes(master_workbook_bytes, db=db)
    assert len(rows) == 2
    assert rows[0].designer == "Binil JR"
    assert rows[0].project_value == "T-100"
    assert rows[0].hours == Decimal("8")
    assert rows[0].is_np_row is False


def test_scan_master_workbook(master_workbook_bytes, test_session_factory, seeded_db):
    upload_id = "scan-test"
    save_master_upload(upload_id, "master.xlsx", master_workbook_bytes)
    with test_session_factory() as db:
        scan = scan_master_workbook(db, upload_id)
    assert scan.row_count == 2
    assert scan.designer_count == 2
    assert scan.designers[0].designer == "Binil JR"
    assert scan.designers[0].user_matched is True
    assert any(not designer.user_matched for designer in scan.designers)


def _poll_master_job(client, headers, job_id: str):
    for _ in range(40):
        response = client.get(
            f"/api/v1/imports/historical-timesheets/master/jobs/{job_id}",
            headers=headers,
        )
        assert response.status_code == 200
        body = response.json()
        if body["status"] in {"completed", "failed", "cancelled"}:
            return body
        time.sleep(0.1)
    raise AssertionError("Master import job did not complete in time")


def test_master_upload_and_import(client, master_workbook_bytes, test_session_factory):
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/master/upload",
        headers=headers,
        files={"file": ("master.xlsx", master_workbook_bytes, "application/vnd.ms-excel")},
    )
    assert upload.status_code == 200, upload.text
    upload_id = upload.json()["upload_id"]
    assert upload.json()["scan"]["row_count"] == 2

    run = client.post(
        "/api/v1/imports/historical-timesheets/master/run",
        headers=headers,
        json={"upload_id": upload_id, "designers": ["Binil JR"]},
    )
    assert run.status_code == 200, run.text
    job = _poll_master_job(client, headers, run.json()["job_id"])
    assert job["status"] == "completed"
    assert job["summary"]["rows_read"] == 1
    assert job["summary"]["rows_imported"] == 1
    assert job["summary"]["rows_skipped"] == 0
    assert job["summary"]["duplicate_check_disabled"] is True

    with test_session_factory() as db:
        entries = db.scalars(select(TimesheetEntry)).all()
        assert len(entries) == 1
        assert Decimal(str(entries[0].hours)) == Decimal("8")


def test_master_import_all_skips_unknown_designer(client, master_workbook_bytes, test_session_factory):
    headers = login(client, "admin@prosohm.com")
    upload = client.post(
        "/api/v1/imports/historical-timesheets/master/upload",
        headers=headers,
        files={"file": ("master.xlsx", master_workbook_bytes, "application/vnd.ms-excel")},
    )
    upload_id = upload.json()["upload_id"]

    run = client.post(
        "/api/v1/imports/historical-timesheets/master/run",
        headers=headers,
        json={"upload_id": upload_id},
    )
    job = _poll_master_job(client, headers, run.json()["job_id"])
    assert job["summary"]["rows_imported"] == 1
    assert job["summary"]["rows_skipped"] == 1

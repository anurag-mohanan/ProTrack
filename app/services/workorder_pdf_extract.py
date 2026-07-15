"""Heuristic customer workorder extraction (PDF + Excel) for Overview autofill.

Layouts differ by customer (CMT PDF, ABC Shop Order PDF, B&B Kick-Off xlsx).
Returns suggestions for human review — does not persist. Work order number is optional.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.core.exceptions import ProTrackValidationError

_MAX_FILE_BYTES = 12 * 1024 * 1024

# Colon/pipe only — hyphens appear inside tool/part numbers (e.g. CMT-2649).
_LABEL_VALUE = re.compile(
    r"(?m)^\s*(?P<label>[A-Za-z#][A-Za-z0-9 /#\.'%]{0,48}?)\s*[:|]\s*(?P<value>[^\n\r]{1,200})",
    re.IGNORECASE,
)


@dataclass
class WorkorderExtract:
    part_description: str | None = None
    work_order_number: str | None = None
    press_tonnage: str | None = None
    plastic_material: str | None = None
    cavity_count: int | None = None
    tool_type: str | None = None
    customer_specs: str | None = None
    warnings: list[str] = field(default_factory=list)
    source_chars: int = 0


def _clean(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = " ".join(str(value).replace("\u00a0", " ").split()).strip(" -:;|")
    return cleaned or None


def _norm_label(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", label.lower()).strip()


def _match_label_bucket(label: str) -> str | None:
    n = _norm_label(label)
    if not n:
        return None

    # Part description / name (priority field)
    if n in {"part description", "part name", "part desc"} or n.startswith("part name"):
        return "part_description"
    if n in {"description", "product name", "component"}:
        return "part_description"

    # Press / tonnage
    if any(
        token in n
        for token in (
            "press tonnage",
            "press tonnage primary",
            "primary press",
            "tonnage",
            "press size",
            "clamp",
            "machine ton",
        )
    ):
        return "press_tonnage"
    if n in {"press"} or n.endswith(" press"):
        return "press_tonnage"

    # Material / resin
    if any(
        token in n
        for token in (
            "plastic type",
            "plastic type main",
            "plastic material",
            "material",
            "mat l",
            "resin",
            "polymer",
            "family grade",
            "grade colour",
            "grade color",
        )
    ):
        return "plastic_material"
    if "mat" in n and ("family" in n or "grade" in n or "colour" in n or "color" in n):
        return "plastic_material"

    # Cavity (exclude steel/finish/temp lines that merely mention cavity)
    if _is_cavity_field_label(n):
        return "cavity_count"

    # Tool type
    if any(token in n for token in ("tool type", "mould type", "mold type", "tooling type", "type of tool")):
        return "tool_type"

    # Optional work-order / job refs (not required)
    if any(
        token in n
        for token in (
            "work order",
            "workorder",
            "shop order",
            "s o",
            "job number",
            "b b job number",
            "order no",
            "order number",
            "customer order",
        )
    ) or n in {"wo", "w o", "so"}:
        return "work_order_number"

    # Specs bucket for extra tooling facts
    if any(
        token in n
        for token in (
            "customer tool",
            "tool number",
            "tool #",
            "plant tool",
            "shrink",
            "program",
            "gate style",
            "manifold",
            "notes",
            "specification",
            "specs",
            "requirement",
            "remarks",
            "general notes",
        )
    ):
        return "customer_specs"
    return None


def _is_cavity_field_label(n: str) -> bool:
    """True only for labels that *mean* cavity count (not STEEL, CAVITY / FINISH, CAVITY)."""
    if any(
        token in n
        for token in (
            "steel",
            "finish",
            "temp",
            "radius",
            "running",
            "soft",
            "optic",
            "cut",
            "type",
            "core soft",
            "cavity steel",
            "cavity soft",
        )
    ):
        return False
    if n in {"cav", "cavity", "cavities", "cavitation", "no of cav", "no of cavity", "no of cavities"}:
        return True
    if any(
        token in n
        for token in (
            "cavitation",
            "cavity count",
            "number of cavity",
            "number of cavities",
            "no of cavities",
            "no of cavity",
        )
    ):
        return True
    if n.endswith(" cavity") or n.startswith("cavity ") or " cavities" in n:
        return True
    return False


def _parse_cavity(value: str) -> int | None:
    cleaned = _clean(value)
    if not cleaned:
        return None
    lower = cleaned.lower()
    word_map = {
        "one": 1,
        "single": 1,
        "two": 2,
        "three": 3,
        "four": 4,
        "five": 5,
        "six": 6,
        "eight": 8,
    }
    if lower in word_map:
        return word_map[lower]

    # "1+1" / "2+2 Cavity"
    plus_parts = re.findall(r"\d+", cleaned)
    if "+" in cleaned and len(plus_parts) >= 2:
        total = sum(int(part) for part in plus_parts)
        return total if 1 <= total <= 128 else None

    match = re.search(r"(\d{1,3})", cleaned)
    if not match:
        return None
    count = int(match.group(1))
    return count if 1 <= count <= 128 else None


def _parse_tonnage(value: str) -> str | None:
    match = re.search(
        r"(\d{2,5}(?:[.,]\d+)?)\s*(t|ton|tons|tonne|tonnes)?\b",
        value,
        re.IGNORECASE,
    )
    if not match:
        return None
    amount = match.group(1).replace(",", "")
    if "." in amount:
        amount = amount.split(".", 1)[0]
    if not amount.isdigit():
        return None
    tons = int(amount)
    if tons < 50 or tons > 10000:
        return None
    return f"{tons}T"


def _normalize_part_description(raw: str) -> tuple[str | None, int | None]:
    """Strip cavity prefixes like ``1+1 Cavity`` from CMT part description lines."""
    text = _clean(raw)
    if not text:
        return None, None
    cavity = None
    match = re.match(
        r"^((?:\d+\s*\+\s*)*\d+)\s+cavity\s+(.+)$",
        text,
        re.IGNORECASE,
    )
    if match:
        cavity = _parse_cavity(match.group(1))
        text = match.group(2).strip()
    return text[:255] or None, cavity


def _clean_material(raw: str) -> str | None:
    value = _clean(raw)
    if not value:
        return None
    # ABC: "1:R RESIN - S375AHW-600R BLACK" / "R RESIN - ..."
    value = re.sub(r"^\d+[a-z]?\)\s*", "", value, flags=re.IGNORECASE)
    value = re.sub(r"^[A-Z]\s*:", "", value).strip()
    value = re.sub(r"^(?:material\s+is\s+)", "", value, flags=re.IGNORECASE)
    value = re.sub(r"^(?:r\s+)?resin\s*-\s*", "", value, flags=re.IGNORECASE)
    return value[:150] or None


def _find_press_tonnage(text: str) -> str | None:
    patterns = (
        r"\b(\d{3,5})\s*[Tt]\s+[Pp]ress\b",
        r"\b[Pp]ress(?:\s*tonnage)?(?:\s*primary)?\s*[:#\-]?\s*(\d{3,5})\s*[Tt]\b",
        r"(?:press\s*)?tonnage(?:\s*primary)?\s*[:\-#]?\s*(\d{2,5}(?:[.,]\d+)?)\s*[Tt](?:on(?:ne)?s?)?\b",
        r"(?im)^(?=.*\bpress\b).{0,80}?\b(\d{3,5})\s*[Tt]\b",
        r"(?im)^.{0,80}?\b(\d{3,5})\s*[Tt]\b(?=.*\bpress\b).{0,40}$",
    )
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
        if not match:
            continue
        parsed = _parse_tonnage(match.group(1) + "T")
        if parsed:
            return parsed
    return None


def _apply_field(result: WorkorderExtract, bucket: str, raw: str) -> None:
    if bucket == "work_order_number":
        value = _clean(raw)
        if value and not result.work_order_number:
            result.work_order_number = value[:100]
        return

    if bucket == "press_tonnage":
        if result.press_tonnage:
            return
        parsed = _parse_tonnage(raw) or _find_press_tonnage(raw)
        if parsed:
            result.press_tonnage = parsed
        return

    if bucket == "plastic_material":
        if result.plastic_material:
            return
        material = _clean_material(raw)
        if material:
            result.plastic_material = material
        return

    if bucket == "cavity_count":
        if result.cavity_count is not None:
            return
        cavity = _parse_cavity(raw)
        if cavity is not None:
            result.cavity_count = cavity
        return

    if bucket == "tool_type":
        value = _clean(raw)
        if value and not result.tool_type:
            result.tool_type = value[:100]
        return

    if bucket == "part_description":
        if result.part_description:
            return
        part, cavity = _normalize_part_description(raw)
        if part:
            result.part_description = part
        if cavity is not None and result.cavity_count is None:
            result.cavity_count = cavity
        return

    if bucket == "customer_specs":
        value = _clean(raw)
        if not value:
            return
        if result.customer_specs:
            if value not in result.customer_specs:
                result.customer_specs = f"{result.customer_specs}\n{value}"[:4000]
        else:
            result.customer_specs = value[:4000]


def _scan_label_value_pairs(text: str, result: WorkorderExtract) -> None:
    for match in _LABEL_VALUE.finditer(text):
        bucket = _match_label_bucket(match.group("label"))
        if bucket:
            _apply_field(result, bucket, match.group("value"))


def _scan_customer_line_patterns(text: str, result: WorkorderExtract) -> None:
    """Formats that omit a colon between label and value (common on CMT / ABC PDFs)."""
    patterns: list[tuple[str, str]] = [
        # CMT Intermediate Approval (labels often lack a colon)
        (r"(?im)^\s*Part Description\s+(.+)$", "part_description"),
        (r"(?im)^\s*Tool Number\s+(.+)$", "customer_specs"),
        (r"(?im)^\s*Job Number\s+(\S+)$", "work_order_number"),
        # ABC Shop Order
        (r"(?im)^\s*PART NAME\s*:?\s*(.+)$", "part_description"),
        (r"(?im)^\s*#\s*Cav\.?\s*:?\s*(\d+)\b", "cavity_count"),
        (r"(?im)^\s*S\.?O\.?\s*#\s*:?\s*(\S+)$", "work_order_number"),
        (
            r"(?im)^\s*\d+[a-z]?\)\s*MAT'?L.*?COLOUR\s*\d*\s*:?\s*(.+)$",
            "plastic_material",
        ),
        (r"(?im)^\s*\d+[a-z]?\)\s*.*?RESIN\s*-\s*(.+)$", "plastic_material"),
        (r"(?im)^\s*\d+\)\s*PRIMARY PRESSES?.*?NUMBERS?\s*:?\s*(.+)$", "customer_specs"),
        # Generic
        (r"(?im)^\s*Part Name\s+(.+)$", "part_description"),
        (r"(?im)^\s*Tool Type\s+(.+)$", "tool_type"),
        (r"(?im)^\s*Press Tonnage(?:\s+Primary)?\s+(.+)$", "press_tonnage"),
        (r"(?im)^\s*Cavitation\s+(.+)$", "cavity_count"),
        (r"(?im)^\s*Plastic Type(?:\s*\(Main\))?\s+(.+)$", "plastic_material"),
        (r"(?im)^\s*Customer Tool Number\s+(.+)$", "customer_specs"),
    ]
    for pattern, bucket in patterns:
        if bucket == "part_description" and result.part_description:
            continue
        if bucket == "press_tonnage" and result.press_tonnage:
            continue
        if bucket == "plastic_material" and result.plastic_material:
            continue
        if bucket == "cavity_count" and result.cavity_count is not None:
            continue
        if bucket == "tool_type" and result.tool_type:
            continue
        match = re.search(pattern, text)
        if match:
            _apply_field(result, bucket, match.group(1))


def _scan_loose_patterns(text: str, result: WorkorderExtract) -> None:
    if not result.press_tonnage:
        found = _find_press_tonnage(text)
        if found:
            result.press_tonnage = found

    if not result.plastic_material:
        match = re.search(
            r"(?:plastic\s*(?:type|material)|material)\s*[:\-#]?\s*([A-Za-z0-9][A-Za-z0-9 \-/%+.]{1,100})",
            text,
            re.IGNORECASE,
        )
        if match:
            _apply_field(result, "plastic_material", match.group(1))

    if result.cavity_count is None:
        match = re.search(
            r"(?:#\s*cav\.?|cavitation|no\.?\s*of\s*cavit(?:y|ies)|cavit(?:y|ies))\s*[:\-#]?\s*([A-Za-z0-9+][A-Za-z0-9 +\-]{0,20})",
            text,
            re.IGNORECASE,
        )
        if match:
            _apply_field(result, "cavity_count", match.group(1))


def _scan_table_cells_pdf(content: bytes, result: WorkorderExtract) -> None:
    try:
        import pdfplumber
    except ImportError:
        return

    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages[:8]:
                for table in page.extract_tables() or []:
                    for row in table or []:
                        if not row:
                            continue
                        cells = [str(cell or "").strip() for cell in row]
                        if len(cells) >= 2:
                            bucket = _match_label_bucket(cells[0])
                            if bucket:
                                _apply_field(result, bucket, " ".join(cells[1:]))
    except Exception:  # noqa: BLE001
        result.warnings.append("Could not read some PDF tables; used page text only.")


def _unescape_pdf_literal(raw: bytes) -> str:
    text = raw.decode("latin-1", errors="ignore")
    text = text.replace(r"\n", "\n").replace(r"\r", "\r").replace(r"\t", "\t")
    text = text.replace(r"\(", "(").replace(r"\)", ")").replace(r"\\", "\\")
    return text


def _extract_text_stdlib(content: bytes) -> str:
    parts: list[str] = []
    for match in re.finditer(rb"\(((?:\\.|[^\\()\r\n])*)\)", content):
        value = _unescape_pdf_literal(match.group(1)).strip()
        if len(value) >= 2 and any(ch.isalnum() for ch in value):
            parts.append(value)

    for match in re.finditer(rb"\[(.*?)\]\s*TJ", content, flags=re.IGNORECASE | re.DOTALL):
        chunk_parts: list[str] = []
        for inner in re.finditer(rb"\(((?:\\.|[^\\()])*)\)", match.group(1)):
            value = _unescape_pdf_literal(inner.group(1)).strip()
            if value:
                chunk_parts.append(value)
        if chunk_parts:
            parts.append("".join(chunk_parts))

    latin = content.decode("latin-1", errors="ignore")
    for line in latin.splitlines():
        cleaned = "".join(ch if 32 <= ord(ch) < 127 else " " for ch in line)
        cleaned = " ".join(cleaned.split())
        if len(cleaned) < 6:
            continue
        lower = cleaned.lower()
        if any(
            token in lower
            for token in (
                "part description",
                "part name",
                "work order",
                "tonnage",
                "material",
                "cavity",
                "cavitation",
                "tool type",
                "plastic",
                "general notes",
                "press",
                "resin",
                "# cav",
            )
        ) or re.search(r"\b\d{3,5}\s*t\b", cleaned, re.IGNORECASE):
            parts.append(cleaned)

    seen: set[str] = set()
    ordered: list[str] = []
    for part in parts:
        key = part.lower()
        if key in seen:
            continue
        seen.add(key)
        ordered.append(part)
    return "\n".join(ordered).strip()


def _extract_text_pdfplumber(content: bytes) -> str:
    import pdfplumber

    chunks: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if not pdf.pages:
            raise ProTrackValidationError("PDF has no pages.")
        for page in pdf.pages[:12]:
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(text)
    return "\n".join(chunks).strip()


def _extract_text_pypdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if not reader.pages:
        raise ProTrackValidationError("PDF has no pages.")
    chunks: list[str] = []
    for page in reader.pages[:12]:
        text = page.extract_text() or ""
        if text.strip():
            chunks.append(text)
    return "\n".join(chunks).strip()


def extract_pdf_text(content: bytes) -> str:
    errors: list[str] = []
    for extractor, label in (
        (_extract_text_pdfplumber, "pdfplumber"),
        (_extract_text_pypdf, "pypdf"),
        (_extract_text_stdlib, "stdlib scrape"),
    ):
        try:
            combined = extractor(content)
            if combined:
                return combined
            errors.append(f"{label} found no text")
        except ImportError:
            errors.append(f"{label} not installed")
        except ProTrackValidationError:
            raise
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{label} failed: {exc}")

    if any("found no text" in item for item in errors):
        raise ProTrackValidationError(
            "No extractable text found. Use a text-based PDF/Excel (not a scanned image) "
            "or enter details manually."
        )
    raise ProTrackValidationError(
        f"Could not read workorder file ({'; '.join(errors)}). "
        "On the API host run: pip install -r requirements.txt then restart the API, "
        "or enter details manually."
    )


def extract_workorder_fields_from_text(text: str) -> WorkorderExtract:
    result = WorkorderExtract(source_chars=len(text))
    _scan_label_value_pairs(text, result)
    _scan_customer_line_patterns(text, result)
    _scan_loose_patterns(text, result)
    if not any(
        [
            result.part_description,
            result.press_tonnage,
            result.plastic_material,
            result.cavity_count is not None,
            result.tool_type,
            result.customer_specs,
        ]
    ):
        result.warnings.append(
            "No tooling fields matched. Check Part Name / Part Description, Press Tonnage, "
            "Plastic Material, Cavity, and Tool Type on the workorder, or enter them manually."
        )
    return result


def _extract_from_excel(content: bytes) -> WorkorderExtract:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ProTrackValidationError(
            "Excel workorder import requires openpyxl. "
            "On the API host run: pip install -r requirements.txt then restart the API."
        ) from exc

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ProTrackValidationError(f"Could not read Excel workorder: {exc}") from exc

    result = WorkorderExtract()
    text_chunks: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=120, max_col=16, values_only=True):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                nonempty = [cell for cell in cells if cell]
                if not nonempty:
                    continue
                text_chunks.append(" | ".join(nonempty))
                # Common B&B layout: label in col A, value in col B, alternates in later cols.
                label = cells[0] if cells else ""
                value = next((cell for cell in cells[1:] if cell), "")
                bucket = _match_label_bucket(label)
                if bucket and value:
                    if bucket == "press_tonnage" and "/" in value:
                        # "650T/720T" alternate list — take primary/first.
                        value = value.split("/", 1)[0]
                    _apply_field(result, bucket, value)
                elif len(nonempty) >= 2:
                    # Fallback when first cell is blank but second is label.
                    maybe_label = nonempty[0]
                    maybe_value = nonempty[1]
                    bucket = _match_label_bucket(maybe_label)
                    if bucket:
                        _apply_field(result, bucket, maybe_value)
    finally:
        workbook.close()

    joined = "\n".join(text_chunks)
    result.source_chars = len(joined)
    # Re-run text heuristics on flattened sheet text for any missed fields.
    text_result = extract_workorder_fields_from_text(joined)
    for field_name in (
        "part_description",
        "press_tonnage",
        "plastic_material",
        "tool_type",
        "customer_specs",
        "work_order_number",
    ):
        if not getattr(result, field_name) and getattr(text_result, field_name):
            setattr(result, field_name, getattr(text_result, field_name))
    if result.cavity_count is None and text_result.cavity_count is not None:
        result.cavity_count = text_result.cavity_count
    result.warnings.extend(text_result.warnings)
    return result


def extract_workorder_fields_from_pdf(content: bytes) -> WorkorderExtract:
    """Backward-compatible PDF entry point."""
    return extract_workorder_fields_from_file(content, filename="workorder.pdf")


def extract_workorder_fields_from_file(
    content: bytes,
    *,
    filename: str | None = None,
) -> WorkorderExtract:
    if not content:
        raise ProTrackValidationError("Workorder file is empty.")
    if len(content) > _MAX_FILE_BYTES:
        raise ProTrackValidationError("Workorder file exceeds the 12 MB limit.")

    suffix = Path(filename or "workorder.pdf").suffix.lower().lstrip(".")
    if suffix in {"xlsx", "xlsm"}:
        return _extract_from_excel(content)
    if suffix and suffix != "pdf":
        raise ProTrackValidationError("Only PDF and Excel (.xlsx/.xlsm) workorders are supported.")

    text = extract_pdf_text(content)
    result = extract_workorder_fields_from_text(text)
    try:
        import pdfplumber  # noqa: F401
    except ImportError:
        result.warnings.append(
            "pdfplumber is not installed on the API host; used fallback PDF text scrape. "
            "Install requirements for better accuracy: pip install -r requirements.txt"
        )
    _scan_table_cells_pdf(content, result)
    return result

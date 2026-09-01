"""Excel letterhead branding — company logo top-right + professional header."""

from pathlib import Path

from openpyxl import Workbook

from app.services.reporting.excel import letterhead as letterhead_module
from app.services.reporting.excel.letterhead import (
    add_company_logo_top_right,
    write_report_letterhead,
)


def test_write_report_letterhead_places_logo_top_right(tmp_path, monkeypatch):
    logo = tmp_path / "company-logo.png"
    # Minimal valid 1x1 PNG
    logo.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    monkeypatch.setattr(letterhead_module, "resolve_company_logo_path", lambda: logo)
    # letterhead re-exports from template — also patch template resolver used at runtime
    from app.services.reporting.excel import template as template_module

    monkeypatch.setattr(template_module, "COMPANY_LOGO_DIR", tmp_path)
    monkeypatch.setattr(template_module, "resolve_company_logo_path", lambda: logo)

    workbook = Workbook()
    sheet = workbook.active
    next_row = write_report_letterhead(
        sheet,
        company_name="Prosohm",
        report_title="Timesheet Report",
        period_label="July 2026",
        col_span=8,
    )
    assert next_row >= 6
    assert sheet.cell(row=1, column=1).value == "PROSOHM"
    assert sheet.cell(row=3, column=1).value == "Timesheet Report"
    assert len(sheet._images) == 1
    anchor = str(sheet._images[0].anchor)
    assert "G" in anchor or "H" in anchor or "7" in anchor or "8" in anchor


def test_add_company_logo_top_right_no_logo_is_safe(tmp_path, monkeypatch):
    from app.services.reporting.excel import template as template_module

    monkeypatch.setattr(template_module, "COMPANY_LOGO_DIR", tmp_path)
    monkeypatch.setattr(template_module, "resolve_company_logo_path", lambda: None)
    workbook = Workbook()
    sheet = workbook.active
    assert add_company_logo_top_right(sheet, col_span=6) is False
    assert len(sheet._images) == 0


def test_resolve_company_logo_path_finds_upload(tmp_path, monkeypatch):
    from app.services.reporting.excel import template as template_module

    logo = tmp_path / "company-logo.png"
    logo.write_bytes(b"not-a-real-png-but-exists")
    monkeypatch.setattr(template_module, "COMPANY_LOGO_DIR", tmp_path)
    assert template_module.resolve_company_logo_path() == logo

"""Timesheet reports must scope projects/designers to the selected reporting period."""

from __future__ import annotations

import uuid
from datetime import date
from decimal import Decimal

from app.models.enums import TeamRelationshipType, TimesheetStatus, WorkCategory
from app.models.models import Customer, Project, Team, TeamMember, Timesheet, TimesheetEntry, User
from app.services.reporting.designer_team_timesheet import build_designer_team_timesheet
from app.services.reporting.excel.designer_team_timesheet import generate_designer_team_timesheet_excel
from tests.conftest import IDS


def _setup_team_with_designer(session):
    team = Team(id=uuid.uuid4(), name="Period Scope Team", is_active=True)
    session.add(team)
    session.flush()
    designer = session.get(User, IDS["user_binil"])
    leader = session.get(User, IDS["user_anurag"])
    assert designer and leader
    designer.team_id = team.id
    leader.team_id = team.id
    session.add_all(
        [
            TeamMember(
                team_id=team.id,
                user_id=leader.id,
                relationship_type=TeamRelationshipType.team_leader,
                is_primary=True,
            ),
            TeamMember(
                team_id=team.id,
                user_id=designer.id,
                relationship_type=TeamRelationshipType.member,
                is_primary=True,
            ),
        ]
    )
    return team, designer


def _add_project(session, team: Team, tool_number: str) -> Project:
    customer = session.get(Customer, IDS["customer"])
    assert customer is not None
    project = Project(
        id=uuid.uuid4(),
        tool_number=tool_number,
        part_description=f"Part for {tool_number}",
        customer_id=customer.id,
        team_id=team.id,
        is_deleted=False,
        is_archived=False,
    )
    session.add(project)
    session.flush()
    return project


def _add_entry(session, designer: User, project: Project, entry_date: date, hours: str) -> None:
    timesheet = Timesheet(
        user_id=designer.id,
        week_start=entry_date,
        status=TimesheetStatus.approved,
    )
    session.add(timesheet)
    session.flush()
    session.add(
        TimesheetEntry(
            timesheet_id=timesheet.id,
            entry_date=entry_date,
            hours=Decimal(hours),
            work_category=WorkCategory.productive,
            is_billable=True,
            customer_id=project.customer_id,
            project_id=project.id,
        )
    )


def test_monthly_report_excludes_projects_without_period_activity(session):
    team, designer = _setup_team_with_designer(session)
    project_july = _add_project(session, team, "R-JULY")
    project_august = _add_project(session, team, "R-AUG")
    project_september = _add_project(session, team, "R-SEP")
    _add_entry(session, designer, project_july, date(2026, 7, 15), "8")
    _add_entry(session, designer, project_august, date(2026, 8, 10), "6")
    _add_entry(session, designer, project_september, date(2026, 9, 5), "5")
    session.commit()

    admin = session.get(User, IDS["user_admin"])
    assert admin
    payload = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 8, 15),
        team_id=team.id,
    )

    tool_numbers = {row.tool_number for row in payload.projects}
    assert tool_numbers == {"R-AUG"}
    august_row = payload.projects[0]
    assert float(august_row.actual_hours) == 6.0
    assert payload.project_count == 1
    assert float(payload.total_project_actual_hours) == 6.0


def test_monthly_report_excludes_zero_hour_designers(session):
    team, designer = _setup_team_with_designer(session)
    inactive = session.get(User, IDS["user_junior_designer"])
    assert inactive
    session.add(
        TeamMember(
            team_id=team.id,
            user_id=inactive.id,
            relationship_type=TeamRelationshipType.member,
            is_primary=False,
        )
    )
    project = _add_project(session, team, "R-ACT")
    _add_entry(session, designer, project, date(2026, 8, 12), "12")
    session.commit()

    admin = session.get(User, IDS["user_admin"])
    assert admin
    payload = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 8, 15),
        team_id=team.id,
    )

    designer_ids = {row.user_id for row in payload.designers}
    assert designer_ids == {designer.id}
    assert payload.designer_count == 1
    assert float(payload.total_designer_hours) == 12.0


def test_cross_month_project_shows_only_in_period_hours(session):
    team, designer = _setup_team_with_designer(session)
    project = _add_project(session, team, "R-MULTI")
    _add_entry(session, designer, project, date(2026, 7, 20), "8")
    _add_entry(session, designer, project, date(2026, 8, 5), "4")
    _add_entry(session, designer, project, date(2026, 9, 2), "3")
    session.commit()

    admin = session.get(User, IDS["user_admin"])
    assert admin

    july = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 7, 1),
        team_id=team.id,
    )
    august = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 8, 15),
        team_id=team.id,
    )

    assert len(july.projects) == 1
    assert float(july.projects[0].actual_hours) == 8.0
    assert len(august.projects) == 1
    assert float(august.projects[0].actual_hours) == 4.0


def test_designer_and_project_totals_reconcile(session):
    team, designer = _setup_team_with_designer(session)
    project_a = _add_project(session, team, "R-A")
    project_b = _add_project(session, team, "R-B")
    _add_entry(session, designer, project_a, date(2026, 8, 3), "10")
    _add_entry(session, designer, project_b, date(2026, 8, 4), "20")
    session.commit()

    admin = session.get(User, IDS["user_admin"])
    assert admin
    payload = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 8, 1),
        team_id=team.id,
    )

    project_sum = sum(float(row.actual_hours) for row in payload.projects)
    designer_sum = sum(float(row.total_hours) for row in payload.designers)
    assert project_sum == designer_sum == 30.0
    assert float(payload.total_designer_hours) == 30.0
    assert float(payload.total_project_actual_hours) == 30.0


def test_excel_matches_period_scoped_payload(session):
    team, designer = _setup_team_with_designer(session)
    project_aug = _add_project(session, team, "R-AUG-XLS")
    project_jul = _add_project(session, team, "R-JUL-XLS")
    _add_entry(session, designer, project_aug, date(2026, 8, 8), "18")
    _add_entry(session, designer, project_jul, date(2026, 7, 8), "8")
    session.commit()

    admin = session.get(User, IDS["user_admin"])
    assert admin
    payload = build_designer_team_timesheet(
        session,
        current_user=admin,
        report_id="monthly-timesheet",
        anchor=date(2026, 8, 1),
        team_id=team.id,
    )
    workbook = __import__("io").BytesIO(generate_designer_team_timesheet_excel(payload))
    from openpyxl import load_workbook

    wb = load_workbook(filename=workbook)
    projects = wb["Project Hours"]
    data_rows = [
        row
        for row in projects.iter_rows(min_row=1, max_row=projects.max_row, values_only=True)
        if row and row[0] in {"R-AUG-XLS", "R-JUL-XLS"}
    ]
    assert len(data_rows) == 1
    assert data_rows[0][0] == "R-AUG-XLS"
